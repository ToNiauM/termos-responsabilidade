"""Migração inicial: popula termos.db a partir de acervo.xlsx e geral.xlsx. Roda UMA vez.

    .venv/bin/python importar_planilhas.py [acervo.xlsx] [geral.xlsx]
"""
import sys
from pathlib import Path

from openpyxl import load_workbook

import db


def _linhas(caminho: Path, aba: str) -> list[dict]:
    wb = load_workbook(caminho, read_only=True, data_only=True)
    ws = wb[aba]
    it = ws.iter_rows(values_only=True)
    cab = [db._texto(c) for c in next(it)]
    saida = [dict(zip(cab, r)) for r in it if any(v is not None for v in r)]
    wb.close()
    return saida


def migrar(conn, acervo: Path, geral: Path) -> dict:
    avisos = []
    db.criar_esquema(conn)
    for t in ("atribuicoes", "pessoas", "localizacoes", "responsaveis"):
        conn.execute(f"DELETE FROM {t}")
    conn.commit()

    resumo_bens = db.importar_bens(conn, geral)  # acha a aba 'base' pelo cabeçalho

    for r in _linhas(acervo, "responsavel"):
        sigla = db._texto(r["ccustos"]).upper()
        if not sigla:
            continue
        conn.execute("INSERT OR REPLACE INTO responsaveis (ccustos, responsavel, email, matricula, funcao) VALUES (?,?,?,?,?)", (
            sigla, db._texto(r.get("responsavel")) or "(preencher)",
            db._texto(r.get("email")), db._texto(r.get("matricula")), db._texto(r.get("funcao"))))

    for r in _linhas(acervo, "ccustos"):
        loc, sigla = db._texto(r["localizacao"]), db._texto(r["ccustos"]).upper()
        if not loc or not sigla or sigla == "TERMOS INDIVIDUAIS":
            continue
        if not db.responsavel(conn, sigla):
            conn.execute("INSERT INTO responsaveis (ccustos, responsavel) VALUES (?, '(preencher)')", (sigla,))
            avisos.append(f"Centro de custo {sigla} sem responsável: criado como '(preencher)'.")
        conn.execute("INSERT OR REPLACE INTO localizacoes VALUES (?, ?)", (loc, sigla))

    nomes = {db._texto(r["responsavel"]).upper() for r in _linhas(geral, "nomes")}
    dados = _linhas(geral, "dados")
    nomes |= {db._texto(r["Nome"]).upper() for r in dados}
    conn.executemany("INSERT OR IGNORE INTO pessoas (nome) VALUES (?)", [(n,) for n in sorted(nomes) if n])

    for r in dados:
        nome, num = db._texto(r["Nome"]).upper(), db._numero(r["Patrimônio"])
        if not nome or num is None:
            continue
        if not db.buscar_bem(conn, int(num)):
            avisos.append(f"Bem {int(num)} de {nome} não existe na base; atribuição ignorada.")
            continue
        conn.execute("INSERT OR IGNORE INTO atribuicoes VALUES (?, ?)", (nome, int(num)))
    conn.commit()

    contar = lambda t: conn.execute(f"SELECT count(*) FROM {t}").fetchone()[0]
    return {"bens": resumo_bens["total"], "responsaveis": contar("responsaveis"),
            "localizacoes": contar("localizacoes"), "pessoas": contar("pessoas"),
            "atribuicoes": contar("atribuicoes"), "sem_centro": db.localizacoes_sem_centro(conn), "avisos": avisos}


if __name__ == "__main__":
    acervo = Path(sys.argv[1] if len(sys.argv) > 1 else "acervo.xlsx")
    geral = Path(sys.argv[2] if len(sys.argv) > 2 else "geral.xlsx")
    db.inicializar()
    conn = db.conectar()
    r = migrar(conn, acervo, geral)
    conn.close()
    for k, v in r.items():
        if k not in ("avisos", "sem_centro"):
            print(f"{k}: {v}")
    print("localizações ativas sem centro de custo:", ", ".join(r["sem_centro"]) or "nenhuma")
    for a in r["avisos"]:
        print("AVISO", a)
