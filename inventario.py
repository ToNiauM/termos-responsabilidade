"""Módulo de inventário: eventos (campanhas), salas, leituras por código de barras, sobras, relatório e
planilhas. Só dados: toda função recebe `conn` primeiro e não importa Flask (mesmo padrão de db.py).

Regras (spec 2026-09-15): `bens` é espelho do SPW e nunca muda aqui; "local sistema" = bens.localizacao,
"local inventário" = sala onde o bem foi lido; divergente = os dois diferem (calculado, nunca gravado)."""
import db
from db import ErroDeNegocio, _agora, _obrigatorio, _texto, _todos, _um, acrescentar_linha

CONSERVACAO = ("Bom", "Regular", "Ruim", "Inservível")
ROTULO_SITUACAO = {"localizado": "Localizado", "divergente": "Divergente", "pendente": "Não localizado"}


class BemNaoEncontrado(ErroDeNegocio):
    """Número lido não existe em `bens`: a tela oferece registrar como sobra."""

    def __init__(self, numero):
        super().__init__(f"Bem {numero} não está na base.")
        self.numero = numero


# ---------------------------------------------------------------- eventos
def evento_aberto(conn) -> dict | None:
    return _um(conn, "SELECT * FROM inventario_eventos WHERE encerrado_em IS NULL")


def eventos(conn) -> list[dict]:
    return _todos(conn, "SELECT * FROM inventario_eventos ORDER BY (encerrado_em IS NULL) DESC, aberto_em DESC, id DESC")


def evento(conn, id: int) -> dict | None:
    e = _um(conn, "SELECT * FROM inventario_eventos WHERE id = ?", id)
    if e:
        e["integrantes"] = [r[0] for r in conn.execute(
            "SELECT nome FROM inventario_integrantes WHERE evento_id = ? ORDER BY nome", (id,))]
        e["resumo"] = resumo(conn, id)
    return e


def _evento_aberto_ou_erro(conn, id: int) -> dict:
    e = _um(conn, "SELECT * FROM inventario_eventos WHERE id = ?", id)
    if not e:
        raise ErroDeNegocio("Evento de inventário não encontrado.")
    if e["encerrado_em"]:
        raise ErroDeNegocio("Evento encerrado: não aceita alterações.")
    return e


_COLS_SNAPSHOT = "numero, situacao, descricao, complemento, classificacao, localizacao"


def _fonte_bens(conn, evento_id: int) -> str:
    """De onde vêm os bens do evento: `bens` (aberto, ou encerrado antes de existir snapshot) ou a subconsulta
    do snapshot gravado no encerramento. evento_id é int e vai inline no SQL (não há injeção)."""
    e = _um(conn, "SELECT encerrado_em FROM inventario_eventos WHERE id = ?", evento_id)
    if not e or not e["encerrado_em"] or not conn.execute(
            "SELECT 1 FROM inventario_bens_encerrados WHERE evento_id = ? LIMIT 1", (evento_id,)).fetchone():
        return "bens"
    return f"(SELECT {_COLS_SNAPSHOT} FROM inventario_bens_encerrados WHERE evento_id = {int(evento_id)})"


def abrir_evento(conn, nome: str, descricao, integrantes: list, salas: list | None = None) -> int:
    """Um evento aberto por vez. salas=None → todas as localizações com bens ATIVO; lista → amostragem."""
    nome = _obrigatorio(nome, "Nome do evento")
    if evento_aberto(conn):
        raise ErroDeNegocio("Já existe um evento de inventário aberto; encerre-o antes de abrir outro.")
    nomes = sorted({" ".join(_texto(n).split()) for n in integrantes if _texto(n).strip()})
    if not nomes:
        raise ErroDeNegocio("Informe ao menos um integrante da comissão.")
    ativas = db.localizacoes_ativas(conn)
    escolhidas = ativas if salas is None else [s for s in ativas if s in set(salas)]
    if not escolhidas:
        raise ErroDeNegocio("Nenhuma sala com bens ativos no escopo do evento.")
    cur = conn.execute("INSERT INTO inventario_eventos (nome, descricao, aberto_em) VALUES (?,?,?)",
                       (nome, _texto(descricao) or None, _agora()))
    eid = cur.lastrowid
    conn.executemany("INSERT INTO inventario_integrantes VALUES (?,?)", [(eid, n) for n in nomes])
    conn.executemany("INSERT INTO inventario_salas (evento_id, localizacao) VALUES (?,?)", [(eid, s) for s in escolhidas])
    conn.commit()
    return eid


def encerrar_evento(conn, id: int) -> None:
    """Grava encerrado_em e congela os bens do evento (ativos das salas do escopo + todo bem lido) em
    inventario_bens_encerrados, para o relatório não mudar quando o export do SPW seguinte for carregado."""
    e = _um(conn, "SELECT * FROM inventario_eventos WHERE id = ?", id)
    if not e:
        raise ErroDeNegocio("Evento de inventário não encontrado.")
    if e["encerrado_em"]:
        return
    conn.execute("UPDATE inventario_eventos SET encerrado_em = ? WHERE id = ?", (_agora(), id))
    conn.execute(f"""INSERT OR IGNORE INTO inventario_bens_encerrados (evento_id, {_COLS_SNAPSHOT})
        SELECT ?, {_COLS_SNAPSHOT} FROM bens
        WHERE (situacao = 'ATIVO' AND localizacao IN (SELECT localizacao FROM inventario_salas WHERE evento_id = ?))
           OR numero IN (SELECT numero FROM inventario_leituras WHERE evento_id = ?)""", (id, id, id))
    conn.commit()


# ---------------------------------------------------------------- salas
def salas(conn, evento_id: int) -> list[dict]:
    """Por sala do escopo: bens ativos (total), localizados aqui, divergentes lidos aqui, pendentes."""
    B = _fonte_bens(conn, evento_id)
    linhas = _todos(conn, f"""
        SELECT s.localizacao, l.ccustos,
          (SELECT count(*) FROM {B} b WHERE b.localizacao = s.localizacao AND b.situacao = 'ATIVO') AS total,
          (SELECT count(*) FROM inventario_leituras r JOIN {B} b ON b.numero = r.numero
             WHERE r.evento_id = s.evento_id AND r.localizacao = s.localizacao
               AND b.localizacao = s.localizacao AND b.situacao = 'ATIVO') AS localizados,
          (SELECT count(*) FROM inventario_leituras r JOIN {B} b ON b.numero = r.numero
             WHERE r.evento_id = s.evento_id AND r.localizacao = s.localizacao AND b.localizacao <> s.localizacao
               AND b.situacao = 'ATIVO') AS divergentes
        FROM inventario_salas s LEFT JOIN localizacoes l ON l.localizacao = s.localizacao
        WHERE s.evento_id = ? ORDER BY s.localizacao""", evento_id)
    for s in linhas:
        s["pendentes"] = s["total"] - s["localizados"]
    return linhas


def _sala_ou_erro(conn, evento_id: int, localizacao: str) -> dict:
    s = _um(conn, "SELECT * FROM inventario_salas WHERE evento_id = ? AND localizacao = ?", evento_id, localizacao)
    if not s:
        raise ErroDeNegocio("Sala fora do escopo deste evento.")
    return s


def resumo(conn, evento_id: int) -> dict:
    """Progresso do evento pela quantidade de bens localizados (não há "concluir sala")."""
    ss = salas(conn, evento_id)
    bens = sum(s["total"] for s in ss)
    lidos = sum(s["localizados"] for s in ss)
    sobras = conn.execute("SELECT count(*) FROM inventario_sobras WHERE evento_id = ?", (evento_id,)).fetchone()[0]
    return {"salas": len(ss), "salas_iniciadas": sum(1 for s in ss if s["localizados"] or s["divergentes"]),
            "bens": bens, "lidos": lidos, "divergentes": sum(s["divergentes"] for s in ss),
            "pendentes": sum(s["pendentes"] for s in ss), "sobras": sobras,
            "pct_bens": round(100 * lidos / bens, 1) if bens else 0.0}


# ---------------------------------------------------------------- leituras
_LEITURA = "r.localizacao AS lido_em_sala, r.lido_em, r.integrante, r.conservacao, r.quem_usa, r.observacao, r.foto_url"


def ler(conn, evento_id: int, localizacao: str, numero: int, integrante: str) -> dict:
    """Registra (ou atualiza) a leitura do bem nesta sala. Qualquer bem cadastrado é aceito em qualquer sala;
    a divergência é só sinalizada (regra do sistema antigo)."""
    _evento_aberto_ou_erro(conn, evento_id)
    _sala_ou_erro(conn, evento_id, localizacao)
    if not conn.execute("SELECT 1 FROM inventario_integrantes WHERE evento_id = ? AND nome = ?", (evento_id, integrante)).fetchone():
        raise ErroDeNegocio("Escolha o integrante da comissão antes de ler.")
    bem = db.buscar_bem(conn, numero)
    if not bem:
        raise BemNaoEncontrado(numero)
    anterior = _um(conn, "SELECT * FROM inventario_leituras WHERE evento_id = ? AND numero = ?", evento_id, numero)
    agora = _agora()
    if anterior:
        conn.execute("UPDATE inventario_leituras SET localizacao = ?, lido_em = ?, integrante = ? WHERE id = ?",
                     (localizacao, agora, integrante, anterior["id"]))
    else:
        conn.execute("INSERT INTO inventario_leituras (evento_id, numero, localizacao, lido_em, integrante) VALUES (?,?,?,?,?)",
                     (evento_id, numero, localizacao, agora, integrante))
    conn.commit()
    return {"situacao": "localizado" if bem["localizacao"] == localizacao else "divergente", "bem": bem,
            "cadastrado_em": bem["localizacao"], "ativo": bem["situacao"] == "ATIVO", "reler": bool(anterior),
            "leitura_anterior": anterior, "lido_em": agora, "integrante": integrante}


def bens_da_sala(conn, evento_id: int, localizacao: str) -> dict:
    """bens: ativos cadastrados na sala (com a leitura do evento, se houver) e situacao_inv;
    trazidos: leituras feitas nesta sala de bens de outra sala ou não ativos; sobras: desta sala."""
    B = _fonte_bens(conn, evento_id)
    bens = _todos(conn, f"""
        SELECT b.*, {_LEITURA} FROM {B} b
        LEFT JOIN inventario_leituras r ON r.numero = b.numero AND r.evento_id = ?
        WHERE b.localizacao = ? AND b.situacao = 'ATIVO' ORDER BY b.numero""", evento_id, localizacao)
    for b in bens:
        b["situacao_inv"] = "pendente" if not b["lido_em"] else ("localizado" if b["lido_em_sala"] == localizacao else "divergente")
    trazidos = _todos(conn, f"""
        SELECT b.*, {_LEITURA} FROM inventario_leituras r JOIN {B} b ON b.numero = r.numero
        WHERE r.evento_id = ? AND r.localizacao = ? AND (b.localizacao <> ? OR b.situacao <> 'ATIVO')
        ORDER BY r.lido_em DESC""", evento_id, localizacao, localizacao)
    sobras = _todos(conn, "SELECT * FROM inventario_sobras WHERE evento_id = ? AND localizacao = ? ORDER BY id DESC",
                    evento_id, localizacao)
    return {"bens": bens, "trazidos": trazidos, "sobras": sobras}


def atualizar_leitura(conn, evento_id: int, numero: int, **campos) -> None:
    """Campos: conservacao, quem_usa, observacao, foto_url (só os presentes são gravados; '' vira NULL)."""
    _evento_aberto_ou_erro(conn, evento_id)
    permitidos = {"conservacao", "quem_usa", "observacao", "foto_url"}
    extra = set(campos) - permitidos
    if extra:
        raise ErroDeNegocio(f"Campo desconhecido: {', '.join(sorted(extra))}.")
    if "conservacao" in campos and campos["conservacao"] and campos["conservacao"] not in CONSERVACAO:
        raise ErroDeNegocio("Estado de conservação inválido.")
    if not _um(conn, "SELECT id FROM inventario_leituras WHERE evento_id = ? AND numero = ?", evento_id, numero):
        raise ErroDeNegocio("Leia o bem antes de preencher os dados.")
    for campo, valor in campos.items():
        conn.execute(f"UPDATE inventario_leituras SET {campo} = ? WHERE evento_id = ? AND numero = ?",
                     (_texto(valor) or None, evento_id, numero))
    conn.commit()


# ---------------------------------------------------------------- sobras
def registrar_sobra(conn, evento_id, localizacao, descricao, complemento, observacao, foto_url, integrante, exigir_foto=True) -> int:
    """Bem sem cadastro encontrado na sala. Foto obrigatória quando as fotos estão ativas (exigir_foto)."""
    _evento_aberto_ou_erro(conn, evento_id)
    _sala_ou_erro(conn, evento_id, localizacao)
    descricao = _obrigatorio(descricao, "Descrição")
    observacao = _obrigatorio(observacao, "Observação")
    integrante = _obrigatorio(integrante, "Integrante")
    if not conn.execute("SELECT 1 FROM inventario_integrantes WHERE evento_id = ? AND nome = ?", (evento_id, integrante)).fetchone():
        raise ErroDeNegocio("Escolha o integrante da comissão antes de ler.")
    if exigir_foto and not _texto(foto_url):
        raise ErroDeNegocio("A sobra precisa de foto.")
    cur = conn.execute("""INSERT INTO inventario_sobras (evento_id, localizacao, descricao, complemento, observacao, foto_url, integrante, criado_em)
                          VALUES (?,?,?,?,?,?,?,?)""",
                       (evento_id, localizacao, descricao, _texto(complemento) or None, observacao, _texto(foto_url), integrante, _agora()))
    conn.commit()
    return cur.lastrowid


def definir_foto_sobra(conn, sobra_id: int, foto_url: str) -> None:
    s = _um(conn, "SELECT evento_id FROM inventario_sobras WHERE id = ?", sobra_id)
    if not s:
        raise ErroDeNegocio("Sobra não encontrada.")
    _evento_aberto_ou_erro(conn, s["evento_id"])
    conn.execute("UPDATE inventario_sobras SET foto_url = ? WHERE id = ?", (_texto(foto_url), sobra_id))
    conn.commit()


def excluir_sobra(conn, evento_id: int, sobra_id: int) -> dict:
    """Só sobras podem ser apagadas (leituras de bens cadastrados, nunca). Devolve a sobra para apagar a foto."""
    _evento_aberto_ou_erro(conn, evento_id)
    s = _um(conn, "SELECT * FROM inventario_sobras WHERE evento_id = ? AND id = ?", evento_id, sobra_id)
    if not s:
        raise ErroDeNegocio("Sobra não encontrada.")
    conn.execute("DELETE FROM inventario_sobras WHERE id = ?", (sobra_id,))
    conn.commit()
    return s


# ---------------------------------------------------------------- relatório e planilha do evento
COLUNAS_XLSX = ["Patrimônio", "Descrição", "Complemento", "Classificação", "Local sistema", "Local inventário",
                "Situação", "Conservação", "Quem usa", "Observação", "Integrante", "Data/hora", "Foto", "Situação do bem"]
COLUNAS_SOBRAS = ["Sala", "Descrição", "Complemento", "Observação", "Integrante", "Data/hora", "Foto"]
_CAMPOS_REL = """b.numero AS numero, b.descricao, b.complemento, b.classificacao, b.localizacao AS local_sistema,
        r.localizacao AS local_inventario, r.lido_em, r.integrante, r.conservacao, r.quem_usa, r.observacao, r.foto_url,
        b.situacao AS situacao_bem"""


def relatorio(conn, evento_id: int, localizacao: str | None = None, situacao: str | None = None) -> list[dict]:
    """Uma linha por bem ativo das salas do escopo (ou da sala pedida), mais os lidos nela vindos de fora
    do escopo (ou de bens que deixaram de estar ATIVO). situacao filtra por localizado | divergente | pendente."""
    if not _um(conn, "SELECT id FROM inventario_eventos WHERE id = ?", evento_id):
        raise ErroDeNegocio("Evento de inventário não encontrado.")
    B = _fonte_bens(conn, evento_id)
    filtro_sala = ""
    params = [evento_id]
    if localizacao:
        filtro_sala = " AND (s.localizacao = ? OR r.localizacao = ?)"
        params += [localizacao, localizacao]
    params.append(evento_id)
    if localizacao:
        params.append(localizacao)
    linhas = _todos(conn, f"""
        SELECT {_CAMPOS_REL},
          CASE WHEN r.id IS NULL THEN 'pendente' WHEN r.localizacao = b.localizacao THEN 'localizado' ELSE 'divergente' END AS situacao_inv
        FROM inventario_salas s JOIN {B} b ON b.localizacao = s.localizacao AND b.situacao = 'ATIVO'
        LEFT JOIN inventario_leituras r ON r.evento_id = s.evento_id AND r.numero = b.numero
        WHERE s.evento_id = ?{filtro_sala}
        UNION ALL
        SELECT {_CAMPOS_REL},
          CASE WHEN r.localizacao = b.localizacao THEN 'localizado' ELSE 'divergente' END AS situacao_inv
        FROM inventario_leituras r JOIN {B} b ON b.numero = r.numero
        WHERE r.evento_id = ? AND (b.localizacao NOT IN (SELECT localizacao FROM inventario_salas WHERE evento_id = r.evento_id) OR b.situacao <> 'ATIVO')
          {"AND r.localizacao = ?" if localizacao else ""}
        ORDER BY local_sistema, numero""", *params)
    if situacao:
        linhas = [x for x in linhas if x["situacao_inv"] == situacao]
    return linhas


def _data_br(iso):
    return f"{iso[8:10]}/{iso[5:7]}/{iso[:4]} {iso[11:16]}" if iso else ""


def exportar_xlsx(conn, evento_id: int, destino, localizacao: str | None = None):
    from openpyxl import Workbook
    e = _um(conn, "SELECT nome FROM inventario_eventos WHERE id = ?", evento_id)
    if not e:
        raise ErroDeNegocio("Evento de inventário não encontrado.")
    wb = Workbook()
    ws = wb.active
    ws.title = "Bens"
    acrescentar_linha(ws, [f"{e['nome']} — gerado em {_data_br(_agora())} — {('sala ' + localizacao) if localizacao else 'todas as salas'}"])
    ws.append(COLUNAS_XLSX)
    for x in relatorio(conn, evento_id, localizacao):
        acrescentar_linha(ws, [x["numero"], x["descricao"], x["complemento"], x["classificacao"], x["local_sistema"], x["local_inventario"],
                                ROTULO_SITUACAO[x["situacao_inv"]], x["conservacao"], x["quem_usa"], x["observacao"], x["integrante"],
                                _data_br(x["lido_em"]), x["foto_url"], x["situacao_bem"]])
    ws2 = wb.create_sheet("Sobras")
    acrescentar_linha(ws2, [f"{e['nome']} — sobras (bens sem cadastro)"])
    ws2.append(COLUNAS_SOBRAS)
    sql = "SELECT * FROM inventario_sobras WHERE evento_id = ?" + (" AND localizacao = ?" if localizacao else "") + " ORDER BY localizacao, id"
    for s in _todos(conn, sql, *([evento_id, localizacao] if localizacao else [evento_id])):
        acrescentar_linha(ws2, [s["localizacao"], s["descricao"], s["complemento"], s["observacao"], s["integrante"], _data_br(s["criado_em"]), s["foto_url"]])
    wb.save(destino)
    return destino


# ---------------------------------------------------------------- planilha de cadastros (migração de inventários)
ABAS = {
    "inv_eventos": ["id", "nome", "descricao", "aberto_em", "encerrado_em"],
    "inv_integrantes": ["evento_id", "nome"],
    "inv_salas": ["evento_id", "localizacao"],
    "inv_leituras": ["evento_id", "numero", "localizacao", "lido_em", "integrante", "conservacao", "quem_usa", "observacao", "foto_url"],
    "inv_sobras": ["evento_id", "localizacao", "descricao", "complemento", "observacao", "foto_url", "integrante", "criado_em"],
}
_TABELA = {aba: "inventario_" + aba[4:] for aba in ABAS}


def exportar_abas(conn, wb) -> None:
    """Só acrescenta as abas inv_* quando já existe algum evento de inventário (planilhas de cadastro puras,
    sem inventário nunca feito, continuam só com as 4 abas de sempre)."""
    if not conn.execute("SELECT 1 FROM inventario_eventos LIMIT 1").fetchone():
        return
    for aba, colunas in ABAS.items():
        ws = wb.create_sheet(aba)
        ws.append(colunas)
        for linha in conn.execute(f"SELECT {', '.join(colunas)} FROM {_TABELA[aba]} ORDER BY {colunas[0]}, {colunas[1]}"):
            acrescentar_linha(ws, list(linha))


def _data_iso(valor, rotulo, linha, problemas, obrigatoria):
    """Aceita datetime do Excel, 'YYYY-MM-DD HH:MM:SS' ou 'YYYY-MM-DD' (vira 00:00:00). Vazio → None."""
    from datetime import datetime
    if valor is None or _texto(valor) == "":
        if obrigatoria:
            problemas.append(f"{linha}: {rotulo} vazia")
        return None
    if isinstance(valor, datetime):
        return valor.strftime("%Y-%m-%d %H:%M:%S")
    t = _texto(valor)
    for fmt in ("%Y-%m-%d %H:%M:%S", "%Y-%m-%d"):
        try:
            return datetime.strptime(t, fmt).strftime("%Y-%m-%d %H:%M:%S")
        except ValueError:
            pass
    problemas.append(f"{linha}: {rotulo} inválida ({t}); use AAAA-MM-DD HH:MM:SS")
    return None


def validar_abas(conn, brutos: dict) -> tuple[dict, list]:
    """brutos: {aba: [linhas dict com _linha]} (aba ausente = []). Devolve ({aba: [tuplas p/ INSERT]}, problemas)."""
    problemas: list[str] = []
    linhas: dict = {aba: [] for aba in ABAS}
    ids, abertos = set(), 0
    for r in brutos["inv_eventos"]:
        rot = f"inv_eventos linha {r['_linha']}"
        try:
            eid = int(r["id"])
        except (TypeError, ValueError):
            problemas.append(f"{rot}: id inválido"); continue
        if eid in ids:
            problemas.append(f"{rot}: id {eid} repetido"); continue
        nome = _texto(r["nome"])
        if not nome:
            problemas.append(f"{rot}: nome vazio"); continue
        aberto = _data_iso(r["aberto_em"], "aberto_em", rot, problemas, True)
        encerrado = _data_iso(r["encerrado_em"], "encerrado_em", rot, problemas, False)
        if encerrado is None and _texto(r["encerrado_em"]) == "":
            abertos += 1
        ids.add(eid)
        linhas["inv_eventos"].append((eid, nome, _texto(r["descricao"]) or None, aberto, encerrado))
    if abertos > 1:
        problemas.append("inv_eventos: mais de um evento aberto (sem encerrado_em)")

    def evento_ok(r, rot):
        try:
            eid = int(r["evento_id"])
        except (TypeError, ValueError):
            problemas.append(f"{rot}: evento_id inválido"); return None
        if eid not in ids:
            problemas.append(f"{rot}: evento {eid} não está na aba inv_eventos"); return None
        return eid

    vistos = set()
    for r in brutos["inv_integrantes"]:
        rot = f"inv_integrantes linha {r['_linha']}"
        eid, nome = evento_ok(r, rot), " ".join(_texto(r["nome"]).split())
        if eid is None or not nome or (eid, nome) in vistos:
            continue
        vistos.add((eid, nome))
        linhas["inv_integrantes"].append((eid, nome))
    vistos = set()
    for r in brutos["inv_salas"]:
        rot = f"inv_salas linha {r['_linha']}"
        eid, loc = evento_ok(r, rot), _texto(r["localizacao"])
        if eid is None or not loc or (eid, loc) in vistos:
            continue
        vistos.add((eid, loc))
        linhas["inv_salas"].append((eid, loc))
    vistos = set()
    for r in brutos["inv_leituras"]:
        rot = f"inv_leituras linha {r['_linha']}"
        eid = evento_ok(r, rot)
        num = db._numero(r["numero"])
        if eid is None:
            continue
        if num is None or num != int(num) or not db.buscar_bem(conn, int(num)):
            problemas.append(f"{rot}: bem {_texto(r['numero']) or '(vazio)'} não existe na base"); continue
        num = int(num)
        if (eid, num) in vistos:
            problemas.append(f"{rot}: bem {num} repetido no evento {eid}"); continue
        cons = _texto(r["conservacao"]) or None
        valido = True
        if cons and cons not in CONSERVACAO:
            problemas.append(f"{rot}: conservação inválida ({cons})"); valido = False
        loc, integ = _texto(r["localizacao"]), _texto(r["integrante"])
        if not loc or not integ:
            problemas.append(f"{rot}: localização e integrante são obrigatórios"); valido = False
        lido = _data_iso(r["lido_em"], "data lido_em", rot, problemas, True)
        if lido is None:
            valido = False
        if not valido:
            continue
        vistos.add((eid, num))
        linhas["inv_leituras"].append((eid, num, loc, lido, integ, cons, _texto(r["quem_usa"]) or None,
                                       _texto(r["observacao"]) or None, _texto(r["foto_url"]) or None))
    for r in brutos["inv_sobras"]:
        rot = f"inv_sobras linha {r['_linha']}"
        eid = evento_ok(r, rot)
        if eid is None:
            continue
        loc, desc, obs, integ = (_texto(r[c]) for c in ("localizacao", "descricao", "observacao", "integrante"))
        if not (loc and desc and obs and integ):
            problemas.append(f"{rot}: localização, descrição, observação e integrante são obrigatórios"); continue
        criado = _data_iso(r["criado_em"], "data criado_em", rot, problemas, True)
        if criado is None:
            continue
        linhas["inv_sobras"].append((eid, loc, desc, _texto(r["complemento"]) or None, obs, _texto(r["foto_url"]), integ, criado))
    return linhas, problemas


def substituir_tabelas(conn, linhas: dict) -> None:
    """Dentro da transação de db.importar_cadastros: apaga e regrava as 5 tabelas (ids de evento preservados)."""
    for aba in reversed(list(ABAS)):
        conn.execute(f"DELETE FROM {_TABELA[aba]}")
    conn.executemany("INSERT INTO inventario_eventos (id, nome, descricao, aberto_em, encerrado_em) VALUES (?,?,?,?,?)", linhas["inv_eventos"])
    conn.executemany("INSERT INTO inventario_integrantes VALUES (?,?)", linhas["inv_integrantes"])
    conn.executemany("INSERT INTO inventario_salas VALUES (?,?)", linhas["inv_salas"])
    conn.executemany("INSERT INTO inventario_leituras (evento_id, numero, localizacao, lido_em, integrante, conservacao, quem_usa, observacao, foto_url) VALUES (?,?,?,?,?,?,?,?,?)", linhas["inv_leituras"])
    conn.executemany("INSERT INTO inventario_sobras (evento_id, localizacao, descricao, complemento, observacao, foto_url, integrante, criado_em) VALUES (?,?,?,?,?,?,?,?)", linhas["inv_sobras"])
