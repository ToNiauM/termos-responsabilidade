"""Banco SQLite: esquema, importação de bens, consultas e cadastros.

Todas as funções recebem a conexão como primeiro argumento; quem abre e fecha é o chamador
(o Flask, por request; os testes, por fixture). Nenhuma função aqui usa Flask.
"""
import sqlite3
from datetime import date, datetime, timedelta
from pathlib import Path

from openpyxl import Workbook, load_workbook

import config
from migracoes_acesso import migrar

ESQUEMA = """
CREATE TABLE IF NOT EXISTS bens (
  numero        INTEGER PRIMARY KEY,
  situacao      TEXT NOT NULL,
  descricao     TEXT NOT NULL,
  complemento   TEXT,
  classificacao TEXT,
  localizacao   TEXT,
  data_entrada  TEXT,
  valor_compra  REAL,
  valor_atual   REAL
);
CREATE TABLE IF NOT EXISTS responsaveis (
  ccustos     TEXT PRIMARY KEY,
  responsavel TEXT NOT NULL,
  email       TEXT,
  matricula   TEXT,
  funcao      TEXT
);
CREATE TABLE IF NOT EXISTS localizacoes (
  localizacao TEXT PRIMARY KEY,
  ccustos     TEXT NOT NULL REFERENCES responsaveis(ccustos) ON UPDATE CASCADE
);
CREATE TABLE IF NOT EXISTS pessoas (
  nome      TEXT PRIMARY KEY,
  email     TEXT,
  matricula TEXT
);
CREATE TABLE IF NOT EXISTS atribuicoes (
  nome   TEXT    NOT NULL REFERENCES pessoas(nome) ON UPDATE CASCADE ON DELETE CASCADE,
  numero INTEGER NOT NULL REFERENCES bens(numero) DEFERRABLE INITIALLY DEFERRED,
  PRIMARY KEY (nome, numero)
);
CREATE TABLE IF NOT EXISTS textos (
  chave TEXT PRIMARY KEY,
  valor TEXT NOT NULL
);
CREATE TABLE IF NOT EXISTS processos_sei (
  id         INTEGER PRIMARY KEY,
  tipo       TEXT NOT NULL CHECK (tipo IN ('ccusto','individual','devolucao')),
  descricao  TEXT NOT NULL,
  numero_sei TEXT NOT NULL,
  vigente    INTEGER NOT NULL DEFAULT 0,
  criado_em  TEXT NOT NULL
);
CREATE UNIQUE INDEX IF NOT EXISTS processos_sei_vigente ON processos_sei(tipo) WHERE vigente = 1;
CREATE TABLE IF NOT EXISTS termos_emitidos (
  id            INTEGER PRIMARY KEY,
  tipo          TEXT NOT NULL,
  chave         TEXT NOT NULL,
  processo_id   INTEGER NOT NULL REFERENCES processos_sei(id),
  documento_sei TEXT,
  bloco_sei     TEXT,
  email_enviado_em TEXT,
  emitido_em    TEXT NOT NULL,
  quantidade    INTEGER NOT NULL,
  valor_total   REAL NOT NULL
);
CREATE TABLE IF NOT EXISTS termos_emitidos_bens (
  termo_id    INTEGER NOT NULL REFERENCES termos_emitidos(id) ON DELETE CASCADE,
  numero      INTEGER NOT NULL,
  descricao   TEXT, complemento TEXT, localizacao TEXT, valor_atual REAL,
  PRIMARY KEY (termo_id, numero)
);
CREATE TABLE IF NOT EXISTS importacoes (
  id            INTEGER PRIMARY KEY,
  importado_em  TEXT NOT NULL,
  arquivo       TEXT,
  total INTEGER NOT NULL, ativos INTEGER NOT NULL,
  novos INTEGER NOT NULL, removidos INTEGER NOT NULL, movidos INTEGER NOT NULL, situacao INTEGER NOT NULL
);
CREATE TABLE IF NOT EXISTS importacoes_mudancas (
  importacao_id INTEGER NOT NULL REFERENCES importacoes(id) ON DELETE CASCADE,
  numero        INTEGER NOT NULL,
  tipo          TEXT NOT NULL CHECK (tipo IN ('novo','removido','movido','situacao')),
  de            TEXT, para TEXT,
  descricao     TEXT
);
CREATE TABLE IF NOT EXISTS robo_execucoes (
  id            INTEGER PRIMARY KEY,
  iniciado_em   TEXT NOT NULL,
  terminado_em  TEXT NOT NULL,
  resultado     TEXT NOT NULL CHECK (resultado IN ('importado','sem_mudanca','erro')),
  hash          TEXT,
  importacao_id INTEGER REFERENCES importacoes(id) ON DELETE SET NULL,
  mensagem      TEXT
);
CREATE TABLE IF NOT EXISTS inventario_eventos (
  id           INTEGER PRIMARY KEY,
  nome         TEXT NOT NULL,
  descricao    TEXT,
  aberto_em    TEXT NOT NULL,
  encerrado_em TEXT
);
CREATE TABLE IF NOT EXISTS inventario_integrantes (
  evento_id INTEGER NOT NULL REFERENCES inventario_eventos(id) ON DELETE CASCADE,
  nome      TEXT NOT NULL,
  PRIMARY KEY (evento_id, nome)
);
CREATE TABLE IF NOT EXISTS inventario_salas (
  evento_id    INTEGER NOT NULL REFERENCES inventario_eventos(id) ON DELETE CASCADE,
  localizacao  TEXT NOT NULL,
  PRIMARY KEY (evento_id, localizacao)
);
CREATE TABLE IF NOT EXISTS inventario_leituras (
  id          INTEGER PRIMARY KEY,
  evento_id   INTEGER NOT NULL REFERENCES inventario_eventos(id) ON DELETE CASCADE,
  numero      INTEGER NOT NULL,
  localizacao TEXT NOT NULL,
  lido_em     TEXT NOT NULL,
  integrante  TEXT NOT NULL,
  conservacao TEXT CHECK (conservacao IN ('Bom','Regular','Ruim','Inservível')),
  quem_usa    TEXT,
  observacao  TEXT,
  fotos_seq   INTEGER NOT NULL DEFAULT 0,   -- maior nfoto já usado nesta leitura (não reaproveita após apagar foto)
  UNIQUE (evento_id, numero)
);
CREATE TABLE IF NOT EXISTS inventario_sobras (
  id          INTEGER PRIMARY KEY,
  evento_id   INTEGER NOT NULL REFERENCES inventario_eventos(id) ON DELETE CASCADE,
  localizacao TEXT NOT NULL,
  descricao   TEXT NOT NULL,
  complemento TEXT,
  observacao  TEXT NOT NULL,
  foto_url    TEXT NOT NULL,
  integrante  TEXT NOT NULL,
  criado_em   TEXT NOT NULL
);
CREATE TABLE IF NOT EXISTS inventario_bens_encerrados (
  evento_id     INTEGER NOT NULL REFERENCES inventario_eventos(id) ON DELETE CASCADE,
  numero        INTEGER NOT NULL,
  situacao      TEXT,
  descricao     TEXT,
  complemento   TEXT,
  classificacao TEXT,
  localizacao   TEXT,
  PRIMARY KEY (evento_id, numero)
);
CREATE TABLE IF NOT EXISTS inventario_fotos (
  evento_id INTEGER NOT NULL,
  numero    INTEGER NOT NULL,
  nfoto     INTEGER NOT NULL,
  url       TEXT NOT NULL,
  criado_em TEXT NOT NULL,
  PRIMARY KEY (evento_id, numero, nfoto),
  FOREIGN KEY (evento_id, numero) REFERENCES inventario_leituras(evento_id, numero) ON DELETE CASCADE
);
CREATE TABLE IF NOT EXISTS usuarios (
  id            INTEGER PRIMARY KEY,
  login         TEXT NOT NULL UNIQUE,
  email         TEXT,                        -- opcional; também serve para entrar
  nome          TEXT NOT NULL,
  senha_hash    TEXT NOT NULL,
  ativo         INTEGER NOT NULL DEFAULT 1,
  trocar_senha  INTEGER NOT NULL DEFAULT 0,
  falhas        INTEGER NOT NULL DEFAULT 0,
  bloqueado_ate TEXT,
  criado_em     TEXT NOT NULL,
  ultimo_acesso TEXT
);
"""


class ErroDeNegocio(Exception):
    """Erro que vira mensagem para o usuário (flash), não traceback."""


TIPOS_TERMO = ("ccusto", "individual", "devolucao")
ROTULO_TIPO = {"ccusto": "termos por centro de custo", "individual": "termos individuais",
               "devolucao": "termos de devolução"}


def _agora() -> str:
    return datetime.now().strftime("%Y-%m-%d %H:%M:%S")


def conectar(caminho: Path | None = None) -> sqlite3.Connection:
    conn = sqlite3.connect(str(caminho or config.caminho_db()), timeout=30)
    conn.row_factory = sqlite3.Row
    conn.execute("PRAGMA foreign_keys = ON")
    return conn


def _colunas(conn, tabela: str) -> list[str]:
    return [r[1] for r in conn.execute(f"PRAGMA table_info({tabela})")]


def criar_esquema(conn: sqlite3.Connection) -> None:
    conn.executescript(ESQUEMA)
    # Bancos criados antes de 2026-09-16: tratamento (Prezado/Prezada) saiu; pessoas ganhou e-mail e matrícula.
    if "tratamento" in _colunas(conn, "responsaveis"):
        conn.execute("ALTER TABLE responsaveis DROP COLUMN tratamento")
    if "email" not in _colunas(conn, "pessoas"):
        conn.execute("ALTER TABLE pessoas ADD COLUMN email TEXT")
        conn.execute("ALTER TABLE pessoas ADD COLUMN matricula TEXT")
    if "bloco_sei" not in _colunas(conn, "termos_emitidos"):
        conn.execute("ALTER TABLE termos_emitidos ADD COLUMN bloco_sei TEXT")
        conn.execute("ALTER TABLE termos_emitidos ADD COLUMN email_enviado_em TEXT")
    if "fotos_seq" not in _colunas(conn, "inventario_leituras"):
        conn.execute("ALTER TABLE inventario_leituras ADD COLUMN fotos_seq INTEGER NOT NULL DEFAULT 0")
    # Fase 3 (2026-09-17): a foto única da leitura virou a tabela inventario_fotos (várias por bem).
    if "foto_url" in _colunas(conn, "inventario_leituras"):
        conn.execute("""INSERT OR IGNORE INTO inventario_fotos (evento_id, numero, nfoto, url, criado_em)
                        SELECT evento_id, numero, 1, foto_url, lido_em FROM inventario_leituras
                        WHERE foto_url IS NOT NULL AND foto_url <> ''""")
        conn.execute("ALTER TABLE inventario_leituras DROP COLUMN foto_url")
    # Fase 4b: e-mail opcional do usuário (bancos criados antes dele não têm a coluna).
    if "email" not in _colunas(conn, "usuarios"):
        conn.execute("ALTER TABLE usuarios ADD COLUMN email TEXT")
    conn.execute("CREATE UNIQUE INDEX IF NOT EXISTS usuarios_email ON usuarios(email) WHERE email IS NOT NULL")
    # Evento com encerrado_em vazio ("" em vez de NULL, visto em produção em 2026-09-17) não é aberto nem encerrado.
    conn.execute("UPDATE inventario_eventos SET encerrado_em = NULL WHERE encerrado_em = ''")
    conn.commit()
    # Fase 5A: perfil único de usuários.perfil vira funções (usuarios_funcoes); a comissão de
    # inventário ganha identidade de usuário quando o nome não é ambíguo. Reserva a própria transação.
    migrar(conn)


def inicializar() -> None:
    """Primeira execução do programa: pastas + esquema vazio."""
    config.preparar_pastas()
    conn = conectar()
    try:
        criar_esquema(conn)
    finally:
        conn.close()


COLUNAS_EXPORT = {
    "Número Bem": "numero",
    "Situação": "situacao",
    "Descrição": "descricao",
    "Complemento": "complemento",
    "Classificação Contábil": "classificacao",
    "Localização": "localizacao",
    "Data Entrada": "data_entrada",
    "Valor Compra": "valor_compra",
    "Valor Atual": "valor_atual",
}


class ImportacaoInvalida(ErroDeNegocio):
    pass


def _texto(v) -> str:
    if v is None:
        return ""
    if isinstance(v, (datetime, date)):
        return v.strftime("%d/%m/%Y")
    return " ".join(str(v).split())


def _numero(v):
    try:
        return float(v) if v not in (None, "") else None
    except (TypeError, ValueError):
        return None


def _aba_do_export(wb):
    for ws in wb.worksheets:
        cabecalho = [_texto(c) for c in next(ws.iter_rows(min_row=1, max_row=1, values_only=True), ())]
        if "Número Bem" in cabecalho:
            return ws, cabecalho
    raise ImportacaoInvalida("Nenhuma aba com a coluna 'Número Bem'. Envie o export do sistema de patrimônio.")


def importar_bens(conn: sqlite3.Connection, arquivo, nome_arquivo: str | None = None) -> dict:
    """Substitui a tabela `bens` pelo conteúdo do export. Tudo ou nada.

    Devolve {"total", "ativos", "sem_centro", "importacao_id", "novos", "removidos", "movidos",
    "situacao"}. Levanta ImportacaoInvalida (e não altera nada) se faltar coluna ou se algum bem
    atribuído a pessoa deixar de existir.
    """
    try:
        wb = load_workbook(arquivo, read_only=True, data_only=True)
    except Exception:
        raise ImportacaoInvalida("Arquivo inválido: envie o export do sistema em .xlsx.")

    try:
        ws, cabecalho = _aba_do_export(wb)
        faltando = [c for c in COLUNAS_EXPORT if c not in cabecalho]
        if faltando:
            raise ImportacaoInvalida("Colunas ausentes no export: " + ", ".join(faltando))
        indice = {campo: cabecalho.index(col) for col, campo in COLUNAS_EXPORT.items()}

        linhas = []
        for r in ws.iter_rows(min_row=2, values_only=True):
            num = _numero(r[indice["numero"]])
            if num is None:
                continue
            linhas.append((
                int(num), _texto(r[indice["situacao"]]), _texto(r[indice["descricao"]]),
                _texto(r[indice["complemento"]]), _texto(r[indice["classificacao"]]),
                _texto(r[indice["localizacao"]]), _texto(r[indice["data_entrada"]]),
                _numero(r[indice["valor_compra"]]), _numero(r[indice["valor_atual"]]),
            ))
    except ImportacaoInvalida:
        raise
    except Exception:
        raise ImportacaoInvalida("Não consegui ler a planilha; o arquivo pode estar corrompido.")
    finally:
        wb.close()

    antes = {r["numero"]: (r["situacao"], r["localizacao"], r["descricao"])
             for r in conn.execute("SELECT numero, situacao, localizacao, descricao FROM bens")}

    try:
        conn.execute("DELETE FROM bens")
        conn.executemany("INSERT INTO bens VALUES (?,?,?,?,?,?,?,?,?)", linhas)
        orfaos = [str(r[0]) for r in conn.execute(
            "SELECT numero FROM atribuicoes WHERE numero NOT IN (SELECT numero FROM bens) ORDER BY numero")]
        if orfaos:
            raise ImportacaoInvalida(
                "O export não traz bens que estão atribuídos a pessoas: " + ", ".join(orfaos)
                + ". Remova a atribuição na aba Pessoas ou use um export completo.")

        mudancas = _mudancas(antes, linhas)
        contagem = {t: sum(1 for m in mudancas if m[1] == t) for t in ("novo", "removido", "movido", "situacao")}
        cur = conn.execute(
            "INSERT INTO importacoes (importado_em, arquivo, total, ativos, novos, removidos, movidos, situacao) VALUES (?,?,?,?,?,?,?,?)",
            (_agora(), nome_arquivo, len(linhas), sum(1 for l in linhas if l[1] == "ATIVO"),
             contagem["novo"], contagem["removido"], contagem["movido"], contagem["situacao"]))
        conn.executemany("INSERT INTO importacoes_mudancas VALUES (?,?,?,?,?,?)",
                         [(cur.lastrowid, n, t, de, para, desc) for n, t, de, para, desc in mudancas])
        importacao_id = cur.lastrowid
        conn.commit()
    except sqlite3.IntegrityError:
        conn.rollback()
        raise ImportacaoInvalida("O export tem número de bem repetido; corrija a planilha e envie de novo.")
    except Exception:
        conn.rollback()
        raise

    total = conn.execute("SELECT count(*) FROM bens").fetchone()[0]
    ativos = conn.execute("SELECT count(*) FROM bens WHERE situacao='ATIVO'").fetchone()[0]
    return {"total": total, "ativos": ativos, "sem_centro": localizacoes_sem_centro(conn),
            "importacao_id": importacao_id, "novos": contagem["novo"], "removidos": contagem["removido"],
            "movidos": contagem["movido"], "situacao": contagem["situacao"]}


def localizacoes_sem_centro(conn: sqlite3.Connection) -> list[str]:
    """Localizações de bens ATIVOS que não têm centro de custo mapeado (ignora bens atribuídos a pessoa)."""
    return [r[0] for r in conn.execute(
        "SELECT DISTINCT localizacao FROM bens WHERE situacao='ATIVO' AND localizacao <> '' "
        "AND localizacao NOT IN (SELECT localizacao FROM localizacoes) "
        "AND numero NOT IN (SELECT numero FROM atribuicoes) ORDER BY localizacao")]


def localizacoes_ativas(conn) -> list[str]:
    """Localizações distintas com bens ATIVO (as "salas" que um inventário pode conferir)."""
    return [r[0] for r in conn.execute(
        "SELECT DISTINCT localizacao FROM bens WHERE situacao = 'ATIVO' AND localizacao <> '' ORDER BY localizacao")]


def _mudancas(antes: dict, linhas: list[tuple]) -> list[tuple]:
    """[(numero, tipo, de, para, descricao)] comparando a tabela anterior com o export.
    linhas = tuplas na ordem de INSERT INTO bens (numero, situacao, descricao, complemento, classificacao, localizacao, ...)."""
    depois = {l[0]: (l[1], l[5], l[2]) for l in linhas}
    m = []
    for n, (sit, loc, desc) in depois.items():
        if n not in antes:
            m.append((n, "novo", None, loc, desc))
            continue
        sit0, loc0, _ = antes[n]
        if sit != sit0:
            m.append((n, "situacao", sit0, sit, desc))
        if loc != loc0:
            m.append((n, "movido", loc0, loc, desc))
    for n, (sit0, loc0, desc0) in antes.items():
        if n not in depois:
            m.append((n, "removido", loc0, None, desc0))
    return sorted(m, key=lambda x: (x[0], x[1]))


def importacoes(conn, limite: int = 20) -> list[dict]:
    return _todos(conn, "SELECT * FROM importacoes ORDER BY importado_em DESC, id DESC LIMIT ?", limite)


def importacao(conn, id: int) -> dict | None:
    i = _um(conn, "SELECT * FROM importacoes WHERE id = ?", id)
    if i:
        i["mudancas"] = _todos(conn, "SELECT * FROM importacoes_mudancas WHERE importacao_id = ? ORDER BY tipo, numero", id)
    return i


# ---------------------------------------------------------------- robô do SPW (importar_spw.py)

def registrar_execucao_robo(conn, iniciado_em: str, resultado: str, hash: str | None = None,
                            importacao_id: int | None = None, mensagem: str | None = None) -> int:
    """Uma linha por execução do robô. resultado: importado | sem_mudanca | erro."""
    cur = conn.execute(
        "INSERT INTO robo_execucoes (iniciado_em, terminado_em, resultado, hash, importacao_id, mensagem) VALUES (?,?,?,?,?,?)",
        (iniciado_em, _agora(), resultado, hash, importacao_id, mensagem))
    conn.commit()
    return cur.lastrowid


def execucoes_robo(conn, limite: int = 10) -> list[dict]:
    return _todos(conn, "SELECT * FROM robo_execucoes ORDER BY iniciado_em DESC, id DESC LIMIT ?", limite)


def ultimo_hash_robo(conn) -> str | None:
    """Hash das linhas da última execução que chegou ao fim (importou ou viu que nada mudou)."""
    r = _um(conn, "SELECT hash FROM robo_execucoes WHERE resultado IN ('importado','sem_mudanca') ORDER BY iniciado_em DESC, id DESC LIMIT 1")
    return r["hash"] if r else None


def importacao_desatualizada(conn, agora: datetime | None = None, dias: int = 4) -> bool:
    """True se nem a última importação de bens nem a última execução saudável do robô (resultado
    'importado' ou 'sem_mudanca') aconteceram nos últimos `dias` dias — ou se nenhuma das duas nunca
    aconteceu. Um robô que só confirma 'sem_mudanca' repetidamente não deve, sozinho, virar alerta."""
    candidatos = []
    ultima = importacoes(conn, 1)
    if ultima:
        candidatos.append(ultima[0]["importado_em"])
    robo = _um(conn, "SELECT iniciado_em FROM robo_execucoes WHERE resultado IN ('importado','sem_mudanca') "
                     "ORDER BY iniciado_em DESC, id DESC LIMIT 1")
    if robo:
        candidatos.append(robo["iniciado_em"])
    if not candidatos:
        return True
    referencia = max(datetime.strptime(c, "%Y-%m-%d %H:%M:%S") for c in candidatos)
    return (agora or datetime.now()) - referencia > timedelta(days=dias)


def historico_do_bem(conn, numero: int) -> dict:
    return {
        "mudancas": _todos(conn, """
            SELECT m.*, i.importado_em FROM importacoes_mudancas m JOIN importacoes i ON i.id = m.importacao_id
            WHERE m.numero = ? ORDER BY i.importado_em DESC, m.tipo""", numero),
        "termos": _todos(conn, """
            SELECT t.id, t.tipo, t.chave, t.emitido_em, t.documento_sei FROM termos_emitidos_bens b
            JOIN termos_emitidos t ON t.id = b.termo_id WHERE b.numero = ? ORDER BY t.emitido_em DESC""", numero),
    }


def _todos(conn, sql, *args) -> list[dict]:
    return [dict(r) for r in conn.execute(sql, args)]


def _um(conn, sql, *args) -> dict | None:
    r = conn.execute(sql, args).fetchone()
    return dict(r) if r else None


def centros(conn) -> list[dict]:
    return _todos(conn, "SELECT * FROM responsaveis ORDER BY ccustos")


def listar_cadastros(conn, aba, filtros):
    """Busca no conjunto completo, ordenação permitida e paginação no servidor."""
    import unicodedata

    def normalizar(valor):
        return "".join(c for c in unicodedata.normalize("NFD", str(valor or "").casefold())
                       if not unicodedata.combining(c))

    conn.create_function("cadastro_busca", 1, normalizar, deterministic=True)
    fontes = {
        "responsaveis": ("SELECT *, ccustos AS chave FROM responsaveis", ["ccustos", "responsavel", "funcao"]),
        "pessoas": ("""SELECT p.nome, p.nome AS chave, p.email, p.matricula, COUNT(a.numero) AS quantidade
                        FROM pessoas p LEFT JOIN atribuicoes a ON a.nome = p.nome GROUP BY p.nome""",
                    ["nome", "email", "quantidade"]),
        "localizacoes": ("""SELECT localizacao, ccustos, localizacao AS chave FROM localizacoes
            UNION ALL SELECT DISTINCT b.localizacao, '', b.localizacao FROM bens b
            WHERE b.situacao = 'ATIVO' AND b.localizacao <> ''
            AND b.localizacao NOT IN (SELECT localizacao FROM localizacoes)
            AND b.numero NOT IN (SELECT numero FROM atribuicoes)""",
                         ["localizacao", "ccustos"]),
        "processos": ("SELECT *, CAST(id AS TEXT) AS chave FROM processos_sei",
                      ["numero_sei", "descricao", "tipo", "vigente", "criado_em"]),
    }
    fonte, colunas = fontes[aba]
    busca_colunas = [c for c in colunas if c not in ("quantidade", "vigente", "criado_em")]
    where, params = [], []
    for palavra in filtros.get("q", "").split():
        where.append("(" + " OR ".join(f"instr(cadastro_busca({c}), ?) > 0" for c in busca_colunas) + ")")
        params.extend([normalizar(palavra)] * len(busca_colunas))
    if aba == "localizacoes":
        if filtros.get("centro"):
            where.append("ccustos = ?")
            params.append(filtros["centro"])
        if filtros.get("situacao") in ("sem_centro", "vinculadas"):
            where.append("ccustos = ''" if filtros["situacao"] == "sem_centro" else "ccustos <> ''")
    if aba == "processos":
        if filtros.get("tipo") in TIPOS_TERMO:
            where.append("tipo = ?")
            params.append(filtros["tipo"])
        if filtros.get("situacao") in ("vigente", "encerrado"):
            where.append("vigente = ?")
            params.append(int(filtros["situacao"] == "vigente"))
    base = f"FROM ({fonte}) AS cadastro" + (" WHERE " + " AND ".join(where) if where else "")
    total = conn.execute("SELECT COUNT(*) " + base, params).fetchone()[0]
    ordem = filtros.get("ordem") if filtros.get("ordem") in colunas else colunas[0]
    direcao = "DESC" if filtros.get("direcao") == "desc" else "ASC"
    try:
        tamanho = int(filtros.get("por_pagina", 20))
    except (TypeError, ValueError):
        tamanho = 20
    tamanho = tamanho if tamanho in (10, 20, 50) else 20
    paginas = max(1, (total + tamanho - 1) // tamanho)
    try:
        pagina = max(1, min(int(filtros.get("pagina", 1)), paginas))
    except (TypeError, ValueError):
        pagina = 1
    inicio = (pagina - 1) * tamanho
    itens = _todos(conn, f"SELECT * {base} ORDER BY {ordem} COLLATE NOCASE {direcao}, chave LIMIT ? OFFSET ?",
                   *params, tamanho, inicio)
    return dict(itens=itens, total=total, pagina=pagina, paginas=paginas, por_pagina=tamanho,
                inicio=inicio + 1 if total else 0, fim=min(inicio + tamanho, total),
                ordem=ordem, direcao=direcao.lower())


def salvar_centro(conn, antigo, dados):
    """Altera sigla, dados e referências do histórico em uma única transação."""
    sigla = _obrigatorio(dados.get("ccustos"), "Centro de custo").upper()
    nome = _obrigatorio(dados.get("responsavel"), "Responsável")
    with conn:
        if not responsavel(conn, antigo):
            raise ErroDeNegocio("Centro de custo não encontrado.")
        if sigla != antigo and responsavel(conn, sigla):
            raise ErroDeNegocio(f"O centro de custo {sigla} já existe.")
        conn.execute("""UPDATE responsaveis SET ccustos=?, responsavel=?, funcao=?, matricula=?, email=?
                        WHERE ccustos=?""", (sigla, nome, _texto(dados.get("funcao")),
                        _texto(dados.get("matricula")), _texto(dados.get("email")), antigo))
        conn.execute("UPDATE termos_emitidos SET chave=? WHERE tipo='ccusto' AND chave=?", (sigla, antigo))
    return sigla


def responsavel(conn, ccustos: str) -> dict | None:
    return _um(conn, "SELECT * FROM responsaveis WHERE ccustos = ?", ccustos)


def bens_do_centro(conn, ccustos: str) -> list[dict]:
    """Bens ATIVOS nas localizações do centro, excluindo os atribuídos a pessoas (setor OU pessoa)."""
    return _todos(conn, """
        SELECT b.* FROM bens b JOIN localizacoes l ON l.localizacao = b.localizacao
        WHERE l.ccustos = ? AND b.situacao = 'ATIVO'
          AND b.numero NOT IN (SELECT numero FROM atribuicoes)
        ORDER BY b.numero""", ccustos)


def pessoas(conn) -> list[str]:
    return [r[0] for r in conn.execute("SELECT nome FROM pessoas ORDER BY nome")]


def bens_da_pessoa(conn, nome: str) -> list[dict]:
    return _todos(conn, """
        SELECT b.* FROM atribuicoes a JOIN bens b ON b.numero = a.numero
        WHERE a.nome = ? ORDER BY b.numero""", nome)


def buscar_bem(conn, numero: int) -> dict | None:
    return _um(conn, "SELECT * FROM bens WHERE numero = ?", numero)


def pessoa_do_bem(conn, numero: int) -> str | None:
    r = conn.execute("SELECT nome FROM atribuicoes WHERE numero = ?", (numero,)).fetchone()
    return r[0] if r else None


def ficha_do_bem(conn, numero: int) -> dict | None:
    """Bem + centro de custo/responsável do setor (pela localização) + pessoa (pela atribuição)."""
    bem = buscar_bem(conn, numero)
    if not bem:
        return None
    setor = _um(conn, """
        SELECT r.ccustos, r.responsavel FROM localizacoes l JOIN responsaveis r ON r.ccustos = l.ccustos
        WHERE l.localizacao = ?""", bem["localizacao"]) or {"ccustos": None, "responsavel": None}
    return {**bem, **setor, "pessoa": pessoa_do_bem(conn, numero)}


def _padrao_like(termo: str) -> str:
    """Texto simples busca "contém". Com * ou %, o padrão é literal (GEX* = começa com GEX)."""
    termo = termo.replace("*", "%")
    return termo if "%" in termo else f"%{termo}%"


def _termos(q: str) -> list[str]:
    """Palavras da pesquisa; o "e" solto é conector ("computador e GEX-LIC"), não termo."""
    return [t for t in q.split() if t.lower() != "e"]


def _clausula(campos: list[str], termos: list[str]) -> tuple[str, list]:
    """Cada termo tem que bater em algum dos campos: (c1 LIKE ? OR c2 LIKE ?) AND (...)."""
    grupo = "(" + " OR ".join(f"MAIUSC({c}) LIKE ?" for c in campos) + ")"
    sql = " AND ".join([grupo] * len(termos)) or "1"
    return sql, [_padrao_like(t).upper() for t in termos for _ in campos]


def pesquisar(conn, q: str, limite: int = 200) -> dict:
    """Busca rápida em centros de custo (sigla/responsável), pessoas (nome) e bens (descrição, complemento,
    localização, centro, pessoa). Sem distinguir maiúsculas, inclusive acentuadas (MAIUSC = str.upper)."""
    conn.create_function("MAIUSC", 1, lambda v: v.upper() if isinstance(v, str) else v)
    termos = _termos(q)
    sql, params = _clausula(["ccustos", "responsavel"], termos)
    centros = _todos(conn, f"SELECT ccustos, responsavel FROM responsaveis WHERE {sql} ORDER BY ccustos", *params)
    for c in centros:
        c["quantidade"] = len(bens_do_centro(conn, c["ccustos"]))
    sql, params = _clausula(["p.nome"], termos)
    pessoas = _todos(conn, f"""
        SELECT p.nome, count(a.numero) AS quantidade FROM pessoas p LEFT JOIN atribuicoes a ON a.nome = p.nome
        WHERE {sql} GROUP BY p.nome ORDER BY p.nome""", *params)
    sql, params = _clausula(["b.descricao", "b.complemento", "b.localizacao", "l.ccustos", "a.nome"], termos)
    bens = _todos(conn, f"""
        SELECT b.*, l.ccustos AS ccustos, a.nome AS pessoa FROM bens b
        LEFT JOIN localizacoes l ON l.localizacao = b.localizacao LEFT JOIN atribuicoes a ON a.numero = b.numero
        WHERE {sql} ORDER BY b.numero LIMIT ?""", *params, limite + 1)
    return {"centros": centros, "pessoas": pessoas, "bens": bens[:limite], "truncado": len(bens) > limite}


def localizacoes_mapeadas(conn) -> list[dict]:
    return _todos(conn, "SELECT localizacao, ccustos FROM localizacoes ORDER BY localizacao")


class CentroEmUso(ErroDeNegocio):
    pass


class BemNaoEncontrado(ErroDeNegocio):
    pass


class JaAtribuido(ErroDeNegocio):
    def __init__(self, numero, pessoa):
        super().__init__(f"O bem {numero} está com {pessoa}.")
        self.pessoa = pessoa


def _obrigatorio(valor, rotulo) -> str:
    v = " ".join(str(valor or "").split())
    if not v:
        raise ErroDeNegocio(f"{rotulo} é obrigatório.")
    return v


def incluir_responsavel(conn, dados: dict) -> None:
    sigla = _obrigatorio(dados.get("ccustos"), "Centro de custo").upper()
    nome = _obrigatorio(dados.get("responsavel"), "Responsável")
    if responsavel(conn, sigla):
        raise ErroDeNegocio(f"O centro de custo {sigla} já existe.")
    conn.execute("INSERT INTO responsaveis (ccustos, responsavel, email, matricula, funcao) VALUES (?,?,?,?,?)", (
        sigla, nome, _texto(dados.get("email")), _texto(dados.get("matricula")), _texto(dados.get("funcao"))))
    conn.commit()


def checar_exclusao_centro(conn, ccustos: str) -> None:
    """Centro só pode ser excluído sem bens ativos sob guarda (bens em localizações dele, não atribuídos)."""
    if not responsavel(conn, ccustos):
        raise ErroDeNegocio(f"Centro de custo {ccustos} não encontrado.")
    n = len(bens_do_centro(conn, ccustos))
    if n:
        bem_str = "bem" if n == 1 else "bens"
        ativo_str = "ativo" if n == 1 else "ativos"
        raise CentroEmUso(f"{ccustos} tem {n} {bem_str} {ativo_str} sob guarda. "
                          "Mova as localizações para outro centro antes de excluir.")


def excluir_responsavel(conn, ccustos: str) -> None:
    checar_exclusao_centro(conn, ccustos)
    conn.execute("DELETE FROM localizacoes WHERE ccustos = ?", (ccustos,))  # voltam a "pendentes"
    conn.execute("DELETE FROM responsaveis WHERE ccustos = ?", (ccustos,))
    conn.commit()


def atualizar_responsavel(conn, ccustos: str, dados: dict) -> None:
    if not responsavel(conn, ccustos):
        raise ErroDeNegocio(f"Centro de custo {ccustos} não encontrado.")
    nome = _obrigatorio(dados.get("responsavel"), "Responsável")
    conn.execute("UPDATE responsaveis SET responsavel=?, email=?, matricula=?, funcao=? WHERE ccustos=?", (
        nome, _texto(dados.get("email")), _texto(dados.get("matricula")), _texto(dados.get("funcao")), ccustos))
    conn.commit()


def renomear_centro(conn, antigo: str, novo: str) -> None:
    novo = _obrigatorio(novo, "Nova sigla").upper()
    if responsavel(conn, novo):
        raise ErroDeNegocio(f"Já existe um centro de custo {novo}.")
    if not responsavel(conn, antigo):
        raise ErroDeNegocio(f"Centro de custo {antigo} não encontrado.")
    conn.execute("UPDATE responsaveis SET ccustos = ? WHERE ccustos = ?", (novo, antigo))  # cascateia
    conn.execute("UPDATE termos_emitidos SET chave = ? WHERE tipo = 'ccusto' AND chave = ?", (novo, antigo))
    conn.commit()


def incluir_localizacao(conn, localizacao: str, ccustos: str) -> None:
    loc = _obrigatorio(localizacao, "Localização")
    if not responsavel(conn, ccustos):
        raise ErroDeNegocio(f"Centro de custo {ccustos} não cadastrado.")
    conn.execute("INSERT OR REPLACE INTO localizacoes VALUES (?, ?)", (loc, ccustos))
    conn.commit()


def excluir_localizacao(conn, localizacao: str) -> None:
    conn.execute("DELETE FROM localizacoes WHERE localizacao = ?", (localizacao,))
    conn.commit()


def mover_localizacoes(conn, localizacoes: list[str], ccustos: str) -> int:
    """De-Para: muda o centro de custo de uma ou várias localizações de uma vez."""
    if not localizacoes:
        raise ErroDeNegocio("Selecione ao menos uma localização.")
    if not responsavel(conn, ccustos):
        raise ErroDeNegocio(f"Centro de custo {ccustos} não cadastrado.")
    marcadores = ",".join("?" * len(localizacoes))
    cur = conn.execute(f"UPDATE localizacoes SET ccustos = ? WHERE localizacao IN ({marcadores})", (ccustos, *localizacoes))
    conn.commit()
    return cur.rowcount


def pessoa(conn, nome: str) -> dict | None:
    return _um(conn, "SELECT * FROM pessoas WHERE nome = ?", nome)


def incluir_pessoa(conn, nome: str, email: str | None = None, matricula: str | None = None) -> str:
    nome = _obrigatorio(nome, "Nome").upper()
    conn.execute("INSERT OR IGNORE INTO pessoas (nome, email, matricula) VALUES (?,?,?)",
                 (nome, _texto(email) or None, _texto(matricula) or None))
    conn.commit()
    return nome


def salvar_pessoa(conn, antigo: str, dados: dict) -> str:
    """Renomeia (mantendo atribuições e histórico) e atualiza e-mail e matrícula. Tudo ou nada."""
    try:
        novo = _renomear_pessoa(conn, antigo, dados.get("nome"))
        conn.execute("UPDATE pessoas SET email = ?, matricula = ? WHERE nome = ?",
                     (_texto(dados.get("email")) or None, _texto(dados.get("matricula")) or None, novo))
        conn.commit()
    except Exception:
        conn.rollback()
        raise
    return novo


def excluir_pessoa(conn, nome: str) -> None:
    conn.execute("DELETE FROM pessoas WHERE nome = ?", (nome,))  # atribuições vão junto (cascade)
    conn.commit()


def renomear_pessoa(conn, antigo: str, novo: str) -> str:
    novo = _renomear_pessoa(conn, antigo, novo)
    conn.commit()
    return novo


def _renomear_pessoa(conn, antigo: str, novo: str) -> str:
    """Sem commit: quem chama decide quando a transação termina."""
    novo = _obrigatorio(novo, "Nome").upper()
    if antigo not in pessoas(conn):
        raise ErroDeNegocio(f"Pessoa {antigo} não encontrada.")
    if novo == antigo:
        return novo
    if novo in pessoas(conn):
        raise ErroDeNegocio(f"Já existe uma pessoa chamada {novo}.")
    conn.execute("UPDATE pessoas SET nome = ? WHERE nome = ?", (novo, antigo))  # cascateia em atribuicoes
    conn.execute("UPDATE termos_emitidos SET chave = ? WHERE tipo IN ('individual','devolucao') AND chave = ?", (novo, antigo))
    return novo


def atribuir(conn, nome: str, numero: int, confirmar: bool = False) -> None:
    """Coloca o bem sob responsabilidade da pessoa. Se já está com outra, exige confirmar=True."""
    if not buscar_bem(conn, numero):
        raise BemNaoEncontrado(f"Bem {numero} não encontrado na base.")
    atual = pessoa_do_bem(conn, numero)
    if atual == nome:
        return
    if atual and not confirmar:
        raise JaAtribuido(numero, atual)
    conn.execute("DELETE FROM atribuicoes WHERE numero = ?", (numero,))
    conn.execute("INSERT INTO atribuicoes VALUES (?, ?)", (nome, numero))
    conn.commit()


def desatribuir(conn, nome: str, numero: int) -> None:
    conn.execute("DELETE FROM atribuicoes WHERE nome = ? AND numero = ?", (nome, numero))
    conn.commit()


# ---------------------------------------------------------------- processos SEI
def processos(conn) -> list[dict]:
    return _todos(conn, "SELECT * FROM processos_sei ORDER BY vigente DESC, criado_em DESC, id DESC")


def processo_vigente(conn, tipo: str) -> dict | None:
    return _um(conn, "SELECT * FROM processos_sei WHERE tipo = ? AND vigente = 1", tipo)


def incluir_processo(conn, tipo: str, descricao: str, numero_sei: str, vigente: bool = True) -> int:
    if tipo not in TIPOS_TERMO:
        raise ErroDeNegocio("Tipo de termo inválido.")
    descricao = _obrigatorio(descricao, "Descrição")
    numero_sei = _obrigatorio(numero_sei, "Número SEI")
    if vigente:
        conn.execute("UPDATE processos_sei SET vigente = 0 WHERE tipo = ?", (tipo,))
    cur = conn.execute("INSERT INTO processos_sei (tipo, descricao, numero_sei, vigente, criado_em) VALUES (?,?,?,?,?)",
                       (tipo, descricao, numero_sei, 1 if vigente else 0, _agora()))
    conn.commit()
    return cur.lastrowid


def _processo(conn, id: int) -> dict:
    return _um(conn, "SELECT * FROM processos_sei WHERE id = ?", id) or _erro("Processo não encontrado.")


def _erro(msg):
    raise ErroDeNegocio(msg)


def marcar_vigente(conn, id: int) -> None:
    p = _processo(conn, id)
    conn.execute("UPDATE processos_sei SET vigente = 0 WHERE tipo = ?", (p["tipo"],))
    conn.execute("UPDATE processos_sei SET vigente = 1 WHERE id = ?", (id,))
    conn.commit()


def encerrar_processo(conn, id: int) -> None:
    _processo(conn, id)
    conn.execute("UPDATE processos_sei SET vigente = 0 WHERE id = ?", (id,))
    conn.commit()


def excluir_processo(conn, id: int) -> None:
    _processo(conn, id)
    if conn.execute("SELECT 1 FROM termos_emitidos WHERE processo_id = ? LIMIT 1", (id,)).fetchone():
        raise ErroDeNegocio("Este processo tem termos registrados; encerre-o em vez de excluir.")
    conn.execute("DELETE FROM processos_sei WHERE id = ?", (id,))
    conn.commit()


# ---------------------------------------------------------------- termos emitidos (foto)
def _numeros_do_termo(conn, termo_id: int) -> list[int]:
    return [r[0] for r in conn.execute("SELECT numero FROM termos_emitidos_bens WHERE termo_id = ? ORDER BY numero", (termo_id,))]


def ultimo_termo(conn, tipo: str, chave: str) -> dict | None:
    return _um(conn, """
        SELECT t.*, p.descricao AS processo, p.numero_sei FROM termos_emitidos t JOIN processos_sei p ON p.id = t.processo_id
        WHERE t.tipo = ? AND t.chave = ? ORDER BY t.emitido_em DESC, t.id DESC LIMIT 1""", tipo, chave)


def registrar_emissao(conn, tipo: str, chave: str, bens: list) -> dict:
    """Foto do termo. Sem processo vigente do tipo → ErroDeNegocio. No mesmo dia, com o mesmo processo e a mesma lista de bens, só atualiza a hora do registro existente."""
    proc = processo_vigente(conn, tipo)
    if not proc:
        raise ErroDeNegocio(f"Cadastre um processo SEI vigente para {ROTULO_TIPO[tipo]} em Cadastros → Processos SEI.")
    agora = _agora()
    numeros = sorted(int(b["numero"]) for b in bens)
    ultimo = ultimo_termo(conn, tipo, chave)
    if (ultimo and ultimo["emitido_em"][:10] == agora[:10] and ultimo["processo_id"] == proc["id"]
            and _numeros_do_termo(conn, ultimo["id"]) == numeros):
        conn.execute("UPDATE termos_emitidos SET emitido_em = ? WHERE id = ?", (agora, ultimo["id"]))
        conn.commit()
        return termo_emitido(conn, ultimo["id"])
    cur = conn.execute(
        "INSERT INTO termos_emitidos (tipo, chave, processo_id, emitido_em, quantidade, valor_total) VALUES (?,?,?,?,?,?)",
        (tipo, chave, proc["id"], agora, len(bens), sum(b["valor_atual"] or 0 for b in bens)))
    conn.executemany("INSERT INTO termos_emitidos_bens VALUES (?,?,?,?,?,?)", [
        (cur.lastrowid, b["numero"], b["descricao"], b["complemento"], b["localizacao"], b["valor_atual"]) for b in bens])
    conn.commit()
    return termo_emitido(conn, cur.lastrowid)


def termos_emitidos(conn, tipo: str | None = None, chave: str | None = None, limite: int = 200) -> list[dict]:
    sql = """SELECT t.*, p.descricao AS processo, p.numero_sei FROM termos_emitidos t
             JOIN processos_sei p ON p.id = t.processo_id WHERE 1"""
    params: list = []
    if tipo:
        sql += " AND t.tipo = ?"
        params.append(tipo)
    if chave:
        sql += " AND MAIUSC(t.chave) LIKE ?"
        params.append(f"%{chave.upper()}%")
    conn.create_function("MAIUSC", 1, lambda v: v.upper() if isinstance(v, str) else v)
    return _todos(conn, sql + " ORDER BY t.emitido_em DESC, t.id DESC LIMIT ?", *params, limite)


def termo_emitido(conn, id: int) -> dict | None:
    t = _um(conn, """
        SELECT t.*, p.descricao AS processo, p.numero_sei FROM termos_emitidos t
        JOIN processos_sei p ON p.id = t.processo_id WHERE t.id = ?""", id)
    if not t:
        return None
    t["bens"] = _todos(conn, "SELECT * FROM termos_emitidos_bens WHERE termo_id = ? ORDER BY numero", id)
    return t


def salvar_documento_sei(conn, id: int, documento: str, bloco: str = "") -> None:
    """Número do documento e do bloco de assinatura no SEI; os dois são necessários para pedir a assinatura."""
    conn.execute("UPDATE termos_emitidos SET documento_sei = ?, bloco_sei = ? WHERE id = ?",
                 (_texto(documento) or None, _texto(bloco) or None, id))
    conn.commit()


def registrar_email(conn, id: int) -> str:
    """Marca a hora em que o e-mail de assinatura foi enviado (controle; o envio é pelo programa de e-mail)."""
    agora = _agora()
    conn.execute("UPDATE termos_emitidos SET email_enviado_em = ? WHERE id = ?", (agora, id))
    conn.commit()
    return agora


def situacao_termo(conn, tipo: str, chave: str, bens_atuais: list) -> dict:
    """Compara a foto do último termo com os bens de hoje. Só entrada/saída conta."""
    ultimo = ultimo_termo(conn, tipo, chave)
    if not ultimo:
        return {"estado": "sem_termo", "ultimo": None, "entraram": 0, "sairam": 0}
    foto = set(_numeros_do_termo(conn, ultimo["id"]))
    atuais = {int(b["numero"]) for b in bens_atuais}
    entraram, sairam = len(atuais - foto), len(foto - atuais)
    return {"estado": "desatualizado" if entraram or sairam else "vigente",
            "ultimo": ultimo, "entraram": entraram, "sairam": sairam}


def situacoes_centros(conn) -> list[dict]:
    """Cada centro com quantidade/valor dos bens sob guarda e a situação do termo."""
    out = []
    for c in centros(conn):
        bens = bens_do_centro(conn, c["ccustos"])
        out.append({**c, "quantidade": len(bens), "valor": sum(b["valor_atual"] or 0 for b in bens),
                    **situacao_termo(conn, "ccusto", c["ccustos"], bens)})
    return out


def situacoes_devolucoes(conn) -> list[dict]:
    """Cada pessoa com os bens que ainda estão com ela e o último termo de devolução emitido."""
    out = []
    for nome in pessoas(conn):
        bens = bens_da_pessoa(conn, nome)
        out.append({"nome": nome, "quantidade": len(bens), "valor": sum(b["valor_atual"] or 0 for b in bens),
                    "ultimo": ultimo_termo(conn, "devolucao", nome)})
    return out


def situacoes_pessoas(conn) -> list[dict]:
    out = []
    for nome in pessoas(conn):
        bens = bens_da_pessoa(conn, nome)
        out.append({"nome": nome, "quantidade": len(bens), "valor": sum(b["valor_atual"] or 0 for b in bens),
                    **situacao_termo(conn, "individual", nome, bens)})
    return out


CADASTROS = {
    "responsaveis": ["ccustos", "responsavel", "email", "matricula", "funcao"],
    "localizacoes": ["localizacao", "ccustos"],
    "pessoas": ["nome", "email", "matricula"],
    "atribuicoes": ["nome", "numero"],
}


def acrescentar_linha(ws, valores: list) -> None:
    """ws.append que grava texto como texto: '=1+1' numa descrição vira fórmula no openpyxl e o Excel
    tenta calcular. Números e datas passam como estão."""
    ws.append(valores)
    for celula in ws[ws.max_row]:
        if isinstance(celula.value, str) and celula.data_type == "f":
            celula.data_type = "s"


def exportar_cadastros(conn, destino: Path) -> Path:
    """Planilha com as 4 tabelas de cadastro, no formato do banco (para backup e edição em massa),
    mais as abas inv_* (migração/backup dos eventos de inventário)."""
    wb = Workbook()
    wb.remove(wb.active)
    for tabela, colunas in CADASTROS.items():
        ws = wb.create_sheet(tabela)
        ws.append(colunas)
        for linha in conn.execute(f"SELECT {', '.join(colunas)} FROM {tabela} ORDER BY {colunas[0]}"):
            acrescentar_linha(ws, list(linha))
    import inventario
    inventario.exportar_abas(conn, wb)
    wb.save(destino)
    return destino


def _ler_aba_cadastro(wb, tabela: str, problemas: list, colunas=None, opcional=False, opcionais=()) -> list[dict] | None:
    colunas = colunas or CADASTROS[tabela]
    if tabela not in wb.sheetnames:
        if opcional:
            return None
        problemas.append(f"aba '{tabela}' não encontrada")
        return []
    ws = wb[tabela]
    it = ws.iter_rows(values_only=True)
    cabecalho = [_texto(c) for c in next(it, ())]
    faltando = [c for c in colunas if c not in cabecalho and c not in opcionais]
    if faltando:
        problemas.append(f"aba '{tabela}': coluna(s) ausente(s): {', '.join(faltando)}")
        return []
    idx = {c: (cabecalho.index(c) if c in cabecalho else None) for c in colunas}
    linhas = []
    for n, r in enumerate(it, start=2):
        if r is None or all(v is None or _texto(v) == "" for v in r):
            continue
        linhas.append({"_linha": n, **{c: (r[i] if i is not None and i < len(r) else None) for c, i in idx.items()}})
    return linhas


def importar_cadastros(conn, arquivo) -> dict:
    """Substitui responsaveis, localizacoes, pessoas e atribuicoes pelo conteúdo da planilha. Tudo ou nada."""
    try:
        wb = load_workbook(arquivo, read_only=True, data_only=True)
    except Exception:
        raise ImportacaoInvalida("Arquivo inválido: envie a planilha de cadastros em .xlsx.")
    import inventario
    problemas: list[str] = []
    try:
        brutos = {t: _ler_aba_cadastro(wb, t, problemas) for t in CADASTROS}
        inv_brutos = {aba: _ler_aba_cadastro(wb, aba, problemas, colunas=cols + (["foto_url"] if aba == "inv_leituras" else []),
                                             opcional=True, opcionais=("foto_url", "fotos_seq") if aba == "inv_leituras" else ())
                     for aba, cols in inventario.ABAS.items()}
        tem_inventario = any(v is not None for v in inv_brutos.values())
    except ImportacaoInvalida:
        raise
    except Exception:
        raise ImportacaoInvalida("Não consegui ler a planilha; o arquivo pode estar corrompido.")
    finally:
        wb.close()
    if problemas:
        raise ImportacaoInvalida("Planilha de cadastros: " + "; ".join(problemas))
    faltam = [aba for aba, v in inv_brutos.items() if v is None and aba not in inventario.ABAS_OPCIONAIS]
    if tem_inventario and faltam:
        raise ImportacaoInvalida("Planilha de cadastros: abas de inventário incompletas (faltam: " + ", ".join(faltam)
                                 + "). Envie as 5 abas inv_* ou nenhuma.")

    responsaveis, siglas = [], set()
    for r in brutos["responsaveis"]:
        sigla, nome = _texto(r["ccustos"]).upper(), _texto(r["responsavel"])
        if not sigla:
            problemas.append(f"responsaveis linha {r['_linha']}: sigla vazia")
        elif sigla in siglas:
            problemas.append(f"responsaveis linha {r['_linha']}: sigla {sigla} repetida")
        elif not nome:
            problemas.append(f"responsaveis linha {r['_linha']}: responsável vazio")
        else:
            siglas.add(sigla)
            responsaveis.append((sigla, nome, _texto(r["email"]), _texto(r["matricula"]), _texto(r["funcao"])))

    localizacoes, locs = [], set()
    for r in brutos["localizacoes"]:
        loc, sigla = _texto(r["localizacao"]), _texto(r["ccustos"]).upper()
        if not loc:
            problemas.append(f"localizacoes linha {r['_linha']}: localização vazia")
        elif loc in locs:
            problemas.append(f"localizacoes linha {r['_linha']}: localização {loc} repetida")
        elif sigla not in siglas:
            problemas.append(f"localizacoes linha {r['_linha']}: centro {sigla or '(vazio)'} não está na aba responsaveis")
        else:
            locs.add(loc)
            localizacoes.append((loc, sigla))

    pessoas_linhas, nomes = [], set()
    for r in brutos["pessoas"]:
        nome = _texto(r["nome"]).upper()
        if not nome:
            problemas.append(f"pessoas linha {r['_linha']}: nome vazio")
        elif nome in nomes:
            problemas.append(f"pessoas linha {r['_linha']}: nome {nome} repetido")
        else:
            nomes.add(nome)
            pessoas_linhas.append((nome, _texto(r["email"]), _texto(r["matricula"])))

    atribuicoes, numeros = [], set()
    for r in brutos["atribuicoes"]:
        nome, num = _texto(r["nome"]).upper(), _numero(r["numero"])
        if nome not in nomes:
            problemas.append(f"atribuicoes linha {r['_linha']}: {nome or '(vazio)'} não está na aba pessoas")
        elif num is None or num != int(num):
            problemas.append(f"atribuicoes linha {r['_linha']}: número inválido")
        elif not buscar_bem(conn, int(num)):
            problemas.append(f"atribuicoes linha {r['_linha']}: bem {int(num)} não existe na base")
        elif int(num) in numeros:
            problemas.append(f"atribuicoes linha {r['_linha']}: bem {int(num)} repetido (um bem, uma pessoa)")
        else:
            numeros.add(int(num))
            atribuicoes.append((nome, int(num)))

    inv_linhas = {}
    if tem_inventario:
        inv_linhas, inv_problemas = inventario.validar_abas(conn, {a: (v or []) for a, v in inv_brutos.items()})
        problemas += inv_problemas
        if inv_brutos["inv_bens_encerrados"] is None:
            inv_linhas["inv_bens_encerrados"] = None          # aba ausente: tabela mantida

    if problemas:
        extra = f" (+{len(problemas) - 20})" if len(problemas) > 20 else ""
        raise ImportacaoInvalida("Planilha de cadastros: " + "; ".join(problemas[:20]) + extra)

    try:
        for t in ("atribuicoes", "pessoas", "localizacoes", "responsaveis"):
            conn.execute(f"DELETE FROM {t}")
        conn.executemany("INSERT INTO responsaveis (ccustos, responsavel, email, matricula, funcao) VALUES (?,?,?,?,?)", responsaveis)
        conn.executemany("INSERT INTO localizacoes VALUES (?,?)", localizacoes)
        conn.executemany("INSERT INTO pessoas (nome, email, matricula) VALUES (?,?,?)", sorted(pessoas_linhas))
        conn.executemany("INSERT INTO atribuicoes VALUES (?,?)", atribuicoes)
        if tem_inventario:
            inventario.substituir_tabelas(conn, inv_linhas)
        conn.commit()
    except Exception:
        conn.rollback()
        raise
    resultado = {"responsaveis": len(responsaveis), "localizacoes": len(localizacoes), "pessoas": len(nomes),
                "atribuicoes": len(atribuicoes), "sem_centro": localizacoes_sem_centro(conn)}
    if tem_inventario:
        resultado.update({aba: len(l) for aba, l in inv_linhas.items() if l is not None})
    return resultado


# ---------------------------------------------------------------- painel e recorte
IMOVEIS = ("SEDE", "TERRENOS")
FAIXAS_IDADE = [("ate5", "até 5 anos"), ("5a10", "5 a 10 anos"), ("10a20", "10 a 20 anos"),
                ("mais20", "mais de 20 anos"), ("semdata", "sem data")]
FAIXAS_VALOR = [("ate100", "até R$ 100"), ("100a500", "R$ 100 a 500"), ("500a1000", "R$ 500 a 1.000"),
                ("1000a5000", "R$ 1.000 a 5.000"), ("5000a20000", "R$ 5.000 a 20.000"), ("mais20000", "acima de R$ 20.000")]
FILTROS = ("situacao", "ccusto", "pessoa", "localizacao", "classificacao", "valor_de", "valor_ate",
           "entrada_de", "entrada_ate", "idade", "faixa", "ano", "valor_status")
_DE = ("FROM bens b LEFT JOIN localizacoes l ON l.localizacao = b.localizacao "
       "LEFT JOIN atribuicoes a ON a.numero = b.numero")
_DATA_ISO = ("CASE WHEN b.data_entrada LIKE '__/__/____' THEN substr(b.data_entrada,7,4)||'-'||"
             "substr(b.data_entrada,4,2)||'-'||substr(b.data_entrada,1,2) END")
_IDADE = f"(julianday('now') - julianday({_DATA_ISO})) / 365.25"
_FAIXA_IDADE = (f"CASE WHEN {_DATA_ISO} IS NULL THEN 'semdata' WHEN {_IDADE} < 5 THEN 'ate5' "
                f"WHEN {_IDADE} < 10 THEN '5a10' WHEN {_IDADE} < 20 THEN '10a20' ELSE 'mais20' END")
_V = "coalesce(b.valor_atual, 0)"
_FAIXA_VALOR = (f"CASE WHEN {_V} < 100 THEN 'ate100' WHEN {_V} < 500 THEN '100a500' WHEN {_V} < 1000 THEN '500a1000' "
                f"WHEN {_V} < 5000 THEN '1000a5000' WHEN {_V} < 20000 THEN '5000a20000' ELSE 'mais20000' END")
_IMOVEIS_SQL = "coalesce(b.classificacao, '') IN " + str(IMOVEIS)


def _where(f: dict) -> tuple[str, list]:
    """WHERE só com os filtros presentes em f (chaves de FILTROS; vazio = sem filtro). "-" é a sentinela do
    balde vazio (sem centro, sem pessoa, sem classificação, sem localização, sem data)."""
    cl, p = [], []
    if f.get("situacao"):
        cl.append("b.situacao = ?"); p.append(f["situacao"])
    if f.get("ccusto") == "-":
        cl.append("l.ccustos IS NULL")
    elif f.get("ccusto"):
        cl.append("l.ccustos = ?"); p.append(f["ccusto"])
    if f.get("pessoa") == "-":
        cl.append("a.nome IS NULL")
    elif f.get("pessoa"):
        cl.append("a.nome = ?"); p.append(f["pessoa"])
    if f.get("localizacao") == "-":
        cl.append("coalesce(b.localizacao, '') = ''")
    elif f.get("localizacao"):
        cl.append("b.localizacao = ?"); p.append(f["localizacao"])
    c = f.get("classificacao")
    if c == "-":
        cl.append("coalesce(b.classificacao, '') = ''")
    elif c == "imoveis":
        cl.append(_IMOVEIS_SQL)
    elif c == "sem-imoveis":
        cl.append(f"NOT {_IMOVEIS_SQL}")
    elif c:
        cl.append("b.classificacao = ?"); p.append(c)
    status_valor = f.get("valor_status")
    if status_valor == "nao_informado":
        cl.append("b.valor_atual IS NULL")
    elif status_valor == "zero":
        cl.append("b.valor_atual = 0")
    elif status_valor:
        raise ErroDeNegocio("Situação do valor inválida.")
    if f.get("valor_de"):
        cl.append("b.valor_atual >= ?"); p.append(float(f["valor_de"]))
    if f.get("valor_ate"):
        cl.append("b.valor_atual <= ?"); p.append(float(f["valor_ate"]))
    if f.get("entrada_de"):
        cl.append(f"{_DATA_ISO} >= ?"); p.append(f["entrada_de"])
    if f.get("entrada_ate"):
        cl.append(f"{_DATA_ISO} <= ?"); p.append(f["entrada_ate"])
    if f.get("idade"):
        cl.append(f"{_FAIXA_IDADE} = ?"); p.append(f["idade"])
    if f.get("faixa"):
        cl.append(f"{_FAIXA_VALOR} = ?"); p.append(f["faixa"])
    if f.get("ano") == "-":
        cl.append("coalesce(substr(b.data_entrada, 7, 4), '') = ''")
    elif f.get("ano"):
        cl.append("substr(b.data_entrada, 7, 4) = ?"); p.append(f["ano"])
    return (" AND ".join(cl) or "1"), p


def _agrupar(conn, chave: str, f: dict, ordem: str = "quantidade DESC, chave") -> list[dict]:
    where, p = _where(f)
    return _todos(conn, f"SELECT {chave} AS chave, count(*) AS quantidade, coalesce(sum(b.valor_atual), 0) AS valor "
                        f"{_DE} WHERE {where} GROUP BY chave ORDER BY {ordem}", *p)


def dimensoes(conn, f: dict) -> dict:
    """Agrupamentos do conjunto filtrado por f. Item: {chave, rotulo, quantidade, valor}.
    'situacao' ignora o filtro de situação (mostra a composição inteira)."""
    resp = {c["ccustos"]: c["responsavel"] for c in centros(conn)}
    mapa = {l["localizacao"]: l["ccustos"] for l in localizacoes_mapeadas(conn)}

    def rot(lista, fn):
        return [{**x, "rotulo": fn(x["chave"])} for x in lista]

    def fixas(faixas, lista):
        por = {x["chave"]: x for x in lista}
        return [{**por.get(k, {"chave": k, "quantidade": 0, "valor": 0}), "rotulo": r} for k, r in faixas]

    return {
        "situacao": rot(_agrupar(conn, "b.situacao", {k: v for k, v in f.items() if k != "situacao"}), str),
        "centro": rot(_agrupar(conn, "coalesce(l.ccustos, '-')", f),
                      lambda k: "sem centro" if k == "-" else f"{k} – {resp.get(k, '')}"),
        "classificacao": rot(_agrupar(conn, "CASE WHEN coalesce(b.classificacao, '') = '' THEN '-' ELSE b.classificacao END", f),
                            lambda k: "sem classificação" if k == "-" else k),
        "localizacao": rot(_agrupar(conn, "CASE WHEN coalesce(b.localizacao, '') = '' THEN '-' ELSE b.localizacao END", f),
                           lambda k: ("sem localização" if k == "-" else k) + (f" ({mapa[k]})" if k in mapa else "")),
        "idade": fixas(FAIXAS_IDADE, _agrupar(conn, _FAIXA_IDADE, f)),
        "ano": rot(_agrupar(conn, "CASE WHEN coalesce(substr(b.data_entrada, 7, 4), '') = '' THEN '-' ELSE substr(b.data_entrada, 7, 4) END",
                           f, ordem="chave"), lambda k: "sem data" if k == "-" else k),
        "faixa": fixas(FAIXAS_VALOR, _agrupar(conn, _FAIXA_VALOR, f)),
        "pessoa": rot([x for x in _agrupar(conn, "a.nome", f) if x["chave"]], str),
    }


def painel(conn) -> dict:
    def um(sql, *p):
        return conn.execute(sql, p).fetchone()[0]
    ativo = "b.situacao = 'ATIVO'"
    d = {
        "ativos": um(f"SELECT count(*) {_DE} WHERE {ativo}"),
        "valor_sem_imoveis": um(f"SELECT coalesce(sum(b.valor_atual), 0) {_DE} WHERE {ativo} AND NOT {_IMOVEIS_SQL}"),
        "imoveis": um(f"SELECT count(*) {_DE} WHERE {ativo} AND {_IMOVEIS_SQL}"),
        "valor_imoveis": um(f"SELECT coalesce(sum(b.valor_atual), 0) {_DE} WHERE {ativo} AND {_IMOVEIS_SQL}"),
        "sem_centro": um(f"SELECT count(*) {_DE} WHERE {ativo} AND l.ccustos IS NULL AND a.nome IS NULL"),
        "ultima_importacao": (importacoes(conn, 1) or [None])[0],
        "robo": (execucoes_robo(conn, 1) or [None])[0],
        "importacao_desatualizada": importacao_desatualizada(conn),
        "centros": situacoes_centros(conn),
        "pessoas": situacoes_pessoas(conn),
    }
    d.update({
        "valor_nao_informado": um(f"SELECT count(*) FROM bens b WHERE {ativo} AND b.valor_atual IS NULL"),
        "valor_zero": um(f"SELECT count(*) FROM bens b WHERE {ativo} AND b.valor_atual = 0"),
    })
    d["a_emitir_centros"] = sum(1 for c in d["centros"] if c["estado"] != "vigente" and c["quantidade"])
    d["a_emitir_pessoas"] = sum(1 for p in d["pessoas"] if p["estado"] != "vigente" and p["quantidade"])
    return d


def recorte(conn, f: dict, limite: int | None = 1000) -> dict:
    where, p = _where(f)
    sql = f"SELECT b.*, l.ccustos AS ccustos, a.nome AS pessoa {_DE} WHERE {where} ORDER BY b.numero"
    bens = _todos(conn, sql + (f" LIMIT {limite + 1}" if limite else ""), *p)
    totais = dict(conn.execute(f"""SELECT
      count(*) AS quantidade,
      coalesce(sum(b.valor_atual),0) AS valor_total,
      coalesce(sum(CASE WHEN {_IMOVEIS_SQL} THEN 1 ELSE 0 END),0) AS imoveis,
      coalesce(sum(CASE WHEN {_IMOVEIS_SQL} THEN b.valor_atual ELSE 0 END),0) AS valor_imoveis,
      coalesce(sum(CASE WHEN l.ccustos IS NULL AND a.nome IS NULL THEN 1 ELSE 0 END),0) AS sem_centro,
      coalesce(sum(CASE WHEN b.valor_atual IS NULL THEN 1 ELSE 0 END),0) AS valor_nao_informado,
      coalesce(sum(CASE WHEN b.valor_atual = 0 THEN 1 ELSE 0 END),0) AS valor_zero
      {_DE} WHERE {where}""", p).fetchone())
    return dict(totais, bens=bens[:limite] if limite else bens,
                truncado=bool(limite) and len(bens) > limite, dimensoes=dimensoes(conn, f))


def exportar_recorte(conn, f: dict, destino):
    wb = Workbook()
    ws = wb.active
    ws.title = "analise"
    ws.append(["Número", "Descrição", "Complemento", "Classificação", "Localização", "Centro de custo", "Pessoa",
               "Situação", "Data entrada", "Valor compra", "Valor atual"])
    for b in recorte(conn, f, limite=None)["bens"]:
        acrescentar_linha(ws, [b["numero"], b["descricao"], b["complemento"], b["classificacao"], b["localizacao"], b["ccustos"],
                                b["pessoa"], b["situacao"], b["data_entrada"], b["valor_compra"], b["valor_atual"]])
    wb.save(destino)
    return destino


def exportar_bens(conn, destino: Path) -> Path:
    """Planilha dos bens no formato do export do SPW (mesmas 9 colunas) — backup reimportável."""
    wb = Workbook()
    ws = wb.active
    ws.title = "base"
    ws.append(list(COLUNAS_EXPORT))
    campos = ", ".join(COLUNAS_EXPORT.values())
    for linha in conn.execute(f"SELECT {campos} FROM bens ORDER BY numero"):
        acrescentar_linha(ws, list(linha))
    wb.save(destino)
    return destino
