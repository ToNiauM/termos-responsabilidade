from docx import Document
from openpyxl import load_workbook

from tests.conftest import semear

import textos


def bens():
    return [
        {"numero": 1001, "descricao": "CADEIRA", "complemento": "GIRATÓRIA", "localizacao": "01 - SALA", "valor_atual": 64.54},
        {"numero": 1002, "descricao": "NOTEBOOK", "complemento": None, "localizacao": "01 - SALA", "valor_atual": 1500.0},
    ]


def linhas(caminho):
    return Document(caminho).tables[0].rows


def test_termo_individual(dados, tmp_path):
    from Script_Termo_Individual import criar_termo_responsabilidade
    destino = criar_termo_responsabilidade("ANA SILVA", bens(), tmp_path / "t.docx")
    tab = linhas(destino)
    assert len(tab) == 4  # cabeçalho + 2 + total
    assert tab[3].cells[3].text == "R$ 1.564,54"
    assert tab[2].cells[2].text == ""  # complemento None vira vazio


def test_termo_centro_e_planilha(dados, tmp_path):
    from Termo_de_Responsabilidade import gerar_planilha_centro, gerar_termo_centro
    resp = {"ccustos": "CCI", "responsavel": "JAQUELINE", "matricula": "46", "funcao": "coordenadora"}
    destino = gerar_termo_centro("CCI", resp, bens(), tmp_path / "c.docx")
    tab = linhas(destino)
    assert len(tab) == 4 and tab[3].cells[4].text == "R$ 1.564,54"
    assert "JAQUELINE" in Document(destino).paragraphs[-1].text
    xlsx = gerar_planilha_centro(bens(), tmp_path / "c.xlsx")
    ws = load_workbook(xlsx).active
    assert ws.max_row == 3 and ws["A1"].value == "numero"


def test_termo_devolucao(dados, tmp_path):
    from termo_devolucao import gerar_termo_devolucao
    destino = gerar_termo_devolucao("ANA SILVA", bens(), tmp_path / "d.docx")
    assert len(linhas(destino)) == 4
    assert gerar_termo_devolucao("ANA SILVA", [], tmp_path / "vazio.docx") is None


def texto(caminho):
    return "\n".join(p.text for p in Document(caminho).paragraphs)


def test_individual_com_textos_alterados(dados, tmp_path):
    from Script_Termo_Individual import criar_termo_responsabilidade
    t = dict(textos.PADRAO, individual_abertura="TESTE {nome} do {orgao_sigla}.", individual_compromissos="a\nb\nc",
             orgao_sigla="XYZ", assinatura_eletronica="Assinado via X")
    destino = criar_termo_responsabilidade("ANA SILVA", bens(), tmp_path / "t.docx", textos=t)
    doc = Document(destino)
    abertura = next(p for p in doc.paragraphs if p.text.startswith("TESTE"))
    assert [r.text for r in abertura.runs] == ["TESTE ", "ANA SILVA", " do XYZ."] and abertura.runs[1].bold
    corpo = texto(destino)
    assert "\na\n" in corpo and "\nc\n" in corpo and corpo.rstrip().endswith("Assinado via X")


def test_centro_com_paragrafos_e_assinatura_dos_textos(dados, tmp_path):
    from Termo_de_Responsabilidade import gerar_termo_centro
    t = dict(textos.PADRAO, ccusto_paragrafos="Primeiro {ccustos}.\n\nSegundo do {orgao_sigla}.",
             ccusto_assinatura="{responsavel}\nChefe", orgao_sigla="XYZ")
    resp = {"ccustos": "CCI", "responsavel": "JAQUELINE", "matricula": "46", "funcao": "coordenadora"}
    destino = gerar_termo_centro("CCI", resp, bens(), tmp_path / "c.docx", textos=t)
    corpo = texto(destino)
    assert "Primeiro CCI.\nSegundo do XYZ." in corpo
    assert Document(destino).paragraphs[-1].text == "JAQUELINE\nChefe"


def test_devolucao_com_recebedor_e_cidade_dos_textos(dados, tmp_path):
    from termo_devolucao import gerar_termo_devolucao
    t = dict(textos.PADRAO, cidade="Goiânia (GO)", recebedor_nome="FULANO", recebedor_cargo="Chefe")
    destino = gerar_termo_devolucao("ANA SILVA", bens(), tmp_path / "d.docx", textos=t)
    corpo = texto(destino)
    assert "Goiânia (GO), " in corpo and "\nFULANO\nChefe\n" in corpo
