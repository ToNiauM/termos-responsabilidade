import sqlite3

import pytest
from openpyxl import Workbook, load_workbook

import db
from tests.conftest import semear


def test_esquema_cria_seis_tabelas(dados):
    nomes = {r["name"] for r in dados.execute("SELECT name FROM sqlite_master WHERE type='table'")}
    assert {"bens", "responsaveis", "localizacoes", "pessoas", "atribuicoes", "textos"} <= nomes


def test_esquema_e_idempotente(dados):
    db.criar_esquema(dados)  # segunda vez não pode falhar


def test_foreign_keys_ligadas(dados):
    semear(dados)
    with pytest.raises(sqlite3.IntegrityError):
        dados.execute("INSERT INTO localizacoes VALUES ('02 - X', 'NAO_EXISTE')")


CABECALHO = ["Número Bem", "Situação", "Descrição", "Complemento", "Classificação Contábil",
             "Localização", "Data Entrada", "Valor Compra", "Valor Atual"]


def xlsx(tmp_path, linhas, cabecalho=CABECALHO, aba="base"):
    wb = Workbook()
    ws = wb.active
    ws.title = aba
    ws.append(cabecalho)
    for l in linhas:
        ws.append(l)
    caminho = tmp_path / "export.xlsx"
    wb.save(caminho)
    return caminho


def test_importar_substitui_bens_e_conta(dados, tmp_path):
    semear(dados)
    arq = xlsx(tmp_path, [
        [1002, "ATIVO", "NOTEBOOK", "DELL NOVO", "EQUIP", "01 - SALA CCI", "06/12/2012", 3000, 1400],
        [2001, "ATIVO", "MONITOR", "LG", "EQUIP", "01 - SALA CCI", "01/01/2020", 900, 800],
        [2002, "DOADO", "CADEIRA", "", "MÓVEIS", "CFC", "01/01/2000", 10, 1],
    ])
    resumo = db.importar_bens(dados, arq)
    assert resumo["total"] == 3 and resumo["ativos"] == 2
    assert [r["numero"] for r in dados.execute("SELECT numero FROM bens ORDER BY numero")] == [1002, 2001, 2002]
    assert dados.execute("SELECT complemento FROM bens WHERE numero=1002").fetchone()[0] == "DELL NOVO"
    # outras tabelas intactas
    assert dados.execute("SELECT count(*) FROM atribuicoes").fetchone()[0] == 1


def test_importar_cabecalho_errado_nao_altera_nada(dados, tmp_path):
    semear(dados)
    arq = xlsx(tmp_path, [[1, "ATIVO"]], cabecalho=["Patrimônio", "Situação"])
    with pytest.raises(db.ImportacaoInvalida):
        db.importar_bens(dados, arq)
    assert dados.execute("SELECT count(*) FROM bens").fetchone()[0] == 4


def test_importar_que_some_com_bem_atribuido_e_revertida(dados, tmp_path):
    semear(dados)  # ANA tem o 1002
    arq = xlsx(tmp_path, [[1001, "ATIVO", "CADEIRA", "", "MÓVEIS", "01 - SALA CCI", "x", 1, 1]])
    with pytest.raises(db.ImportacaoInvalida) as e:
        db.importar_bens(dados, arq)
    assert "1002" in str(e.value)
    assert dados.execute("SELECT count(*) FROM bens").fetchone()[0] == 4


def test_importar_usa_aba_pelo_cabecalho_e_converte_data(dados, tmp_path):
    from datetime import datetime
    wb = Workbook()
    wb.active.title = "outra"
    wb.active.append(["Nada", "aqui"])
    ws = wb.create_sheet("base")
    ws.append(CABECALHO)
    ws.append([3001, "ATIVO", "MESA", None, "MÓVEIS", "02 - X", datetime(2020, 3, 9), 100, 90])
    arq = tmp_path / "e.xlsx"
    wb.save(arq)
    db.importar_bens(dados, arq)
    r = dados.execute("SELECT data_entrada, complemento FROM bens WHERE numero=3001").fetchone()
    assert r["data_entrada"] == "09/03/2020" and r["complemento"] == ""


def test_localizacoes_sem_centro(dados):
    semear(dados)
    assert db.localizacoes_sem_centro(dados) == ["99 - SEM MAPA"]


def test_localizacoes_sem_centro_ignora_bens_atribuidos(dados):
    semear(dados)
    dados.execute("INSERT INTO bens VALUES (2003,'ATIVO','TABLET','','EQUIPAMENTOS','TERMOS INDIVIDUAIS','01/02/2023',900,800)")
    assert db.localizacoes_sem_centro(dados) == ["99 - SEM MAPA", "TERMOS INDIVIDUAIS"]
    dados.execute("INSERT INTO atribuicoes VALUES ('ANA SILVA', 2003)")
    assert db.localizacoes_sem_centro(dados) == ["99 - SEM MAPA"]


def test_importar_arquivo_invalido_levanta_importacao_invalida(dados, tmp_path):
    semear(dados)
    arq = tmp_path / "x.xlsx"
    arq.write_bytes(b"nada")
    with pytest.raises(db.ImportacaoInvalida):
        db.importar_bens(dados, arq)
    assert dados.execute("SELECT count(*) FROM bens").fetchone()[0] == 4


def test_importar_numero_repetido_e_revertida(dados, tmp_path):
    semear(dados)
    arq = xlsx(tmp_path, [
        [1001, "ATIVO", "CADEIRA", "", "MÓVEIS", "01 - SALA CCI", "01/01/2000", 10, 5],
        [1001, "ATIVO", "CADEIRA CÓPIA", "", "MÓVEIS", "01 - SALA CCI", "01/01/2000", 10, 5],
    ])
    with pytest.raises(db.ImportacaoInvalida) as e:
        db.importar_bens(dados, arq)
    assert "repetido" in str(e.value)
    assert dados.execute("SELECT count(*) FROM bens").fetchone()[0] == 4


# ---------------------------------------------------------------- importações
def test_importar_registra_mudancas(dados, tmp_path):
    semear(dados)
    arq = xlsx(tmp_path, [
        [1001, "ATIVO", "CADEIRA", "GIRATÓRIA", "MÓVEIS", "02 - OUTRA SALA", "31/12/1996", 75.94, 64.54],   # movido
        [1002, "ATIVO", "NOTEBOOK", "DELL", "EQUIPAMENTOS", "01 - SALA CCI", "06/12/2012", 3000, 1500],       # igual
        [1003, "ATIVO", "MESA", "ANTIGA", "MÓVEIS", "03 - DEPÓSITO", "06/12/2012", 100, 10],                  # situação + movido
        [5000, "ATIVO", "LUMINÁRIA", "", "MÓVEIS", "01 - SALA CCI", "01/01/2020", 10, 9],                     # novo; 1004 some
    ])
    r = db.importar_bens(dados, arq, nome_arquivo="export.xlsx")
    assert (r["novos"], r["removidos"], r["movidos"], r["situacao"]) == (1, 1, 2, 1)
    imp = db.importacao(dados, r["importacao_id"])
    assert imp["arquivo"] == "export.xlsx" and imp["total"] == 4 and imp["novos"] == 1
    tipos = sorted((m["numero"], m["tipo"], m["de"], m["para"]) for m in imp["mudancas"])
    assert tipos == [(1001, "movido", "01 - SALA CCI", "02 - OUTRA SALA"), (1003, "movido", "01 - SALA CCI", "03 - DEPÓSITO"),
                     (1003, "situacao", "BAIXADO", "ATIVO"), (1004, "removido", "99 - SEM MAPA", None), (5000, "novo", None, "01 - SALA CCI")]
    assert [i["id"] for i in db.importacoes(dados)] == [r["importacao_id"]]
    h = db.historico_do_bem(dados, 1003)
    assert [m["tipo"] for m in h["mudancas"]] == ["movido", "situacao"] and h["mudancas"][0]["importado_em"] == imp["importado_em"]
    assert h["termos"] == []


def test_importar_com_falha_nao_registra_importacao(dados, tmp_path):
    semear(dados)
    arq = xlsx(tmp_path, [[1001, "ATIVO", "CADEIRA", "", "MÓVEIS", "01 - SALA CCI", "x", 1, 1]])   # some o 1002 (atribuído)
    with pytest.raises(db.ImportacaoInvalida):
        db.importar_bens(dados, arq)
    assert db.importacoes(dados) == []


def test_historico_do_bem_lista_termos(dados):
    semear(dados)
    db.incluir_processo(dados, "ccusto", "T", "1")
    db.registrar_emissao(dados, "ccusto", "CCI", db.bens_do_centro(dados, "CCI"))
    h = db.historico_do_bem(dados, 1001)
    assert len(h["termos"]) == 1 and h["termos"][0]["chave"] == "CCI" and h["termos"][0]["tipo"] == "ccusto"


def test_bens_do_centro_exclui_baixados_atribuidos_e_sem_mapa(dados):
    semear(dados)
    assert [b["numero"] for b in db.bens_do_centro(dados, "CCI")] == [1001]


def test_bens_da_pessoa(dados):
    semear(dados)
    bens = db.bens_da_pessoa(dados, "ANA SILVA")
    assert [b["numero"] for b in bens] == [1002] and bens[0]["descricao"] == "NOTEBOOK"


def test_centros_responsavel_pessoas(dados):
    semear(dados)
    assert [c["ccustos"] for c in db.centros(dados)] == ["CCI"]
    assert db.responsavel(dados, "CCI")["responsavel"] == "JAQUELINE PORTELA"
    assert db.responsavel(dados, "XX") is None
    assert db.pessoas(dados) == ["ANA SILVA"]


def test_ficha_do_bem_com_setor_e_pessoa(dados):
    semear(dados)
    f = db.ficha_do_bem(dados, 1002)
    assert f["ccustos"] == "CCI" and f["responsavel"] == "JAQUELINE PORTELA" and f["pessoa"] == "ANA SILVA"
    f = db.ficha_do_bem(dados, 1004)
    assert f["ccustos"] is None and f["pessoa"] is None
    assert db.ficha_do_bem(dados, 9999) is None
    assert db.buscar_bem(dados, 1001)["descricao"] == "CADEIRA"
    assert db.localizacoes_mapeadas(dados) == [{"localizacao": "01 - SALA CCI", "ccustos": "CCI"}]


def test_renomear_centro_cascateia(dados):
    semear(dados)
    db.renomear_centro(dados, "CCI", "GESERV")
    assert db.localizacoes_mapeadas(dados)[0]["ccustos"] == "GESERV"
    assert db.responsavel(dados, "CCI") is None and db.responsavel(dados, "GESERV")


def test_excluir_centro_em_uso_falha(dados):
    semear(dados)  # CCI tem o bem 1001 ativo em "01 - SALA CCI"
    with pytest.raises(db.CentroEmUso) as e:
        db.excluir_responsavel(dados, "CCI")
    assert "1 bem" in str(e.value)
    db.desatribuir(dados, "ANA SILVA", 1002)  # agora 1001 e 1002 respondem pelo setor
    with pytest.raises(db.CentroEmUso) as e:
        db.excluir_responsavel(dados, "CCI")
    assert "2 bens" in str(e.value)


def test_incluir_responsavel_e_localizacao(dados):
    semear(dados)
    db.incluir_responsavel(dados, {"ccustos": " decom ", "responsavel": "THIAGO",
                                   "email": "t@cfc", "matricula": "481", "funcao": "gerente"})
    assert db.responsavel(dados, "DECOM")["responsavel"] == "THIAGO"
    db.incluir_localizacao(dados, "99 - SEM MAPA", "DECOM")
    assert db.localizacoes_sem_centro(dados) == []
    with pytest.raises(db.ErroDeNegocio):
        db.incluir_responsavel(dados, {"ccustos": "", "responsavel": "X"})


def test_atribuir_bem_livre_e_transferencia(dados):
    semear(dados)
    db.incluir_pessoa(dados, "  BRUNO LIMA ")
    db.atribuir(dados, "BRUNO LIMA", 1001)
    assert [b["numero"] for b in db.bens_da_pessoa(dados, "BRUNO LIMA")] == [1001]
    with pytest.raises(db.BemNaoEncontrado):
        db.atribuir(dados, "BRUNO LIMA", 9999)
    with pytest.raises(db.JaAtribuido) as e:
        db.atribuir(dados, "BRUNO LIMA", 1002)
    assert e.value.pessoa == "ANA SILVA"
    db.atribuir(dados, "BRUNO LIMA", 1002, confirmar=True)
    assert db.bens_da_pessoa(dados, "ANA SILVA") == []
    assert db.pessoa_do_bem(dados, 1002) == "BRUNO LIMA"
    db.atribuir(dados, "BRUNO LIMA", 1002)  # já é dele: não é erro


def test_desatribuir_e_excluir_pessoa(dados):
    semear(dados)
    db.desatribuir(dados, "ANA SILVA", 1002)
    assert [b["numero"] for b in db.bens_do_centro(dados, "CCI")] == [1001, 1002]
    db.atribuir(dados, "ANA SILVA", 1002)
    db.excluir_pessoa(dados, "ANA SILVA")
    assert db.pessoas(dados) == [] and db.pessoa_do_bem(dados, 1002) is None


def test_excluir_centro_sem_bens_apaga_mapeamentos(dados):
    semear(dados)
    db.incluir_responsavel(dados, {"ccustos": "VAZIO", "responsavel": "X"})
    db.incluir_localizacao(dados, "99 - SEM MAPA", "VAZIO")   # 1004 é ATIVO aqui...
    with pytest.raises(db.CentroEmUso):
        db.excluir_responsavel(dados, "VAZIO")
    dados.execute("UPDATE bens SET situacao='BAIXADO' WHERE numero=1004")
    dados.commit()
    db.excluir_responsavel(dados, "VAZIO")                       # ...sem bens ativos: some, e a sala volta a pendente
    assert db.responsavel(dados, "VAZIO") is None
    assert db.localizacoes_mapeadas(dados) == [{"localizacao": "01 - SALA CCI", "ccustos": "CCI"}]
    with pytest.raises(db.ErroDeNegocio):
        db.excluir_responsavel(dados, "NAO_EXISTE")


def test_atualizar_responsavel(dados):
    semear(dados)
    db.atualizar_responsavel(dados, "CCI", {"responsavel": " carlos ", "email": "c@cfc",
                                           "matricula": "99", "funcao": "gerente"})
    r = db.responsavel(dados, "CCI")
    assert (r["responsavel"], r["funcao"], r["matricula"]) == ("carlos", "gerente", "99")
    with pytest.raises(db.ErroDeNegocio):
        db.atualizar_responsavel(dados, "CCI", {"responsavel": ""})
    with pytest.raises(db.ErroDeNegocio):
        db.atualizar_responsavel(dados, "NAO_EXISTE", {"responsavel": "X"})


def test_mover_localizacoes(dados):
    semear(dados)
    db.incluir_responsavel(dados, {"ccustos": "PRES", "responsavel": "Y"})
    db.incluir_localizacao(dados, "99 - SEM MAPA", "CCI")
    n = db.mover_localizacoes(dados, ["01 - SALA CCI", "99 - SEM MAPA"], "PRES")
    assert n == 2 and {l["ccustos"] for l in db.localizacoes_mapeadas(dados)} == {"PRES"}
    assert [b["numero"] for b in db.bens_do_centro(dados, "PRES")] == [1001, 1004]
    with pytest.raises(db.ErroDeNegocio):
        db.mover_localizacoes(dados, [], "PRES")
    with pytest.raises(db.ErroDeNegocio):
        db.mover_localizacoes(dados, ["01 - SALA CCI"], "NAO_EXISTE")


def test_renomear_pessoa_mantem_atribuicoes(dados):
    semear(dados)
    assert db.renomear_pessoa(dados, "ANA SILVA", " ana  souza ") == "ANA SOUZA"
    assert db.pessoas(dados) == ["ANA SOUZA"] and db.pessoa_do_bem(dados, 1002) == "ANA SOUZA"
    db.incluir_pessoa(dados, "BRUNO")
    with pytest.raises(db.ErroDeNegocio):
        db.renomear_pessoa(dados, "ANA SOUZA", "bruno")
    with pytest.raises(db.ErroDeNegocio):
        db.renomear_pessoa(dados, "NINGUEM", "X")
    assert db.renomear_pessoa(dados, "ANA SOUZA", "ANA SOUZA") == "ANA SOUZA"


def test_exportar_cadastros_quatro_abas(dados, tmp_path):
    semear(dados)
    arq = db.exportar_cadastros(dados, tmp_path / "c.xlsx")
    wb = load_workbook(arq)
    assert wb.sheetnames == ["responsaveis", "localizacoes", "pessoas", "atribuicoes"]
    assert [c.value for c in wb["responsaveis"][1]] == ["ccustos", "responsavel", "email", "matricula", "funcao"]
    assert [c.value for c in wb["pessoas"][1]] == ["nome", "email", "matricula"]
    assert [c.value for c in wb["atribuicoes"][2]] == ["ANA SILVA", 1002]
    assert wb["localizacoes"].max_row == 2 and wb["pessoas"].max_row == 2


def test_importar_a_propria_exportacao_e_idempotente(dados, tmp_path):
    semear(dados)
    arq = db.exportar_cadastros(dados, tmp_path / "c.xlsx")
    resumo = db.importar_cadastros(dados, arq)
    assert resumo == {"responsaveis": 1, "localizacoes": 1, "pessoas": 1, "atribuicoes": 1, "sem_centro": ["99 - SEM MAPA"]}
    assert db.ficha_do_bem(dados, 1002)["pessoa"] == "ANA SILVA"


def cadastros_xlsx(tmp_path, **abas):
    wb = Workbook()
    wb.remove(wb.active)
    cabecalhos = {"responsaveis": ["ccustos", "responsavel", "email", "matricula", "funcao"],
                  "localizacoes": ["localizacao", "ccustos"], "pessoas": ["nome", "email", "matricula"], "atribuicoes": ["nome", "numero"]}
    for aba, cab in cabecalhos.items():
        ws = wb.create_sheet(aba)
        ws.append(cab)
        for linha in abas.get(aba, []):
            ws.append(linha)
    caminho = tmp_path / "cad.xlsx"
    wb.save(caminho)
    return caminho


def test_importar_cadastros_substitui_e_normaliza(dados, tmp_path):
    semear(dados)
    arq = cadastros_xlsx(tmp_path,
                         responsaveis=[["geserv ", " carlos", "c@cfc", 7, "gerente"], ["pres", "MARIA", "", "", ""]],
                         localizacoes=[["01 - SALA CCI", "GESERV"], ["99 - SEM MAPA", "pres"]],
                         pessoas=[[" bruno lima ", "b@cfc", 12]], atribuicoes=[["bruno lima", "1001"]])
    resumo = db.importar_cadastros(dados, arq)
    assert resumo["responsaveis"] == 2 and resumo["atribuicoes"] == 1 and resumo["sem_centro"] == []
    assert db.responsavel(dados, "CCI") is None and db.responsavel(dados, "GESERV")["matricula"] == "7"
    f = db.ficha_do_bem(dados, 1001)
    assert f["ccustos"] == "GESERV" and f["pessoa"] == "BRUNO LIMA"
    assert db.pessoas(dados) == ["BRUNO LIMA"]
    assert dict(db.pessoa(dados, "BRUNO LIMA")) == {"nome": "BRUNO LIMA", "email": "b@cfc", "matricula": "12"}


@pytest.mark.parametrize("abas, trecho", [
    ({"localizacoes": [["01 - SALA CCI", "NAOEXISTE"]]}, "NAOEXISTE"),
    ({"atribuicoes": [["ANA SILVA", 1001]]}, "pessoas"),                     # pessoa fora da aba pessoas
    ({"pessoas": [["ANA"]], "atribuicoes": [["ANA", 9999]]}, "9999"),
    ({"pessoas": [["ANA"], ["BRUNO"]], "atribuicoes": [["ANA", 1001], ["BRUNO", 1001]]}, "repetido"),
    ({"responsaveis": [["", "X", "", "", ""]]}, "sigla"),
    ({"responsaveis": [["A", "", "", "", ""]]}, "responsável"),
    ({"responsaveis": [["A", "X", "", "", ""], ["a", "Y", "", "", ""]]}, "repetid"),
    ({"pessoas": [["ANA"]], "atribuicoes": [["ANA", "1001.5"]]}, "inválido"),
])
def test_importar_cadastros_invalidos_nao_alteram_nada(dados, tmp_path, abas, trecho):
    semear(dados)
    with pytest.raises(db.ImportacaoInvalida) as e:
        db.importar_cadastros(dados, cadastros_xlsx(tmp_path, **abas))
    assert trecho in str(e.value)
    assert db.centros(dados)[0]["ccustos"] == "CCI" and db.pessoas(dados) == ["ANA SILVA"]


def test_importar_cadastros_sem_aba_ou_coluna(dados, tmp_path):
    semear(dados)
    wb = Workbook()
    wb.active.title = "responsaveis"
    wb.active.append(["ccustos", "responsavel"])
    arq = tmp_path / "x.xlsx"
    wb.save(arq)
    with pytest.raises(db.ImportacaoInvalida) as e:
        db.importar_cadastros(dados, arq)
    assert "localizacoes" in str(e.value) or "coluna" in str(e.value)


def test_importar_planilha_truncada_levanta_importacao_invalida(dados, tmp_path):
    semear(dados)
    arq = xlsx(tmp_path, [[1002, "ATIVO", "NOTEBOOK", "DELL", "EQUIP", "01 - SALA CCI", "06/12/2012", 3000, 1400]])
    bruto = arq.read_bytes()
    truncado = tmp_path / "trunc.xlsx"
    truncado.write_bytes(bruto[: int(len(bruto) * 0.6)])
    with pytest.raises(db.ImportacaoInvalida):
        db.importar_bens(dados, truncado)
    with pytest.raises(db.ImportacaoInvalida):
        db.importar_cadastros(dados, truncado)


def test_exportar_bens_formato_spw_reimportavel(dados, tmp_path):
    semear(dados)
    arq = db.exportar_bens(dados, tmp_path / "bens.xlsx")
    ws = load_workbook(arq)["base"]
    assert [c.value for c in ws[1]] == CABECALHO
    assert ws.max_row == 5 and ws["A2"].value == 1001 and ws["F2"].value == "01 - SALA CCI"
    resumo = db.importar_bens(dados, arq)
    assert resumo["total"] == 4 and resumo["ativos"] == 3


def test_pesquisar_texto_simples_e_contem(dados):
    semear(dados)
    r = db.pesquisar(dados, "cci")
    assert [c["ccustos"] for c in r["centros"]] == ["CCI"] and r["centros"][0]["quantidade"] == 1  # 1001; 1002 é da ANA, 1003 baixado
    assert [b["numero"] for b in r["bens"]] == [1001, 1002, 1003]  # pela localização "01 - SALA CCI"
    assert r["pessoas"] == []
    r = db.pesquisar(dados, "ana")
    assert r["pessoas"] == [{"nome": "ANA SILVA", "quantidade": 1}] and r["centros"] == []
    assert db.pesquisar(dados, "jaqueline")["centros"][0]["ccustos"] == "CCI"


def test_pesquisar_acentos_e_pessoa_do_bem(dados):
    semear(dados)
    r = db.pesquisar(dados, "armário")
    assert [b["numero"] for b in r["bens"]] == [1004]
    assert db.pesquisar(dados, "notebook")["bens"][0]["pessoa"] == "ANA SILVA"
    assert db.pesquisar(dados, "cadeira")["bens"][0]["pessoa"] is None


def test_pesquisar_curingas_sao_literais(dados):
    semear(dados)
    dados.execute("INSERT INTO responsaveis (ccustos, responsavel) VALUES ('GEX-ITEC','X'), ('GEX-LIC','Y'), ('AGEX','Z')")
    assert [c["ccustos"] for c in db.pesquisar(dados, "gex")["centros"]] == ["AGEX", "GEX-ITEC", "GEX-LIC"]
    assert [c["ccustos"] for c in db.pesquisar(dados, "gex*")["centros"]] == ["GEX-ITEC", "GEX-LIC"]
    assert [c["ccustos"] for c in db.pesquisar(dados, "%itec")["centros"]] == ["GEX-ITEC"]
    assert [b["numero"] for b in db.pesquisar(dados, "note*ok")["bens"]] == [1002]
    assert [b["numero"] for b in db.pesquisar(dados, "*book")["bens"]] == [1002]
    assert db.pesquisar(dados, "book*")["bens"] == []


def test_pesquisar_varias_palavras_todas_tem_que_bater(dados):
    semear(dados)
    assert [b["numero"] for b in db.pesquisar(dados, "notebook cci")["bens"]] == [1002]  # centro pela localização
    assert [b["numero"] for b in db.pesquisar(dados, "cci ana")["bens"]] == [1002]  # pessoa
    assert db.pesquisar(dados, "cadeira ana")["bens"] == []
    assert [b["numero"] for b in db.pesquisar(dados, "mesa e cci")["bens"]] == [1003]  # "e" é conector
    r = db.pesquisar(dados, "cci jaqueline")
    assert [c["ccustos"] for c in r["centros"]] == ["CCI"] and r["bens"] == []  # nenhum bem casa "jaqueline"
    assert db.pesquisar(dados, "cadeira")["bens"][0]["ccustos"] == "CCI"
    assert db.pesquisar(dados, "armário")["bens"][0]["ccustos"] is None  # localização sem centro


def test_pesquisar_limite(dados):
    semear(dados)
    r = db.pesquisar(dados, "sala", limite=2)
    assert [b["numero"] for b in r["bens"]] == [1001, 1002] and r["truncado"] is True
    assert db.pesquisar(dados, "sala")["truncado"] is False


# ---------------------------------------------------------------- processos SEI
def test_esquema_cria_tabelas_novas(dados):
    nomes = {r[0] for r in dados.execute("SELECT name FROM sqlite_master WHERE type='table'")}
    assert {"processos_sei", "termos_emitidos", "termos_emitidos_bens", "importacoes", "importacoes_mudancas"} <= nomes


def test_processos_um_vigente_por_tipo(dados):
    a = db.incluir_processo(dados, "ccusto", "Termos 2025", "1111")
    b = db.incluir_processo(dados, "ccusto", "Termos 2026", "2222")
    c = db.incluir_processo(dados, "individual", "Individuais 2026", "3333")
    assert db.processo_vigente(dados, "ccusto")["id"] == b
    assert db.processo_vigente(dados, "individual")["id"] == c
    assert db.processo_vigente(dados, "devolucao") is None
    db.marcar_vigente(dados, a)
    assert db.processo_vigente(dados, "ccusto")["id"] == a
    db.encerrar_processo(dados, a)
    assert db.processo_vigente(dados, "ccusto") is None
    assert [p["id"] for p in db.processos(dados)][0] == c          # vigentes primeiro
    with pytest.raises(sqlite3.IntegrityError):
        dados.execute("INSERT INTO processos_sei (tipo, descricao, numero_sei, vigente, criado_em) VALUES ('individual','x','9',1,'2026-01-01 00:00:00')")


def test_processos_validacao_e_exclusao(dados):
    with pytest.raises(db.ErroDeNegocio):
        db.incluir_processo(dados, "outro", "x", "1")
    with pytest.raises(db.ErroDeNegocio):
        db.incluir_processo(dados, "ccusto", "", "1")
    with pytest.raises(db.ErroDeNegocio):
        db.incluir_processo(dados, "ccusto", "x", "  ")
    i = db.incluir_processo(dados, "ccusto", "x", "1", vigente=False)
    assert db.processo_vigente(dados, "ccusto") is None
    db.excluir_processo(dados, i)
    assert db.processos(dados) == []


# ---------------------------------------------------------------- termos emitidos
def test_registrar_emissao_exige_processo_e_grava_foto(dados):
    semear(dados)
    bens = db.bens_do_centro(dados, "CCI")
    with pytest.raises(db.ErroDeNegocio):
        db.registrar_emissao(dados, "ccusto", "CCI", bens)
    db.incluir_processo(dados, "ccusto", "Termos 2026", "2222")
    t = db.registrar_emissao(dados, "ccusto", "CCI", bens)
    assert t["quantidade"] == 1 and t["valor_total"] == 64.54 and t["numero_sei"] == "2222"
    assert [b["numero"] for b in t["bens"]] == [1001] and t["bens"][0]["descricao"] == "CADEIRA"
    assert db.ultimo_termo(dados, "ccusto", "CCI")["id"] == t["id"]
    assert db.termos_emitidos(dados)[0]["id"] == t["id"]
    assert db.termos_emitidos(dados, tipo="individual") == []
    assert db.termos_emitidos(dados, chave="cc")[0]["id"] == t["id"]
    db.salvar_documento_sei(dados, t["id"], " 0451234 ")
    assert db.termo_emitido(dados, t["id"])["documento_sei"] == "0451234"


def test_registrar_emissao_mesmo_dia_mesma_lista_nao_duplica(dados):
    semear(dados)
    db.incluir_processo(dados, "ccusto", "Termos 2026", "2222")
    a = db.registrar_emissao(dados, "ccusto", "CCI", db.bens_do_centro(dados, "CCI"))
    b = db.registrar_emissao(dados, "ccusto", "CCI", db.bens_do_centro(dados, "CCI"))
    assert a["id"] == b["id"] and len(db.termos_emitidos(dados)) == 1
    dados.execute("INSERT INTO bens VALUES (1005,'ATIVO','LUMINÁRIA','','MÓVEIS','01 - SALA CCI','01/01/2020',10,9)")
    c = db.registrar_emissao(dados, "ccusto", "CCI", db.bens_do_centro(dados, "CCI"))
    assert c["id"] != a["id"] and c["quantidade"] == 2


def test_registrar_emissao_com_outro_processo_cria_termo_novo(dados):
    semear(dados)
    db.incluir_processo(dados, "ccusto", "Termos 2026", "2222")
    a = db.registrar_emissao(dados, "ccusto", "CCI", db.bens_do_centro(dados, "CCI"))
    db.incluir_processo(dados, "ccusto", "Termos 2026 (novo)", "3333")   # passa a ser o vigente
    b = db.registrar_emissao(dados, "ccusto", "CCI", db.bens_do_centro(dados, "CCI"))
    assert b["id"] != a["id"] and b["numero_sei"] == "3333" and db.termo_emitido(dados, a["id"])["numero_sei"] == "2222"
    assert len(db.termos_emitidos(dados)) == 2


def test_situacao_termo(dados):
    semear(dados)
    bens = db.bens_do_centro(dados, "CCI")
    assert db.situacao_termo(dados, "ccusto", "CCI", bens)["estado"] == "sem_termo"
    db.incluir_processo(dados, "ccusto", "Termos 2026", "2222")
    db.registrar_emissao(dados, "ccusto", "CCI", bens)
    assert db.situacao_termo(dados, "ccusto", "CCI", db.bens_do_centro(dados, "CCI"))["estado"] == "vigente"
    dados.execute("INSERT INTO bens VALUES (1005,'ATIVO','LUMINÁRIA','','MÓVEIS','01 - SALA CCI','01/01/2020',10,9)")
    dados.execute("UPDATE bens SET situacao = 'BAIXADO' WHERE numero = 1001")
    s = db.situacao_termo(dados, "ccusto", "CCI", db.bens_do_centro(dados, "CCI"))
    assert (s["estado"], s["entraram"], s["sairam"]) == ("desatualizado", 1, 1)
    centros = db.situacoes_centros(dados)
    assert centros[0]["ccustos"] == "CCI" and centros[0]["estado"] == "desatualizado" and centros[0]["quantidade"] == 1
    pessoas = db.situacoes_pessoas(dados)
    assert pessoas == [{"nome": "ANA SILVA", "quantidade": 1, "valor": 1500.0, "estado": "sem_termo",
                        "ultimo": None, "entraram": 0, "sairam": 0}]


def test_renomear_leva_historico_junto(dados):
    semear(dados)
    db.incluir_processo(dados, "ccusto", "T", "1")
    db.incluir_processo(dados, "individual", "I", "2")
    db.registrar_emissao(dados, "ccusto", "CCI", db.bens_do_centro(dados, "CCI"))
    db.registrar_emissao(dados, "individual", "ANA SILVA", db.bens_da_pessoa(dados, "ANA SILVA"))
    db.renomear_centro(dados, "CCI", "GEX")
    db.renomear_pessoa(dados, "ANA SILVA", "ANA SOUZA")
    assert db.ultimo_termo(dados, "ccusto", "GEX") and db.ultimo_termo(dados, "ccusto", "CCI") is None
    assert db.ultimo_termo(dados, "individual", "ANA SOUZA") and db.ultimo_termo(dados, "individual", "ANA SILVA") is None


# ---------------------------------------------------------------- painel e recorte
def _semear_painel(conn):
    semear(conn)
    conn.execute("INSERT INTO bens VALUES (2001,'ATIVO','SEDE','','SEDE','','01/01/1990',1,60000000)")
    conn.execute("INSERT INTO bens VALUES (2002,'ATIVO','MICRO','HP','EQUIPAMENTOS','01 - SALA CCI','15/06/2024',4000,3500)")
    conn.execute("INSERT INTO bens VALUES (2003,'ATIVO','TABLET','','EQUIPAMENTOS','TERMOS INDIVIDUAIS','01/02/2023',900,800)")
    conn.execute("INSERT INTO atribuicoes VALUES ('ANA SILVA', 2003)")
    conn.commit()


def test_dimensoes_sobre_ativos(dados):
    _semear_painel(dados)
    d = db.dimensoes(dados, {"situacao": "ATIVO"})
    assert {x["chave"]: x["quantidade"] for x in d["situacao"]} == {"ATIVO": 6, "BAIXADO": 1}   # situação ignora o próprio filtro
    centro = {x["chave"]: (x["quantidade"], x["rotulo"]) for x in d["centro"]}
    assert centro["CCI"] == (3, "CCI – JAQUELINE PORTELA") and centro["-"] == (3, "sem centro")
    assert {x["chave"]: x["quantidade"] for x in d["classificacao"]} == {"MÓVEIS": 2, "EQUIPAMENTOS": 3, "SEDE": 1}
    assert [x["chave"] for x in d["idade"]] == ["ate5", "5a10", "10a20", "mais20", "semdata"]
    assert {x["chave"]: x["quantidade"] for x in d["idade"] if x["quantidade"]} == {"ate5": 2, "10a20": 2, "mais20": 2}
    assert [x["chave"] for x in d["ano"]] == ["1990", "1996", "2012", "2023", "2024"]
    assert {x["chave"]: x["quantidade"] for x in d["faixa"] if x["quantidade"]} == {"ate100": 1, "100a500": 1, "500a1000": 1, "1000a5000": 2, "mais20000": 1}
    assert d["pessoa"] == [{"chave": "ANA SILVA", "rotulo": "ANA SILVA", "quantidade": 2, "valor": 2300.0}]
    assert any(x["chave"] == "01 - SALA CCI" and "(CCI)" in x["rotulo"] and x["quantidade"] == 3 for x in d["localizacao"])


def test_painel_cards(dados):
    _semear_painel(dados)
    p = db.painel(dados)
    assert p["ativos"] == 6 and p["imoveis"] == 1 and p["valor_imoveis"] == 60000000
    assert round(p["valor_sem_imoveis"], 2) == 64.54 + 1500 + 250.5 + 3500 + 800
    assert p["sem_centro"] == 2 and p["sem_valor"] == 0 and p["ultima_importacao"] is None
    assert p["a_emitir_centros"] == 1 and p["a_emitir_pessoas"] == 1
    assert p["centros"][0]["ccustos"] == "CCI" and "dimensoes" in p
    assert db.recorte(dados, {"situacao": "ATIVO", "ccusto": "-", "pessoa": "-"})["quantidade"] == p["sem_centro"]


def test_recorte_filtros_e_drill_down(dados):
    _semear_painel(dados)
    r = db.recorte(dados, {"situacao": "ATIVO", "ccusto": "CCI"})
    assert [b["numero"] for b in r["bens"]] == [1001, 1002, 2002] and r["quantidade"] == 3 and r["bens"][1]["pessoa"] == "ANA SILVA"
    assert [b["numero"] for b in db.recorte(dados, {"situacao": "ATIVO", "ccusto": "CCI", "faixa": "1000a5000"})["bens"]] == [1002, 2002]
    assert [b["numero"] for b in db.recorte(dados, {"ccusto": "-"})["bens"]] == [1004, 2001, 2003]
    assert [b["numero"] for b in db.recorte(dados, {"classificacao": "imoveis"})["bens"]] == [2001]
    assert 2001 not in [b["numero"] for b in db.recorte(dados, {"classificacao": "sem-imoveis"})["bens"]]
    assert [b["numero"] for b in db.recorte(dados, {"valor_de": "1000", "valor_ate": "2000"})["bens"]] == [1002]
    assert [b["numero"] for b in db.recorte(dados, {"entrada_de": "2012-01-01", "entrada_ate": "2012-12-31"})["bens"]] == [1002, 1003, 1004]
    assert [b["numero"] for b in db.recorte(dados, {"ano": "2024"})["bens"]] == [2002]
    assert [b["numero"] for b in db.recorte(dados, {"idade": "mais20"})["bens"]] == [1001, 2001]
    assert [b["numero"] for b in db.recorte(dados, {"pessoa": "ANA SILVA"})["bens"]] == [1002, 2003]
    assert [b["numero"] for b in db.recorte(dados, {"localizacao": "99 - SEM MAPA"})["bens"]] == [1004]
    assert db.recorte(dados, {})["quantidade"] == 7          # sem filtro = tudo (a rota põe ATIVO por padrão)
    r = db.recorte(dados, {}, limite=2)
    assert len(r["bens"]) == 2 and r["truncado"] and r["quantidade"] == 7


def test_dimensoes_balde_vazio_tem_sentinela_e_filtra(dados):
    semear(dados)
    dados.execute("INSERT INTO bens VALUES (3001,'ATIVO','SEM NADA','','','','',1,1)")
    d = db.dimensoes(dados, {"situacao": "ATIVO"})
    assert any(x["chave"] == "-" and x["rotulo"] == "sem classificação" and x["quantidade"] == 1 for x in d["classificacao"])
    assert any(x["chave"] == "-" and x["rotulo"] == "sem localização" for x in d["localizacao"])
    assert any(x["chave"] == "-" and x["rotulo"] == "sem data" for x in d["ano"])
    assert [b["numero"] for b in db.recorte(dados, {"classificacao": "-"})["bens"]] == [3001]
    assert [b["numero"] for b in db.recorte(dados, {"localizacao": "-"})["bens"]] == [3001]
    assert [b["numero"] for b in db.recorte(dados, {"ano": "-"})["bens"]] == [3001]


def test_exportar_recorte_xlsx(dados, tmp_path):
    _semear_painel(dados)
    from openpyxl import load_workbook
    ws = load_workbook(db.exportar_recorte(dados, {"ccusto": "CCI"}, tmp_path / "r.xlsx")).active
    linhas = list(ws.iter_rows(values_only=True))
    assert linhas[0][:3] == ("Número", "Descrição", "Complemento") and len(linhas) == 5   # cabeçalho + 4 bens (inclui 1003 BAIXADO)
    assert linhas[1][5] == "CCI"


def test_migracao_de_banco_antigo(tmp_path):
    """Banco de antes de 2026-09-16: responsaveis com tratamento, pessoas só com nome, termos sem bloco."""
    import sqlite3
    caminho = tmp_path / "antigo.db"
    velho = sqlite3.connect(caminho)
    velho.executescript("""
        CREATE TABLE responsaveis (ccustos TEXT PRIMARY KEY, tratamento TEXT, responsavel TEXT NOT NULL, email TEXT, matricula TEXT, funcao TEXT);
        CREATE TABLE pessoas (nome TEXT PRIMARY KEY);
        CREATE TABLE processos_sei (id INTEGER PRIMARY KEY, tipo TEXT NOT NULL, descricao TEXT NOT NULL, numero_sei TEXT NOT NULL, vigente INTEGER NOT NULL DEFAULT 0, criado_em TEXT NOT NULL);
        CREATE TABLE termos_emitidos (id INTEGER PRIMARY KEY, tipo TEXT NOT NULL, chave TEXT NOT NULL, processo_id INTEGER NOT NULL,
            documento_sei TEXT, emitido_em TEXT NOT NULL, quantidade INTEGER NOT NULL, valor_total REAL NOT NULL);
        INSERT INTO responsaveis VALUES ('CCI', 'Prezada', 'JAQUELINE', 'j@cfc', '46', 'coordenadora');
        INSERT INTO pessoas VALUES ('ANA SILVA');
    """)
    velho.commit()
    velho.close()
    conn = db.conectar(caminho)
    db.criar_esquema(conn)
    db.criar_esquema(conn)   # idempotente
    assert dict(db.responsavel(conn, "CCI")) == {"ccustos": "CCI", "responsavel": "JAQUELINE", "email": "j@cfc", "matricula": "46", "funcao": "coordenadora"}
    assert dict(db.pessoa(conn, "ANA SILVA")) == {"nome": "ANA SILVA", "email": None, "matricula": None}
    assert "bloco_sei" in db._colunas(conn, "termos_emitidos") and "email_enviado_em" in db._colunas(conn, "termos_emitidos")


def test_pessoa_email_matricula_e_salvar(dados):
    semear(dados)
    assert db.incluir_pessoa(dados, " bruno lima ", " b@cfc ", "0012") == "BRUNO LIMA"
    assert dict(db.pessoa(dados, "BRUNO LIMA")) == {"nome": "BRUNO LIMA", "email": "b@cfc", "matricula": "0012"}
    novo = db.salvar_pessoa(dados, "BRUNO LIMA", {"nome": "bruno souza", "email": "", "matricula": "7"})
    assert novo == "BRUNO SOUZA" and db.pessoa(dados, "BRUNO LIMA") is None
    assert dict(db.pessoa(dados, "BRUNO SOUZA")) == {"nome": "BRUNO SOUZA", "email": None, "matricula": "7"}
    with pytest.raises(db.ErroDeNegocio):
        db.salvar_pessoa(dados, "BRUNO SOUZA", {"nome": "ana silva"})


def test_salvar_pessoa_e_tudo_ou_nada(dados):
    import sqlite3
    semear(dados)
    dados.execute("CREATE TRIGGER falha BEFORE UPDATE OF email ON pessoas BEGIN SELECT RAISE(ABORT, 'falha simulada'); END")
    with pytest.raises(sqlite3.IntegrityError):
        db.salvar_pessoa(dados, "ANA SILVA", {"nome": "ana souza", "email": "a@cfc", "matricula": "1"})
    outra = db.conectar()   # o que outra conexão enxerga = o que foi commitado
    assert db.pessoa(outra, "ANA SILVA") is not None and db.pessoa(outra, "ANA SOUZA") is None
    assert db.pessoa_do_bem(outra, 1002) == "ANA SILVA"
    outra.close()


def test_bloco_sei_e_registro_de_email(dados):
    semear(dados)
    db.incluir_processo(dados, "ccusto", "T", "1111")
    t = db.registrar_emissao(dados, "ccusto", "CCI", db.bens_do_centro(dados, "CCI"))
    db.salvar_documento_sei(dados, t["id"], "0451234", " 55 ")
    t = db.termo_emitido(dados, t["id"])
    assert (t["documento_sei"], t["bloco_sei"], t["email_enviado_em"]) == ("0451234", "55", None)
    quando = db.registrar_email(dados, t["id"])
    assert db.termo_emitido(dados, t["id"])["email_enviado_em"] == quando


def test_exportar_grava_texto_literal_e_numeros_como_numeros(dados, tmp_path):
    from openpyxl import load_workbook
    semear(dados)
    dados.execute("INSERT INTO bens VALUES (1005,'ATIVO','=1+1','=SOMA(A1)','MÓVEIS','01 - SALA CCI','01/01/2020',10,9)")
    dados.commit()
    wb = load_workbook(db.exportar_bens(dados, tmp_path / "bens.xlsx"))
    linha = [c for c in wb["base"].iter_rows(min_row=2) if c[0].value == 1005][0]
    assert linha[2].value == "=1+1" and linha[2].data_type == "s"       # texto, não fórmula
    assert linha[3].value == "=SOMA(A1)" and linha[3].data_type == "s"
    assert linha[0].data_type == "n" and linha[8].data_type == "n"       # número continua número
    wb = load_workbook(db.exportar_recorte(dados, {}, tmp_path / "recorte.xlsx"))
    assert all(c.data_type != "f" for linha in wb["recorte"].iter_rows(min_row=2) for c in linha)
