import io
from datetime import date
from urllib.parse import unquote

import pytest
from openpyxl import Workbook, load_workbook

import db
from tests.conftest import semear, confirmar_revisao, logar, ADMIN_LOGIN, ADMIN_NOME, ADMIN_SENHA, SENHA_PADRAO
from tests.test_db import CABECALHO
from tests.test_permissoes import NEGADO


def _acesso_admin(dados, chave):
    import usuarios
    usuarios.salvar_acesso_sei(dados, usuarios.por_login(dados, ADMIN_LOGIN)["id"], "antonio.junior", "S3nha!", "GELIC")


def _acesso_spw_admin(dados, chave):
    import usuarios
    usuarios.salvar_acesso_spw(dados, usuarios.por_login(dados, ADMIN_LOGIN)["id"], "antonio.junior", "S3nha!")


def test_home_e_busca_de_bem(cliente):
    assert cliente.get("/").status_code == 200
    r = cliente.get("/bem?numero=1002")
    assert b"NOTEBOOK" in r.data and b"ANA SILVA" in r.data and b"JAQUELINE PORTELA" in r.data
    r = cliente.get("/bem?numero=9999", follow_redirects=True)
    assert "não encontrado".encode() in r.data


def test_inicio_nao_calcula_graficos(cliente, monkeypatch):
    import db
    def proibido(*args, **kwargs): raise AssertionError('Início não usa dimensões')
    monkeypatch.setattr(db, 'dimensoes', proibido)
    html = cliente.get('/').get_data(as_text=True)
    conteudo = html.split('id="main-content"', 1)[1].split('</main>', 1)[0]
    assert '<form' not in conteudo
    assert 'echarts.min.js' not in html
    assert 'Termos por centro de custo' not in conteudo
    assert 'Sobre o patrimônio' in conteudo and 'Ver guia' in conteudo


def test_bem_mostra_fotos_por_evento(cliente):
    assert b"Fotos do invent" not in cliente.get("/bem?numero=1001").data
    eid = _abrir(cliente)
    cliente.post(f"/inventario/{eid}/sala/01 - SALA CCI/ler", json={"numero": "1001"})
    import db, inventario
    conn = db.conectar()
    inventario.adicionar_foto(conn, eid, 1001, lambda c: "https://x/a.webp")
    inventario.adicionar_foto(conn, eid, 1001, lambda c: "https://x/b.webp")
    cliente.post(f"/inventario/{eid}/encerrar", data={"confirmar": "1"})
    cliente.post("/inventario/abrir", data={"nome": "Inv 2", "usuarios": _ids("Fulano"), "escopo": "todas"})
    eid2 = inventario.evento_aberto(conn)["id"]
    cliente.post(f"/inventario/{eid2}/sala/01 - SALA CCI/ler", json={"numero": "1001"})
    inventario.adicionar_foto(conn, eid2, 1001, lambda c: "https://x/c.webp")
    r = cliente.get("/bem?numero=1001")
    assert b"Fotos do invent" in r.data and r.data.count(b'class="dsgov-miniatura"') == 3
    assert r.data.index(b"Inv 2") < r.data.index(b">Inv<") and r.data.index(b"https://x/c.webp") < r.data.index(b"https://x/a.webp") < r.data.index(b"https://x/b.webp")
    assert b"encerrado em" in r.data and b"lido em" in r.data
    assert b"Fotos do invent" not in cliente.get("/bem?numero=1002").data


def test_bem_so_mostra_fotos_dos_eventos_visiveis(cliente, dados):
    """Quem soma Consulta e Inventário vê, na ficha do bem, só as fotos dos eventos de que participa."""
    import inventario
    import usuarios
    uid = usuarios.criar(dados, "mista", "Mista", SENHA_PADRAO, ["consulta", "inventariante"], trocar_senha=False)
    eid = _abrir(cliente, comissao=["Fulano"])                       # evento alheio a 'Mista'
    cliente.post(f"/inventario/{eid}/sala/01 - SALA CCI/ler", json={"numero": "1001"})
    inventario.adicionar_foto(dados, eid, 1001, lambda c: "https://x/alheia.webp")
    cliente.post("/sair"); logar(cliente, "mista", SENHA_PADRAO)
    r = cliente.get("/bem?numero=1001")
    assert r.status_code == 200 and b"Fotos do invent" not in r.data and b"alheia.webp" not in r.data
    cliente.post("/sair"); logar(cliente, ADMIN_LOGIN, ADMIN_SENHA)
    cliente.post(f"/inventario/{eid}/comissao", data={"usuarios": _ids("Fulano") + [uid]})
    cliente.post("/sair"); logar(cliente, "mista", SENHA_PADRAO)
    r = cliente.get("/bem?numero=1001")
    assert b"Fotos do invent" in r.data and b"alheia.webp" in r.data


def test_termo_ccusto_documento_docx_planilha(cliente):
    cliente.post("/cadastros/processos/incluir", data={"tipo": "ccusto", "descricao": "T", "numero_sei": "1111", "vigente": "1"})
    r = cliente.post("/gerar", data={"ccusto": "CCI"})
    assert r.status_code == 302 and r.headers["Location"].endswith("/termo/ccusto/CCI")
    assert cliente.get("/termo/ccusto/CCI").status_code == 200
    doc = cliente.get("/termo/ccusto/CCI/documento").data.decode()
    assert "width:90%" in doc and "1001" in doc and "1002" not in doc  # 1002 está com ANA
    assert cliente.get("/termo/ccusto/CCI/docx").headers["Content-Disposition"].endswith('Termo_de_Responsabilidade_CCI.docx')
    assert cliente.get("/termo/ccusto/CCI/planilha").headers["Content-Disposition"].endswith('planilha_CCI.xlsx')


def test_termo_individual(cliente):
    cliente.post("/cadastros/processos/incluir", data={"tipo": "individual", "descricao": "T", "numero_sei": "1111", "vigente": "1"})
    r = cliente.post("/gerar-individual", data={"nome": "ANA SILVA"})
    assert r.status_code == 302
    doc = cliente.get("/termo/individual/ANA SILVA/documento").data.decode()
    assert "width:90%" in doc and "NOTEBOOK" in doc
    assert cliente.get("/termo/individual/ANA SILVA/docx").status_code == 200
    assert cliente.get("/termo/individual/NINGUEM").status_code == 404


def test_termo_sem_processo_vigente_nao_emite(cliente):
    r = cliente.get("/termo/ccusto/CCI")
    assert b"Cadastre um processo SEI vigente" in r.data and b'id="copiar"' not in r.data
    r = cliente.get("/termo/ccusto/CCI/docx", follow_redirects=True)
    assert b"Cadastre um processo SEI vigente" in r.data
    assert cliente.get("/termo/ccusto/CCI/documento").status_code == 200    # prévia continua
    r = cliente.post("/termo/ccusto/CCI/registrar")
    assert r.status_code == 409 and "erro" in r.get_json()


def test_termo_com_processo_registra_ao_baixar_e_ao_copiar(cliente):
    import db
    cliente.post("/cadastros/processos/incluir", data={"tipo": "ccusto", "descricao": "T", "numero_sei": "2222", "vigente": "1"})
    assert cliente.get("/termo/ccusto/CCI/planilha").status_code == 200
    assert db.termos_emitidos(db.conectar()) == []     # planilha não registra emissão
    r = cliente.get("/termo/ccusto/CCI")
    assert b'id="copiar"' in r.data and b"Nenhum termo registrado" in r.data
    assert cliente.get("/termo/ccusto/CCI/docx").status_code == 200
    r = cliente.get("/termo/ccusto/CCI")
    assert "Último termo registrado".encode() in r.data and b"Bens iguais aos de hoje" in r.data
    j = cliente.post("/termo/ccusto/CCI/registrar").get_json()
    assert j["id"] and j["emitido_em"][:4] == str(date.today().year)
    assert len(db.termos_emitidos(db.conectar())) == 1     # docx + registrar no mesmo dia = 1 registro


def test_termo_devolucao_registra_no_processo_de_devolucao(cliente):
    cliente.post("/cadastros/processos/incluir", data={"tipo": "devolucao", "descricao": "D", "numero_sei": "3333", "vigente": "1"})
    cliente.post("/termo_devolucao", data={"nome": "ANA SILVA", "numero_bem": "1001"})
    assert cliente.get("/termo/devolucao/ANA SILVA/docx").status_code == 200
    import db
    t = db.termos_emitidos(db.conectar(), tipo="devolucao")[0]
    assert t["chave"] == "ANA SILVA" and t["quantidade"] == 1


def test_upload_importa_e_lista_sem_centro(cliente):
    wb = Workbook()
    ws = wb.active
    ws.append(CABECALHO)
    ws.append([1002, "ATIVO", "NOTEBOOK", "DELL", "EQ", "01 - SALA CCI", "x", 1, 1])
    ws.append([5000, "ATIVO", "TV", "", "EQ", "77 - NOVA SALA", "x", 1, 1])
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    r = cliente.post("/upload", data={"arquivo": (buf, "export.xlsx")}, content_type="multipart/form-data",
                     follow_redirects=True)
    assert "2 bens importados (2 ativos)".encode() in r.data and b"77 - NOVA SALA" in r.data


def test_upload_mostra_mudancas_e_detalhe(cliente, tmp_path):
    from tests.test_db import xlsx
    arq = xlsx(tmp_path, [
        [1001, "ATIVO", "CADEIRA", "GIRATÓRIA", "MÓVEIS", "02 - OUTRA", "31/12/1996", 75.94, 64.54],
        [1002, "ATIVO", "NOTEBOOK", "DELL", "EQUIPAMENTOS", "01 - SALA CCI", "06/12/2012", 3000, 1500],
    ])
    with open(arq, "rb") as f:
        r = cliente.post("/upload", data={"arquivo": (f, "export.xlsx")}, content_type="multipart/form-data", follow_redirects=True)
    assert b"2 bens importados" in r.data and b"1 movido" in r.data and b"2 removido" in r.data and b"export.xlsx" in r.data
    import db
    iid = db.importacoes(db.conectar())[0]["id"]
    r = cliente.get(f"/importacoes/{iid}")
    assert r.status_code == 200 and b"02 - OUTRA" in r.data and b"1004" in r.data and b"removido" in r.data
    r = cliente.get("/bem?numero=1001")
    assert b"02 - OUTRA" in r.data and b"movido" in r.data


def test_upload_invalido_mostra_erro(cliente):
    r = cliente.post("/upload", data={"arquivo": (io.BytesIO(b"nada"), "x.txt")}, content_type="multipart/form-data",
                     follow_redirects=True)
    assert b".xlsx" in r.data


def test_cadastro_responsaveis_editar_renomear_excluir(cliente):
    r = cliente.post("/cadastros/responsaveis/incluir", data={"ccustos": "decom",
                     "responsavel": "THIAGO", "email": "", "matricula": "481", "funcao": "gerente"}, follow_redirects=True)
    assert b"DECOM" in r.data
    r = cliente.get("/cadastros/responsaveis/CCI/editar")
    assert b"JAQUELINE PORTELA" in r.data and b"01 - SALA CCI" in r.data
    r = cliente.post("/cadastros/responsaveis/CCI/editar", data={"ccustos": "geserv",
                     "responsavel": "CARLOS", "email": "", "matricula": "7", "funcao": "gerente"}, follow_redirects=True)
    assert b"GESERV" in r.data and b"CARLOS" in r.data and b">CCI<" not in r.data
    assert cliente.get("/cadastros/responsaveis/CCI/editar").status_code == 404
    # excluir: pede confirmação; com bens sob guarda, bloqueia
    r = cliente.post("/cadastros/responsaveis/excluir", data={"ccustos": "GESERV"}, follow_redirects=True)
    assert b"sob guarda" in r.data
    r = cliente.post("/cadastros/responsaveis/excluir", data={"ccustos": "DECOM"}, follow_redirects=True)
    assert b"Confirmar exclus" in r.data
    r = confirmar_revisao(cliente, r, "/cadastros/responsaveis/excluir")
    assert b">DECOM<" not in r.data


def test_editar_responsavel_invalido_nao_renomeia(cliente, dados):
    r = cliente.post("/cadastros/responsaveis/CCI/editar", data={"ccustos": "GESERV", "responsavel": ""},
                     follow_redirects=True)
    assert b"obrigat" in r.data
    assert db.responsavel(dados, "CCI") is not None
    assert db.responsavel(dados, "GESERV") is None


def test_cadastro_localizacoes_mover(cliente):
    cliente.post("/cadastros/responsaveis/incluir", data={"ccustos": "PRES", "responsavel": "Y"})
    cliente.post("/cadastros/localizacoes/incluir", data={"localizacao": "99 - SEM MAPA", "ccustos": "CCI"})
    r = cliente.post("/cadastros/localizacoes/mover", data={"localizacoes": ["01 - SALA CCI", "99 - SEM MAPA"], "ccustos_destino": "PRES"},
                     follow_redirects=True)
    assert b"2 localiza" in r.data
    assert b"PRES" not in cliente.get("/bem?numero=1001").data
    r = confirmar_revisao(cliente, r, "/cadastros/localizacoes/mover")
    assert b"PRES" in cliente.get("/bem?numero=1001").data
    r = cliente.post("/cadastros/localizacoes/mover", data={"ccustos_destino": "PRES"}, follow_redirects=True)
    assert b"ao menos uma" in r.data


def test_cadastro_pessoas_editar_nome(cliente):
    r = cliente.get("/cadastros/pessoas/ANA SILVA/editar")
    assert b"ANA SILVA" in r.data
    r = cliente.post("/cadastros/pessoas/ANA SILVA/editar", data={"nome": "ana souza"}, follow_redirects=True)
    assert b"ANA SOUZA" in r.data and b"NOTEBOOK" in r.data
    assert cliente.get("/cadastros/pessoas/NINGUEM/editar").status_code == 404


def test_cadastro_localizacoes(cliente):
    r = cliente.get("/cadastros/localizacoes")
    assert b"99 - SEM MAPA" in r.data
    r = cliente.post("/cadastros/localizacoes/incluir", data={"localizacao": "99 - SEM MAPA", "ccustos": "CCI"}, follow_redirects=True)
    assert r.data.count(b"99 - SEM MAPA") >= 1 and b"pendente" not in r.data.lower()
    r = cliente.post("/cadastros/localizacoes/excluir", data={"localizacao": "99 - SEM MAPA"})
    confirmar_revisao(cliente, r, "/cadastros/localizacoes/excluir")
    assert b"99 - SEM MAPA" in cliente.get("/cadastros/localizacoes").data


def test_cadastro_pessoas_atribuir_com_confirmacao(cliente):
    cliente.post("/cadastros/pessoas/incluir", data={"nome": "bruno lima"})
    r = cliente.get("/cadastros/pessoas?nome=BRUNO LIMA")
    assert b"BRUNO LIMA" in r.data
    r = cliente.post("/cadastros/pessoas/atribuir", data={"nome": "BRUNO LIMA", "numero": "1002"}, follow_redirects=True)
    assert b"ANA SILVA" in r.data and b"confirmar" in r.data  # pede confirmação
    r = confirmar_revisao(cliente, r, "/cadastros/pessoas/atribuir")
    assert b"NOTEBOOK" in r.data
    r = cliente.post("/cadastros/pessoas/desatribuir", data={"nome": "BRUNO LIMA", "numero": "1002"}, follow_redirects=True)
    r = confirmar_revisao(cliente, r, "/cadastros/pessoas/desatribuir")
    assert b"NOTEBOOK" not in r.data
    r = cliente.post("/cadastros/pessoas/atribuir", data={"nome": "BRUNO LIMA", "numero": "9999"}, follow_redirects=True)
    assert "não encontrado".encode() in r.data
    r = cliente.post("/cadastros/pessoas/excluir", data={"nome": "BRUNO LIMA"}, follow_redirects=True)
    assert b"Confirmar" in r.data
    confirmar_revisao(cliente, r, "/cadastros/pessoas/excluir")
    assert b"BRUNO LIMA" not in cliente.get("/cadastros/pessoas").data


def test_termo_devolucao_fluxo(cliente):
    cliente.post("/cadastros/processos/incluir", data={"tipo": "devolucao", "descricao": "T", "numero_sei": "1111", "vigente": "1"})
    r = cliente.post("/termo_devolucao", data={"nome": "ANA SILVA", "numero_bem": "1001"}, follow_redirects=True)
    assert b"CADEIRA" in r.data
    r = cliente.post("/termo_devolucao", data={"nome": "ANA SILVA", "numero_bem": "9999"}, follow_redirects=True)
    assert "não encontrado".encode() in r.data
    r = cliente.post("/termo_devolucao", data={"nome": "ANA SILVA", "gerar": "1"})
    assert unquote(r.headers["Location"]).endswith("/termo/devolucao/ANA SILVA")
    doc = cliente.get("/termo/devolucao/ANA SILVA/documento").data.decode()
    assert "TERMO DE DEVOLUÇÃO" in doc and "CADEIRA" in doc and "width:90%" in doc
    assert cliente.get("/termo/devolucao/ANA SILVA/docx").status_code == 200
    r = cliente.post("/termo_devolucao", data={"nome": "ANA SILVA", "remover": "1001"}, follow_redirects=True)
    assert b"CADEIRA" not in r.data


def test_termo_devolucao_normaliza_numero(cliente):
    cliente.post("/termo_devolucao", data={"nome": "ANA SILVA", "numero_bem": "1001"})
    r = cliente.post("/termo_devolucao", data={"nome": "ANA SILVA", "numero_bem": "01001"}, follow_redirects=True)
    assert r.data.count(b"CADEIRA") == 1
    r = cliente.post("/termo_devolucao", data={"nome": "ANA SILVA", "remover": "1001"}, follow_redirects=True)
    assert b"CADEIRA" not in r.data


def test_termo_devolucao_troca_de_pessoa_limpa_lista(cliente, dados):
    db.incluir_pessoa(dados, "BRUNO LIMA")
    cliente.post("/termo_devolucao", data={"nome": "ANA SILVA", "numero_bem": "1001"})
    r = cliente.post("/termo_devolucao", data={"nome": "BRUNO LIMA", "numero_bem": "1002"}, follow_redirects=True)
    assert b"CADEIRA" not in r.data and b"NOTEBOOK" in r.data


def test_textos_salvar_reflete_no_documento_e_restaurar(cliente):
    r = cliente.get("/textos")
    assert r.status_code == 200 and b"Compromissos" in r.data
    import textos
    dados_form = dict(textos.PADRAO, individual_abertura="TESTE {nome}.", orgao_nome="Órgão X", unidade_sigla="SN")
    r = cliente.post("/textos", data=dados_form, follow_redirects=True)
    assert "Textos salvos".encode() in r.data and "Órgão X".encode() in r.data  # header usa orgao_nome
    assert b'<div class="header-subtitle">SN</div>' in r.data  # e unidade_sigla no subtítulo
    doc = cliente.get("/termo/individual/ANA SILVA/documento").data.decode()
    assert "TESTE <b>Ana Silva</b>." in doc
    r = cliente.post("/textos", data={"restaurar": "individual_abertura"}, follow_redirects=True)
    assert "Padrão restaurado".encode() in r.data
    doc = cliente.get("/termo/individual/ANA SILVA/documento").data.decode()
    assert "Pelo presente termo" in doc and "Órgão X".encode() in cliente.get("/").data


def test_textos_marcador_invalido_nao_grava_nada(cliente):
    import textos
    dados_form = dict(textos.PADRAO, individual_abertura="Eu {nomee}", cidade="Goiânia (GO)")
    r = cliente.post("/textos", data=dados_form, follow_redirects=True)
    assert b"nomee" in r.data
    doc = cliente.get("/termo/devolucao/ANA SILVA/documento").data.decode()
    assert "Goiânia" not in doc and "Brasília (DF)" in doc


def test_exportar_e_importar_cadastros(cliente):
    r = cliente.get("/cadastros/exportar")
    assert r.status_code == 200 and r.headers["Content-Disposition"].endswith("cadastros.xlsx")
    r = cliente.post("/importar-cadastros", data={"arquivo": (io.BytesIO(r.data), "cadastros.xlsx")},
                     content_type="multipart/form-data", follow_redirects=True)
    assert b"1 centro" in r.data and b"1 pessoa" in r.data
    r = cliente.post("/importar-cadastros", data={"arquivo": (io.BytesIO(b"nada"), "x.xlsx")},
                     content_type="multipart/form-data", follow_redirects=True)
    assert "inválido".encode() in r.data


def test_importar_cadastros_reimporta_inventario(cliente):
    eid = _abrir(cliente)
    cliente.post(f"/inventario/{eid}/sala/01 - SALA CCI/ler", json={"numero": "1001"})
    r = cliente.get("/cadastros/exportar")
    r = cliente.post("/importar-cadastros", data={"arquivo": (io.BytesIO(r.data), "cadastros.xlsx")},
                     content_type="multipart/form-data", follow_redirects=True)
    assert "Inventário: 1 evento".encode() in r.data


def test_exportar_bens(cliente):
    r = cliente.get("/bens/exportar")
    assert r.status_code == 200 and r.headers["Content-Disposition"].endswith("bens.xlsx")
    assert b"Baixar" in cliente.get("/upload").data and b"Carregar" in cliente.get("/upload").data


def test_pesquisa_numero_vai_para_ficha_e_texto_lista(cliente):
    r = cliente.get("/pesquisa?q=1002")
    assert r.status_code == 302 and r.headers["Location"].endswith("/bem?numero=1002")
    r = cliente.get("/pesquisa?q=cci")
    assert r.status_code == 200
    assert b"JAQUELINE PORTELA" in r.data and b"CADEIRA" in r.data and b"/termo/ccusto/CCI" in r.data
    assert b"/pesquisa?ccusto=CCI" in r.data
    r = cliente.get("/pesquisa?q=ana")
    assert b"ANA SILVA" in r.data and b"/pesquisa?pessoa=ANA" in r.data
    r = cliente.get("/pesquisa?q=zzz")
    assert "Nada encontrado".encode() in r.data
    assert cliente.get("/pesquisa?q=9999").status_code == 200  # número inexistente cai na busca por texto


def test_pesquisa_filtra_bens_de_centro_e_pessoa(cliente):
    r = cliente.get("/pesquisa?ccusto=CCI")
    assert b"CADEIRA" in r.data and b"NOTEBOOK" not in r.data and b"/termo/ccusto/CCI" in r.data
    r = cliente.get("/pesquisa?pessoa=ANA SILVA")
    assert b"NOTEBOOK" in r.data and b"CADEIRA" not in r.data and b"/termo/individual/ANA" in r.data


def test_cabecalho_tem_campo_de_pesquisa(cliente):
    assert b'action="/pesquisa"' in cliente.get("/").data


def test_cadastro_de_processos_sei(cliente):
    r = cliente.get("/cadastros/processos")
    assert r.status_code == 200 and b"Processos SEI" in r.data
    r = cliente.post("/cadastros/processos/incluir", data={"tipo": "ccusto", "descricao": "Termos 2026", "numero_sei": "2222", "vigente": "1"}, follow_redirects=True)
    assert b"Termos 2026" in r.data and b"2222" in r.data and b"vigente" in r.data
    r = cliente.post("/cadastros/processos/incluir", data={"tipo": "ccusto", "descricao": "", "numero_sei": "1"}, follow_redirects=True)
    assert "Descrição".encode() in r.data
    import db
    pid = db.processos(db.conectar())[0]["id"]
    r = cliente.post("/cadastros/processos/encerrar", data={"id": pid}, follow_redirects=True)
    r = confirmar_revisao(cliente, r, "/cadastros/processos/encerrar")
    assert b"encerrado" in r.data
    r = cliente.post("/cadastros/processos/vigente", data={"id": pid}, follow_redirects=True)
    r = confirmar_revisao(cliente, r, "/cadastros/processos/vigente")
    assert b"vigente" in r.data
    r = cliente.post("/cadastros/processos/excluir", data={"id": pid}, follow_redirects=True)
    r = confirmar_revisao(cliente, r, "/cadastros/processos/excluir")
    assert "exclu".encode() in r.data and b"Termos 2026" not in r.data


def test_termos_emitidos_lista_detalhe_e_documento_sei(cliente):
    cliente.post("/cadastros/processos/incluir", data={"tipo": "ccusto", "descricao": "T", "numero_sei": "2222", "vigente": "1"})
    cliente.get("/termo/ccusto/CCI/docx")
    r = cliente.get("/termos-emitidos")
    assert r.status_code == 200 and b"CCI" in r.data and b"2222" in r.data
    assert b"CCI" not in cliente.get("/termos-emitidos?tipo=individual").data.split(b"<tbody>")[1]
    assert b"CCI" in cliente.get("/termos-emitidos?chave=cc").data
    import db
    tid = db.termos_emitidos(db.conectar())[0]["id"]
    r = cliente.get(f"/termos-emitidos/{tid}")
    assert b"CADEIRA" in r.data and b"1001" in r.data and b'name="documento_sei"' in r.data
    r = cliente.post(f"/termos-emitidos/{tid}/documento", data={"documento_sei": "0459999"}, follow_redirects=True)
    assert b"0459999" in r.data and b"informe documento e bloco" in r.data and b"mailto:" not in r.data
    # com documento + bloco e e-mail do responsável (CCI tem j@cfc.org.br): link mailto com assunto e corpo
    r = cliente.post(f"/termos-emitidos/{tid}/documento", data={"documento_sei": "0459999", "bloco_sei": "77"}, follow_redirects=True)
    html = r.data.decode()
    assert "mailto:j@cfc.org.br?subject=" in html and "bloco%20de%20assinatura%2077" in html
    assert "Prezado%28a%29%20Jaqueline%2C" in html and "Jaqueline%20Portela" not in html    # só o primeiro nome
    assert "E-mail não enviado" in html
    r = cliente.post(f"/termos-emitidos/{tid}/email", follow_redirects=True)
    assert "E-mail enviado em" in r.data.decode() and "Enviar email novamente" in r.data.decode()
    assert b"0459999" in cliente.get("/termos-emitidos").data and b"77" in cliente.get("/termos-emitidos").data
    assert cliente.get("/termos-emitidos/999").status_code == 404
    assert b"Termos emitidos" in cliente.get("/").data    # menu


def test_painel_na_tela_inicial(cliente):
    r = cliente.get("/")
    assert r.status_code == 200
    assert b"dsgov-atalhos" in r.data and "Realizar inventário".encode() in r.data and b'href="/inventario"' in r.data   # carrossel de atalhos
    assert b"bens ativos" in r.data
    assert b"Nenhuma" in r.data       # última importação: nenhuma
    assert b"sem centro nem pessoa" in r.data
    assert "Sobre o patrimônio".encode() in r.data and b'href="/ajuda"' in r.data and b"Ver guia" in r.data


def test_atalhos_do_inicio_seguem_as_permissoes(cliente, usuarios_exemplo):
    """Admin recebe os seis atalhos; Operador e Consulta não recebem 'Realizar inventário' (exige
    inventario.ler); Consulta não recebe o indicador de importação (exige importacao_tela)."""
    r = cliente.get("/")
    for rotulo in ("Termo por centro de custo", "Termo individual", "Termo de devolução",
                   "Termos emitidos", "Realizar inventário", "Análise"):
        assert rotulo.encode() in r.data
    assert r.data.count(b'class="br-card dsgov-atalho"') == 6

    cliente.post("/sair"); logar(cliente, *usuarios_exemplo["operador"])
    r = cliente.get("/")
    assert r.data.count(b'class="br-card dsgov-atalho"') == 5
    assert "Realizar inventário".encode() not in r.data
    assert "Nenhuma importação registrada".encode() in r.data   # operador continua com o KPI de importação

    cliente.post("/sair"); logar(cliente, *usuarios_exemplo["consulta"])
    r = cliente.get("/")
    assert r.data.count(b'class="br-card dsgov-atalho"') == 5
    assert "Realizar inventário".encode() not in r.data
    assert "Nenhuma importação registrada".encode() not in r.data   # consulta não recebe o KPI de importação


def test_inventariante_sozinho_e_redirecionada_do_inicio(cliente):
    cliente.post("/sair"); logar(cliente, "beltrana", SENHA_PADRAO)
    assert cliente.get("/").headers["Location"] == "/inventario"


def test_listas_mostram_situacao_do_termo(cliente):
    r = cliente.get("/centro-custos")
    assert b"sem termo" in r.data and b"/termo/ccusto/CCI" in r.data
    r = cliente.get("/termos-individuais")
    assert b"ANA SILVA" in r.data and b"sem termo" in r.data and b"/termo/individual/ANA" in r.data
    cliente.post("/cadastros/processos/incluir", data={"tipo": "ccusto", "descricao": "T", "numero_sei": "2", "vigente": "1"})
    cliente.get("/termo/ccusto/CCI/docx")
    r = cliente.get("/centro-custos")
    assert b"vigente" in r.data and b"sem doc./bloco" in r.data
    import db
    tid = db.termos_emitidos(db.conectar())[0]["id"]
    cliente.post(f"/termos-emitidos/{tid}/documento", data={"documento_sei": "1", "bloco_sei": "2"})
    assert "não enviado".encode() in cliente.get("/centro-custos").data
    cliente.post(f"/termos-emitidos/{tid}/email")
    r = cliente.get("/centro-custos")
    assert b"enviado " in r.data and f"/termos-emitidos/{tid}".encode() in r.data


def test_recorte_tela_filtros_termo_e_xlsx(cliente):
    r = cliente.get("/analise")
    assert r.status_code == 200 and b"Bens ATIVO" in r.data and b"1001" in r.data and b"1003" not in r.data
    assert b'data-grafico="g-situacao"' not in r.data and b'data-grafico="g-centro"' in r.data
    r = cliente.get("/analise?situacao=ATIVO&ccusto=CCI")
    assert b"centro de custo CCI" in r.data and b"/termo/ccusto/CCI" in r.data and b"sem termo" in r.data
    assert b'data-grafico="g-centro"' not in r.data and b"/analise?situacao=ATIVO&amp;ccusto=CCI&amp;faixa=" in r.data
    r = cliente.get("/analise?pessoa=ANA SILVA")
    assert b"/termo/individual/ANA" in r.data and b"NOTEBOOK" in r.data and b"CADEIRA" not in r.data
    r = cliente.get("/analise?situacao=&valor_de=1.000,00&valor_ate=2000")
    assert b"NOTEBOOK" in r.data and b"CADEIRA" not in r.data and b"todas as situa" in r.data
    r = cliente.get("/analise?situacao=&valor_de=1.000&valor_ate=1.600")     # ponto de milhar
    assert b"NOTEBOOK" in r.data and b"CADEIRA" not in r.data
    r = cliente.get("/analise?situacao=&valor_de=1000.5&valor_ate=1600")     # ponto decimal
    assert b"NOTEBOOK" in r.data
    r = cliente.get("/analise?situacao=&valor_de=1.000.000")
    assert b"NOTEBOOK" not in r.data and b"Nenhum bem" in r.data
    r = cliente.get("/analise?valor_de=abc", follow_redirects=True)
    assert "Valor inválido".encode() in r.data
    r = cliente.get("/analise/xlsx?ccusto=CCI")
    assert r.status_code == 200 and r.headers["Content-Disposition"].endswith("analise.xlsx")
    assert b"An\xc3\xa1lise" in cliente.get("/").data       # menu
    r = cliente.get("/analise?situacao=")
    assert b"ccusto=CCI&amp;situacao=" in r.data or b"ccusto=CCI&situacao=" in r.data   # drill-down mantém "todas"
    assert b'href="/analise/xlsx?situacao="' in r.data
    ws = load_workbook(io.BytesIO(cliente.get("/analise/xlsx?situacao=").data)).active
    assert ws.max_row - 1 == 4     # 4 bens da semente (inclui 1003 BAIXADO); a tela mostra o mesmo total


@pytest.mark.parametrize("sufixo", ["", "/xlsx"])
def test_alias_preserva_query(cliente, sufixo):
    query = "situacao=&ccusto=GEX%2BLIC&ccusto=CCI&valor_status=zero"
    r = cliente.get("/recorte" + sufixo + "?" + query)
    assert r.status_code == 301
    assert r.headers["Location"] == "/analise" + sufixo + "?" + query


def _ids(*nomes):
    """IDs das contas com esses nomes: o formulário web da comissão manda IDs de usuário, nunca nomes."""
    import db
    conn = db.conectar()
    return [conn.execute("SELECT id FROM usuarios WHERE nome = ?", (n,)).fetchone()[0] for n in nomes]


def inventario_do_teste(eid):
    """Evento recarregado do banco (comissão exibida, por nome)."""
    import db, inventario
    return inventario.evento(db.conectar(), eid)


def _abrir(cliente, comissao=("Fulano", "Beltrana")):
    """Abre evento com a comissão dada (nomes de usuários existentes). O admin logado é 'Fulano'."""
    cliente.post("/inventario/abrir", data={"nome": "Inv", "usuarios": _ids(*comissao), "escopo": "todas"})
    import db, inventario
    return inventario.evento_aberto(db.conectar())["id"]


def test_inventario_eventos_abrir_e_encerrar(cliente):
    r = cliente.get("/inventario")
    assert r.status_code == 200 and b"Abrir evento" in r.data and b"Nenhum evento aberto" in r.data
    r = cliente.post("/inventario/abrir", data={"nome": "Inventário 2026", "descricao": "Portaria 1", "usuarios": _ids("Fulano", "Beltrana"), "escopo": "todas"}, follow_redirects=True)
    assert "Inventário 2026".encode() in r.data and b"01 - SALA CCI" in r.data and b"99 - SEM MAPA" in r.data
    assert "Inventário".encode() in cliente.get("/").data                          # menu
    r = cliente.post("/inventario/abrir", data={"nome": "Outro", "usuarios": _ids("Fulano"), "escopo": "todas"}, follow_redirects=True)
    assert b"Outro" in r.data                                          # chave única: abrir fecha o anterior, não recusa
    import db, inventario
    eid = inventario.evento_aberto(db.conectar())["id"]
    assert inventario.evento(db.conectar(), eid)["nome"] == "Outro"
    assert b"Comiss" in cliente.get(f"/inventario/{eid}").data
    r = cliente.post(f"/inventario/{eid}/encerrar", data={}, follow_redirects=True)
    assert b"Confirmar encerramento" in r.data
    r = cliente.post(f"/inventario/{eid}/encerrar", data={"confirmar": "1"}, follow_redirects=True)
    assert b"encerrado" in r.data
    assert cliente.get("/inventario/999").status_code == 404


def test_inventario_abrir_com_amostragem(cliente):
    r = cliente.post("/inventario/abrir", data={"nome": "Amostra", "usuarios": _ids("Beltrana"), "escopo": "escolher", "salas": ["99 - SEM MAPA"]}, follow_redirects=True)
    assert b"99 - SEM MAPA" in r.data and b"01 - SALA CCI" not in r.data.split(b"<tbody>")[1]


def test_inventario_sala_leitura_json(cliente):
    eid = _abrir(cliente, comissao=["Beltrana"])
    r = cliente.get(f"/inventario/{eid}/sala/01 - SALA CCI")
    assert r.status_code == 200 and b'id="leitura"' in r.data and b"html5-qrcode" in r.data and "não faz parte da comissão".encode() in r.data
    r = cliente.post(f"/inventario/{eid}/sala/01 - SALA CCI/ler", json={"numero": "1001"})
    assert r.status_code == 403 and r.get_json()["erro"] == NEGADO           # sem vínculo: barrado antes da view
    cliente.post(f"/inventario/{eid}/comissao", data={"usuarios": _ids("Fulano", "Beltrana")})
    j = cliente.post(f"/inventario/{eid}/sala/01 - SALA CCI/ler", json={"numero": "001001"}).get_json()
    assert j["situacao"] == "localizado" and j["numero"] == 1001 and j["descricao"] == "CADEIRA" and j["reler"] is False
    j = cliente.post(f"/inventario/{eid}/sala/01 - SALA CCI/ler", json={"numero": "1004"}).get_json()
    assert j["situacao"] == "divergente" and j["cadastrado_em"] == "99 - SEM MAPA"
    r = cliente.post(f"/inventario/{eid}/sala/01 - SALA CCI/ler", json={"numero": "99999"})
    assert r.status_code == 404 and r.get_json()["numero"] == 99999
    r = cliente.post(f"/inventario/{eid}/sala/01 - SALA CCI/ler", json={"numero": "abc"})
    assert r.status_code == 404
    for numero in ("ABC1001", "S/N-1001", "12.345", "1001x"):
        r = cliente.post(f"/inventario/{eid}/sala/01 - SALA CCI/ler", json={"numero": numero})
        assert r.status_code == 404
    j = cliente.post(f"/inventario/{eid}/sala/01 - SALA CCI/ler", json={"numero": "1001"}).get_json()
    assert j["reler"] and j["leitura_anterior"]["integrante"] == "Fulano"
    r = cliente.get(f"/inventario/{eid}/sala/01 - SALA CCI")
    assert b"Localizado" in r.data and b"Divergente" in r.data and b"1004" in r.data
    assert r.data.index(b'data-numero="1002"') < r.data.index(b'data-numero="1001"')   # pendente antes de localizado
    r = cliente.post(f"/inventario/{eid}/leitura/1001", json={"conservacao": "Ruim", "quem_usa": "Ciclana"})
    assert r.status_code == 200 and r.get_json()["ok"]
    assert cliente.post(f"/inventario/{eid}/leitura/1001", json={"conservacao": "Péssimo"}).status_code == 409
    assert cliente.post(f"/inventario/{eid}/leitura/1002", json={"observacao": "x"}).status_code == 409   # não lido


def test_inventario_fotos_e_sobras(cliente, monkeypatch, tmp_path):
    import io
    from PIL import Image
    import fotos
    eid = _abrir(cliente)
    buf = io.BytesIO(); Image.new("RGB", (30, 20), (1, 2, 3)).save(buf, "PNG"); imagem = buf.getvalue()
    # fotos desativadas: sobra sem foto é aceita; foto de bem recusada
    for v in fotos.VARIAVEIS:
        monkeypatch.delenv(v, raising=False)
    r = cliente.get(f"/inventario/{eid}/sala/01 - SALA CCI")
    assert b"Fotos desativadas" in r.data
    r = cliente.post(f"/inventario/{eid}/sala/01 - SALA CCI/sobra", data={"descricao": "VENTILADOR", "observacao": "sem plaqueta"}, follow_redirects=True)
    assert b"VENTILADOR" in r.data
    cliente.post(f"/inventario/{eid}/sala/01 - SALA CCI/ler", json={"numero": "1001"})
    r = cliente.post(f"/inventario/{eid}/leitura/1001/foto", data={"foto": (io.BytesIO(imagem), "a.png")}, content_type="multipart/form-data")
    assert r.status_code == 409 and "desativadas" in r.get_json()["erro"]
    # fotos ativas: cliente falso
    for v in fotos.VARIAVEIS:
        monkeypatch.setenv(v, "x")
    monkeypatch.setenv("R2_PUBLIC_URL", "https://f.exemplo.org")
    enviados = []
    monkeypatch.setattr(fotos, "enviar", lambda chave, dados: enviados.append(chave) or f"https://f.exemplo.org/{chave}")
    apagados = []
    monkeypatch.setattr(fotos, "apagar", lambda url: apagados.append(url))
    r = cliente.post(f"/inventario/{eid}/leitura/1001/foto", data={"foto": (io.BytesIO(imagem), "a.png")}, content_type="multipart/form-data")
    assert r.status_code == 200 and [f["nfoto"] for f in r.get_json()["fotos"]] == [1] and enviados == ["inv/1-1001.webp"]
    r = cliente.post(f"/inventario/{eid}/leitura/1001/foto", data={"foto": (io.BytesIO(imagem), "a.png")}, content_type="multipart/form-data")
    assert [f["url"] for f in r.get_json()["fotos"]] == ["https://f.exemplo.org/inv/1-1001.webp", "https://f.exemplo.org/inv/2-1001.webp"]
    r = cliente.post(f"/inventario/{eid}/leitura/1002/foto", data={"foto": (io.BytesIO(imagem), "a.png")}, content_type="multipart/form-data")
    assert r.status_code == 409 and "Leia o bem" in r.get_json()["erro"]                  # 1002 não foi lido
    r = cliente.post(f"/inventario/{eid}/leitura/1001/foto/1/excluir", data={"volta": "01 - SALA CCI"}, follow_redirects=True)
    assert apagados == ["https://f.exemplo.org/inv/1-1001.webp"] and b"Foto removida" in r.data
    assert cliente.post(f"/inventario/{eid}/leitura/1001/foto/1/excluir", follow_redirects=True).status_code == 404
    r = cliente.post(f"/inventario/{eid}/sala/01 - SALA CCI/sobra", data={"descricao": "CADEIRA VELHA", "observacao": "x"}, follow_redirects=True)
    assert "precisa de foto".encode() in r.data
    r = cliente.post(f"/inventario/{eid}/sala/01 - SALA CCI/sobra", data={"descricao": "CADEIRA VELHA", "observacao": "x", "foto": (io.BytesIO(imagem), "b.jpg")}, content_type="multipart/form-data", follow_redirects=True)
    assert b"CADEIRA VELHA" in r.data and enviados[-1].startswith("inv/sobra-") and enviados[-1].endswith(".webp")
    # falha no envio: sobra não fica registrada; foto de bem não fica registrada
    monkeypatch.setattr(fotos, "enviar", lambda chave, dados: (_ for _ in ()).throw(RuntimeError("bucket fora")))
    r = cliente.post(f"/inventario/{eid}/sala/01 - SALA CCI/sobra", data={"descricao": "MESA VELHA", "observacao": "x", "foto": (io.BytesIO(imagem), "c.jpg")}, content_type="multipart/form-data", follow_redirects=True)
    assert b"MESA VELHA" not in r.data and "não registrada".encode() in r.data
    r = cliente.post(f"/inventario/{eid}/leitura/1001/foto", data={"foto": (io.BytesIO(imagem), "a.png")}, content_type="multipart/form-data")
    assert r.status_code == 409 and "Falha ao enviar" in r.get_json()["erro"]
    import db, inventario
    assert [f["nfoto"] for f in inventario.fotos_do_bem_no_evento(db.conectar(), eid, 1001)] == [2]
    sobras = inventario.bens_da_sala(db.conectar(), eid, "01 - SALA CCI")["sobras"]
    assert [s["descricao"] for s in sobras] == ["CADEIRA VELHA", "VENTILADOR"]
    r = cliente.post(f"/inventario/{eid}/sobra/{sobras[0]['id']}/excluir", follow_redirects=True)
    assert b"CADEIRA VELHA" not in r.data and len(apagados) == 2
    # evento encerrado: não mexe em foto no bucket nem aceita novo envio
    monkeypatch.setattr(fotos, "enviar", lambda chave, dados: enviados.append(chave) or f"https://f.exemplo.org/{chave}")
    cliente.post(f"/inventario/{eid}/encerrar", data={"confirmar": "1"})
    n_apagados, n_enviados = len(apagados), len(enviados)
    r = cliente.post(f"/inventario/{eid}/leitura/1001/foto/2/excluir", follow_redirects=True)
    assert len(apagados) == n_apagados and b"Evento encerrado" in r.data
    r = cliente.post(f"/inventario/{eid}/leitura/1001/foto", data={"foto": (io.BytesIO(imagem), "d.png")}, content_type="multipart/form-data")
    assert r.status_code == 409 and len(enviados) == n_enviados


def test_sala_mostra_varias_fotos_e_camera(cliente, monkeypatch):
    import fotos
    eid = _abrir(cliente)
    for v in fotos.VARIAVEIS:
        monkeypatch.setenv(v, "x")
    cliente.post(f"/inventario/{eid}/sala/01 - SALA CCI/ler", json={"numero": "1001"})
    import db, inventario
    conn = db.conectar()
    inventario.adicionar_foto(conn, eid, 1001, lambda c: "https://x/1.webp")
    inventario.adicionar_foto(conn, eid, 1001, lambda c: "https://x/2.webp")
    r = cliente.get(f"/inventario/{eid}/sala/01 - SALA CCI")
    linha = r.data.split(b'data-numero="1001"')[1].split(b"</tr>")[0]
    assert linha.count(b'class="dsgov-miniatura"') == 1 and b'src="https://x/1.webp"' in linha and b'src="https://x/2.webp"' not in linha   # só a primeira; as demais no cadastro do bem
    assert b'<span class="br-tag small mr-1" title="Todas as fotos no cadastro do bem">+1</span>' in linha
    assert f"/inventario/{eid}/leitura/1001/foto/1/excluir".encode() in linha and b"/foto/2/excluir" not in linha
    assert b'class="foto-input"' in linha and b"foto-input\" hidden disabled" not in linha         # câmera continua, habilitada
    linha2 = r.data.split(b'data-numero="1002"')[1].split(b"</tr>")[0]
    assert b"dsgov-miniatura" not in linha2 and b'class="foto-input" hidden disabled' in linha2      # não lido: câmera desabilitada
    assert b"/foto/0/excluir" in r.data                                                              # molde da URL para o JS
    assert b"as fotos s" in r.data                                                                   # confirm do Desmarcar fala em fotos
    cliente.post(f"/inventario/{eid}/encerrar", data={"confirmar": "1"})
    r = cliente.get(f"/inventario/{eid}/sala/01 - SALA CCI")
    linha = r.data.split(b'data-numero="1001"')[1].split(b"</tr>")[0]
    assert linha.count(b'class="dsgov-miniatura"') == 1 and b">+1<" in linha and b"/excluir" not in linha and b"foto-input" not in linha   # (o JS da página ainda cita foto-input; por isso a checagem é só na linha)


def test_foto_com_url_nao_http_nao_vira_link(cliente):
    eid = _abrir(cliente)
    cliente.post(f"/inventario/{eid}/sala/01 - SALA CCI/ler", json={"numero": "1001"})
    import db, inventario
    conn = db.conectar()
    inventario.adicionar_foto(conn, eid, 1001, lambda c: "javascript:alert(1)")
    r = cliente.get(f"/inventario/{eid}/sala/01 - SALA CCI")
    assert b'href="javascript:alert' not in r.data
    r = cliente.get("/bem?numero=1001")
    assert b'href="javascript:alert' not in r.data


def test_inventario_relatorio_xlsx_e_card_do_painel(cliente):
    assert b"Nenhum invent" in cliente.get("/").data
    eid = _abrir(cliente)
    cliente.post(f"/inventario/{eid}/sala/01 - SALA CCI/ler", json={"numero": "1001"})
    r = cliente.get("/")
    assert b"Invent\xc3\xa1rio em andamento" in r.data and f"/inventario/{eid}".encode() in r.data
    r = cliente.get(f"/inventario/{eid}/relatorio")
    assert r.status_code == 200 and b"1001" in r.data and b"Localizado" in r.data and b"1002" in r.data
    assert b"1002" not in cliente.get(f"/inventario/{eid}/relatorio?situacao=localizado").data.split(b"<tbody>")[1]
    assert b"1004" not in cliente.get(f"/inventario/{eid}/relatorio?localizacao=01 - SALA CCI").data.split(b"<tbody>")[1]
    r = cliente.get(f"/inventario/{eid}/xlsx?localizacao=01 - SALA CCI")
    assert r.status_code == 200 and r.headers["Content-Disposition"].endswith(".xlsx")


def test_inventario_relatorio_filtros_ordem_modal_e_xlsx_com_fotos(cliente):
    import io
    from openpyxl import load_workbook
    eid = _abrir(cliente)
    cliente.post(f"/inventario/{eid}/sala/01 - SALA CCI/ler", json={"numero": "1001"})
    cliente.post(f"/inventario/{eid}/leitura/1001", json={"quem_usa": "José"})
    import db, inventario
    inventario.adicionar_foto(db.conectar(), eid, 1001, lambda chave: "https://x/1001.webp")
    inventario.adicionar_foto(db.conectar(), eid, 1001, lambda chave: "https://x/1001b.webp")
    r = cliente.get(f"/inventario/{eid}/relatorio?integrante=Fulano&busca=jose&ordem=numero&dir=desc")
    corpo = r.data.split(b"<tbody>")[1]
    assert r.status_code == 200 and b">1001<" in corpo and b">1002<" not in corpo
    assert "Integrante Fulano".encode() in r.data and b"Busca &#34;jose&#34;" in r.data and b"fa-sort-down" in r.data
    assert b"ordem=numero&amp;dir=asc" in r.data or b"dir=asc&amp;ordem=numero" in r.data      # clique de novo inverte
    assert b'data-foto="https://x/1001.webp"' in r.data and b'id="scrim-foto"' in r.data and b'name="fotos"' in r.data
    assert b'<span class="br-tag small">+1</span>' in r.data and b'data-foto="https://x/1001b.webp"' not in r.data
    assert b'name="ordem" value="numero"' in r.data and b"1 linha(s)" in r.data
    r = cliente.get(f"/inventario/{eid}/relatorio?foto=sem&conservacao=-")
    assert b">1001<" not in r.data.split(b"<tbody>")[1] and b"Sem foto" in r.data and "Não informada".encode() in r.data
    r_xss = cliente.get(f"/inventario/{eid}/relatorio?busca=<b>x</b>")
    assert b"<b>x</b>" not in r_xss.data
    r = cliente.get(f"/inventario/{eid}/xlsx?integrante=Fulano&fotos=1")
    assert r.status_code == 200 and r.headers["Content-Disposition"].endswith(".xlsx")
    ws = load_workbook(io.BytesIO(r.data))["Bens"]
    assert ws.cell(row=3, column=1).value == "Todas as salas · Integrante Fulano" and ws.max_row == 6
    assert ws.cell(row=6, column=13).value == '=_xlfn.IMAGE("https://x/1001.webp")'
    ws = load_workbook(io.BytesIO(cliente.get(f"/inventario/{eid}/xlsx").data))["Bens"]
    assert ws.cell(row=6, column=13).value == "https://x/1001.webp" and ws.max_row == 8      # 5 de cabeçalho + 1001, 1002, 1004


def test_inventario_painel(cliente):
    eid = _abrir(cliente)
    cliente.post(f"/inventario/{eid}/sala/01 - SALA CCI/ler", json={"numero": "1001"})
    r = cliente.get(f"/inventario/{eid}/painel")
    assert r.status_code == 200 and b'id="g-situacao"' in r.data and b'id="g-andares"' in r.data and b'id="g-salas"' not in r.data
    assert b"bens no escopo" in r.data and b"echarts-dsgov.js" in r.data and f"/inventario/{eid}/painel?andar=01".encode() in r.data
    assert b"1 (33.3%)" in r.data                                                       # localizados com % (1 de 3 bens ativos no escopo)
    r = cliente.get(f"/inventario/{eid}/painel?andar=01")
    assert b'id="g-salas"' in r.data and b"Salas do andar 01" in r.data and b"todos os andares" in r.data
    assert cliente.get("/inventario/999/painel").status_code == 404
    assert f"/inventario/{eid}/painel".encode() in cliente.get(f"/inventario/{eid}").data    # botão Painel
    cliente.post(f"/inventario/{eid}/encerrar", data={"confirmar": "1"})
    assert b"Evento encerrado" in cliente.get(f"/inventario/{eid}/painel").data


def test_inventario_lote_marcar_e_desmarcar(cliente, monkeypatch):
    import fotos
    eid = _abrir(cliente)
    r = cliente.get(f"/inventario/{eid}/sala/01 - SALA CCI")
    assert b'id="form-lote"' in r.data and b'name="numeros"' in r.data and b'value="marcar"' in r.data and b'value="desmarcar"' in r.data
    r = cliente.post(f"/inventario/{eid}/sala/01 - SALA CCI/lote", data={"acao": "marcar", "numeros": ["1001", "1002", "99999"]}, follow_redirects=True)
    assert b"2 bem(ns) marcado(s)" in r.data and b"99999" in r.data and r.data.count(b">Localizado<") == 2
    r = cliente.post(f"/inventario/{eid}/sala/01 - SALA CCI/lote", data={"acao": "marcar"}, follow_redirects=True)
    assert b"Selecione ao menos um bem" in r.data
    import db, inventario
    inventario.adicionar_foto(db.conectar(), eid, 1001, lambda chave: "https://x/1001.webp")
    apagadas = []
    monkeypatch.setattr(fotos, "apagar", apagadas.append)
    r = cliente.post(f"/inventario/{eid}/sala/01 - SALA CCI/lote", data={"acao": "desmarcar", "numeros": ["1001"]}, follow_redirects=True)
    assert b"desfeita" in r.data and apagadas == ["https://x/1001.webp"] and r.data.count(b">Localizado<") == 1
    r = cliente.post(f"/inventario/{eid}/sala/01 - SALA CCI/lote", data={"acao": "desmarcar", "numeros": ["1001"]}, follow_redirects=True)
    assert b"Nenhuma leitura para desfazer" in r.data
    cliente.post(f"/inventario/{eid}/comissao", data={"usuarios": _ids("Beltrana")})
    r = cliente.post(f"/inventario/{eid}/sala/01 - SALA CCI/lote", data={"acao": "marcar", "numeros": ["1002"]}, follow_redirects=True)
    assert r.status_code == 403 and NEGADO.encode() in r.data
    cliente.post(f"/inventario/{eid}/comissao", data={"usuarios": _ids("Fulano", "Beltrana")})
    cliente.post(f"/inventario/{eid}/encerrar", data={"confirmar": "1"})
    r = cliente.post(f"/inventario/{eid}/sala/01 - SALA CCI/lote", data={"acao": "marcar", "numeros": ["1002"]}, follow_redirects=True)
    assert b"Evento encerrado" in r.data
    assert b'id="form-lote"' not in cliente.get(f"/inventario/{eid}/sala/01 - SALA CCI").data


def test_pessoa_com_email_e_matricula(cliente):
    r = cliente.post("/cadastros/pessoas/incluir", data={"nome": "bruno lima", "email": "nao-e-email", "matricula": "1"})
    assert b"Informe um e-mail v" in r.data
    r = cliente.post("/cadastros/pessoas/incluir", data={"nome": "bruno lima", "email": "b@cfc.org.br", "matricula": "0012"},
                     follow_redirects=True)
    assert b"BRUNO LIMA" in r.data and b"b@cfc.org.br" in r.data      # tela da pessoa mostra o e-mail
    assert b"b@cfc.org.br" in cliente.get("/cadastros/pessoas").data   # lista também
    r = cliente.get("/cadastros/pessoas/BRUNO LIMA/editar")
    assert b'value="b@cfc.org.br"' in r.data and b'value="0012"' in r.data
    r = cliente.post("/cadastros/pessoas/BRUNO LIMA/editar", data={"nome": "bruno lima", "email": "", "matricula": "0012"},
                     follow_redirects=True)
    import db
    assert db.pessoa(db.conectar(), "BRUNO LIMA")["email"] is None
    assert b"tratamento" not in cliente.get("/cadastros/responsaveis/CCI/editar").data


def test_termo_individual_sem_email_avisa(cliente):
    cliente.post("/cadastros/processos/incluir", data={"tipo": "individual", "descricao": "T", "numero_sei": "3333", "vigente": "1"})
    cliente.get("/termo/individual/ANA SILVA/docx")
    import db
    tid = db.termos_emitidos(db.conectar(), tipo="individual")[0]["id"]
    cliente.post(f"/termos-emitidos/{tid}/documento", data={"documento_sei": "1", "bloco_sei": "2"})
    r = cliente.get(f"/termos-emitidos/{tid}")
    assert b"Sem e-mail cadastrado para ANA SILVA" in r.data and b"mailto:" not in r.data
    cliente.post("/cadastros/pessoas/ANA SILVA/editar", data={"nome": "ana silva", "email": "a@cfc.org.br"})
    r = cliente.get(f"/termos-emitidos/{tid}")
    assert b"mailto:a@cfc.org.br" in r.data and b"Prezado%28a%29%20Ana%2C" in r.data


def test_devolucao_situacao_por_pessoa_com_sinal_de_email(cliente):
    r = cliente.get("/termo_devolucao")
    assert "Situação dos termos de devolução".encode() in r.data and b"ANA SILVA" in r.data and b"Bens com ela" in r.data
    cliente.post("/cadastros/processos/incluir", data={"tipo": "devolucao", "descricao": "D", "numero_sei": "4444", "vigente": "1"})
    cliente.post("/termo_devolucao", data={"nome": "ANA SILVA", "numero_bem": "1002"})
    cliente.get("/termo/devolucao/ANA SILVA/docx")
    r = cliente.get("/termo_devolucao")
    assert b"sem doc./bloco" in r.data
    import db
    tid = db.termos_emitidos(db.conectar(), tipo="devolucao")[0]["id"]
    cliente.post(f"/termos-emitidos/{tid}/documento", data={"documento_sei": "1", "bloco_sei": "2"})
    cliente.post(f"/termos-emitidos/{tid}/email")
    r = cliente.get("/termo_devolucao")
    assert b"enviado " in r.data and f"/termos-emitidos/{tid}".encode() in r.data


def test_devolucao_sugere_bens_da_pessoa_sem_travar(cliente):
    # só escolher a pessoa (sem número) mostra os bens que estão com ela, sem erro
    r = cliente.post("/termo_devolucao", data={"nome": "ANA SILVA"}, follow_redirects=True)
    assert b"Bens com ANA SILVA" in r.data and b"1002" in r.data and b"Adicionar todos" in r.data and b"Erro." not in r.data
    # um bem que NÃO está com ela também pode entrar
    r = cliente.post("/termo_devolucao", data={"nome": "ANA SILVA", "numero_bem": "1001"}, follow_redirects=True)
    assert b"Bens a devolver" in r.data and b"1001" in r.data
    # adicionar todos: 1002 sai da sugestão e entra na lista
    r = cliente.post("/termo_devolucao", data={"nome": "ANA SILVA", "todos": "1"}, follow_redirects=True)
    assert b"Bens com ANA SILVA" not in r.data and r.data.count(b">1002<") == 1 and b">1001<" in r.data
    # sem pessoa e sem número: pede a pessoa
    cliente.post("/termo_devolucao", data={"limpar": "1"})
    with cliente.session_transaction() as s:
        s.pop("nome_devolucao", None)
    r = cliente.post("/termo_devolucao", data={"nome": ""}, follow_redirects=True)
    assert b"Escolha a pessoa que devolve" in r.data


def test_docx_tipo_invalido_da_404_e_head_nao_registra(cliente):
    assert cliente.get("/termo/xyz/CCI/docx").status_code == 404
    assert cliente.get("/termo/xyz/CCI").status_code == 404
    assert cliente.post("/termo/xyz/CCI/registrar").status_code == 404
    db.incluir_processo(db.conectar(), "ccusto", "Termos", "1111")
    assert cliente.head("/termo/ccusto/CCI/docx").status_code == 200
    assert db.termos_emitidos(db.conectar()) == []
    assert cliente.get("/termo/ccusto/CCI/docx").status_code == 200
    assert len(db.termos_emitidos(db.conectar())) == 1


def test_leitura_com_json_que_nao_e_objeto_da_400(cliente):
    eid = _abrir(cliente)
    r = cliente.post(f"/inventario/{eid}/sala/01 - SALA CCI/ler", json=[1, 2])
    assert r.status_code == 400 and "objeto JSON" in r.get_json()["erro"]
    r = cliente.post(f"/inventario/{eid}/leitura/1001", json="texto")
    assert r.status_code == 400 and "objeto JSON" in r.get_json()["erro"]


def test_upload_acima_de_20mb_da_mensagem_e_nao_500(cliente):
    grande = io.BytesIO(b"x" * (20 * 1024 * 1024 + 1))
    r = cliente.post("/upload", data={"arquivo": (grande, "export.xlsx")}, content_type="multipart/form-data",
                     headers={"Referer": "http://localhost/upload"}, follow_redirects=True)
    assert r.status_code == 200 and "Arquivo muito grande".encode() in r.data


def test_form_sobra_nao_e_o_proprio_br_card(cliente):
    eid = _abrir(cliente)
    html = cliente.get(f"/inventario/{eid}/sala/01 - SALA CCI").get_data(as_text=True)
    assert 'id="card-sobra"' in html and 'id="form-sobra"' in html
    assert 'class="br-card mt-3" id="form-sobra"' not in html   # BRCard trocaria o id do form


def _grupo(menu, rotulo):
    """Isola o <div class="menu-folder">...</div> cujo título (primeiro <span class="content">) é `rotulo`.
    Sem divs aninhadas dentro de um grupo: o primeiro `</div>` depois do título fecha o próprio grupo."""
    for parte in menu.split(b'<div class="menu-folder')[1:]:
        titulo = parte.split(b'<span class="content">', 2)[1].split(b'</span>')[0]
        if titulo == rotulo.encode():
            fim = parte.index(b'</div>')
            return b'<div class="menu-folder' + parte[:fim + len(b'</div>')]
    raise AssertionError(f"grupo {rotulo!r} não encontrado no menu")


def test_menu_inventario_e_grupo_com_telas_do_evento_aberto(cliente):
    r = cliente.get("/").data
    menu = r.split(b'id="main-navigation"')[1].split(b"menu-footer")[0]
    # Termos de Responsabilidade, Inventário e Cadastros viraram grupos na árvore por função (5C)
    assert menu.count(b"menu-folder") == 3
    inv = _grupo(menu, "Inventário")
    assert b">Eventos<" in inv and b"/painel" not in inv
    assert b'href="/inventario"' in inv
    # a Home não pertence ao grupo Inventário: pasta fechada, sem a classe active
    assert b'<div class="menu-folder">' in menu and b'aria-expanded="false"' in inv
    assert b'href="javascript:void(0)" role="treeitem" aria-expanded=' in inv   # título é link: pasta fecha/abre ao clicar (drop-menu do DSGov)
    eid = _abrir(cliente)
    menu = cliente.get("/").data.split(b'id="main-navigation"')[1].split(b"menu-footer")[0]
    inv = _grupo(menu, "Inventário")
    assert f'href="/inventario/{eid}"'.encode() in inv and f'href="/inventario/{eid}/painel"'.encode() in inv and f'href="/inventario/{eid}/relatorio"'.encode() in inv
    assert b">Inv<" in inv and b">Painel<" in inv and b">Relat" in inv
    cliente.post(f"/inventario/{eid}/encerrar", data={"confirmar": "1"})
    menu = cliente.get("/").data.split(b'id="main-navigation"')[1].split(b"menu-footer")[0]
    assert b"/painel" not in _grupo(menu, "Inventário")


def test_falha_no_bucket_mantem_a_foto_e_avisa(cliente, monkeypatch):
    import fotos
    eid = _abrir(cliente)
    cliente.post(f"/inventario/{eid}/sala/01 - SALA CCI/ler", json={"numero": "1001"})
    import db, inventario
    inventario.adicionar_foto(db.conectar(), eid, 1001, lambda c: "https://x/1.webp")
    monkeypatch.setattr(fotos, "apagar", lambda url: (_ for _ in ()).throw(RuntimeError("bucket fora")))
    origem = {"Referer": f"/inventario/{eid}/sala/01 - SALA CCI"}          # o handler global volta para a página de origem
    r = cliente.post(f"/inventario/{eid}/leitura/1001/foto/1/excluir", data={"volta": "01 - SALA CCI"}, headers=origem, follow_redirects=True)
    assert "Não foi possível apagar a foto no bucket".encode() in r.data and b'src="https://x/1.webp"' in r.data
    r = cliente.post(f"/inventario/{eid}/sala/01 - SALA CCI/lote", data={"acao": "desmarcar", "numeros": ["1001"]}, headers=origem, follow_redirects=True)
    assert "Não foi possível apagar".encode() in r.data and r.data.count(b">Localizado<") == 1
    assert [f["nfoto"] for f in inventario.fotos_do_bem_no_evento(db.conectar(), eid, 1001)] == [1]


def test_inventario_comissao_por_usuarios(cliente, dados, usuarios_exemplo):
    fulano, beltrana = _ids("Fulano", "Beltrana")
    leitor = _ids("Consulta Teste")[0]
    r = cliente.get("/inventario")
    assert b'name="usuarios"' in r.data and b'name="integrantes"' not in r.data                     # web manda IDs
    assert b"Fulano (admin)" in r.data and b"Beltrana (beltrana)" in r.data
    assert b"Operador Teste" not in r.data and b"Consulta Teste" not in r.data                      # só admin/inventariante
    r = cliente.post("/inventario/abrir", data={"nome": "Inv", "usuarios": [leitor], "escopo": "todas"}, follow_redirects=True)
    assert "ao menos um usuário ativo com função de inventário".encode() in r.data                  # ID oculto recusado
    r = cliente.post("/inventario/abrir", data={"nome": "Inv", "usuarios": ["Beltrana"], "escopo": "todas"}, follow_redirects=True)
    assert "Selecione integrantes válidos".encode() in r.data                                       # nome no lugar do ID
    eid = _abrir(cliente, comissao=["Beltrana"])
    r = cliente.get(f"/inventario/{eid}")
    assert b"Comiss" in r.data and b'href="/inventario/%d/comissao"' % eid in r.data and b"Quem est" not in r.data
    r = cliente.get(f"/inventario/{eid}/sala/01 - SALA CCI")
    assert "não faz parte da comissão".encode() in r.data and b'id="leitura" type="text" inputmode="none" autocomplete="off" enterkeyhint="done" placeholder="Aproxime o leitor\xe2\x80\xa6" disabled' in r.data
    r = cliente.get(f"/inventario/{eid}/comissao")
    assert r.status_code == 200 and b'value="%d" checked' % beltrana in r.data and b'value="%d"/' % fulano in r.data
    assert b"sem conta vinculada" not in r.data                                                     # toda a comissão tem conta
    r = cliente.post(f"/inventario/{eid}/comissao", data={"usuarios": [leitor]}, follow_redirects=True)
    assert "ao menos um usuário ativo com função de inventário".encode() in r.data
    assert inventario_do_teste(eid)["integrantes"] == ["Beltrana"]                                   # comissão intacta
    r = cliente.post(f"/inventario/{eid}/comissao", data={"usuarios": _ids("Fulano", "Beltrana")}, follow_redirects=True)
    assert "Comissão atualizada".encode() in r.data
    j = cliente.post(f"/inventario/{eid}/sala/01 - SALA CCI/ler", json={"numero": "1001"}).get_json()
    assert j["situacao"] == "localizado" and j["integrante"] == "Fulano"
    assert b"lendo como <strong>Fulano</strong>" in cliente.get(f"/inventario/{eid}/sala/01 - SALA CCI").data
    r = cliente.get(f"/inventario/{eid}/comissao")                                                  # aviso de quem já leu
    assert r.data.count("já tem leituras neste evento".encode()) == 1
    assert "Fulano (admin) <span class=\"text-gray-70 text-down-01\">· já tem leituras".encode() in r.data
    cliente.post("/sair"); logar(cliente, *usuarios_exemplo["operador"])
    assert cliente.get(f"/inventario/{eid}/comissao").status_code == 403                         # só admin
    r = cliente.post(f"/inventario/{eid}/sala/01 - SALA CCI/ler", json={"numero": "1002"})
    assert r.status_code == 403                                                                  # operador não confere inventário
    cliente.post("/sair"); logar(cliente, *usuarios_exemplo["inventariante"])
    j = cliente.post(f"/inventario/{eid}/sala/01 - SALA CCI/ler", json={"numero": "1002"}).get_json()
    assert j["integrante"] == "Beltrana"
    assert cliente.get("/inventario/%d/integrante" % eid).status_code in (404, 405)


def test_relatorio_nao_linka_ficha_do_bem_para_quem_nao_a_abre(cliente, dados):
    """`consulta_inventarios` alcança o relatório mas não a ficha do bem (ACERVO): o número vira texto."""
    import usuarios
    usuarios.criar(dados, "chefe", "Chefe", SENHA_PADRAO, ["consulta_inventarios"], trocar_senha=False)
    eid = _abrir(cliente)
    cliente.post(f"/inventario/{eid}/sala/01 - SALA CCI/ler", json={"numero": "1001"})
    r = cliente.get(f"/inventario/{eid}/relatorio")
    assert r.status_code == 200 and b'href="/bem?numero=1001"' in r.data
    cliente.post("/sair"); logar(cliente, "chefe", SENHA_PADRAO)
    r = cliente.get(f"/inventario/{eid}/relatorio")
    assert r.status_code == 200 and b">1001<" in r.data.split(b"<tbody>")[1] and b"/bem?numero=" not in r.data
    assert cliente.get("/bem?numero=1001").status_code == 403


def test_sala_somente_consulta_para_quem_nao_confere(cliente, dados):
    """Evento ABERTO visto por quem não lê bens: controles desligados, sem dizer que o evento foi encerrado."""
    import usuarios
    usuarios.criar(dados, "chefe", "Chefe", SENHA_PADRAO, ["consulta_inventarios"], trocar_senha=False)
    eid = _abrir(cliente)
    cliente.post("/sair"); logar(cliente, "chefe", SENHA_PADRAO)
    r = cliente.get(f"/inventario/{eid}")
    assert r.status_code == 200 and b"Ver sala 01 - SALA CCI" in r.data and b"Conferir sala" not in r.data
    r = cliente.get(f"/inventario/{eid}/sala/01 - SALA CCI")
    assert r.status_code == 200 and b'id="form-lote"' not in r.data               # nada de marcar/desmarcar em lote
    assert b'placeholder="Aproxime o leitor\xe2\x80\xa6" disabled' in r.data
    assert "Evento encerrado".encode() not in r.data                              # o evento está aberto
    assert "Somente consulta: seu usuário não confere bens neste evento".encode() in r.data
    r = cliente.post(f"/inventario/{eid}/sala/01 - SALA CCI/ler", json={"numero": "1001"})
    assert r.status_code == 403 and r.get_json()["erro"] == NEGADO


def test_sala_esconde_escrita_de_quem_confere_mas_nao_esta_na_comissao(cliente, monkeypatch):
    """Admin tem a função de conferir (CONFERENCIA), mas não está na comissão deste evento aberto: a tela
    precisa se comportar como somente consulta mesmo assim (antes, `fechado` só olhava a função)."""
    import fotos
    for v in fotos.VARIAVEIS:
        monkeypatch.setenv(v, "x")
    eid = _abrir(cliente, comissao=("Beltrana",))   # Fulano (admin logado) fica de fora da comissão
    import db, inventario
    conn = db.conectar()
    inventario.ler(conn, eid, "01 - SALA CCI", 1001, "Beltrana")
    inventario.adicionar_foto(conn, eid, 1001, lambda c: "https://x/1.webp")
    inventario.registrar_sobra(conn, eid, "01 - SALA CCI", "VENTILADOR", "", "achado", "https://x/s.webp", "Beltrana")
    r = cliente.get(f"/inventario/{eid}/sala/01 - SALA CCI")
    html = r.data
    assert r.status_code == 200
    assert b'id="form-sobra"' not in html and b'id="form-lote"' not in html and b'name="numeros" type="checkbox"' not in html
    assert b'class="foto-input"' not in html
    assert b'aria-label="Excluir foto 1"' not in html and b'aria-label="Excluir sobra"' not in html
    assert b"Somente consulta" in html and b"Evento encerrado" not in html
    # a Beltrana, que está na comissão, continua vendo os controles de escrita
    cliente.post("/sair"); logar(cliente, "beltrana", SENHA_PADRAO)
    r = cliente.get(f"/inventario/{eid}/sala/01 - SALA CCI")
    html = r.data
    assert b'id="form-sobra"' in html and b'id="form-lote"' in html and b'name="numeros" type="checkbox"' in html
    assert b'class="foto-input"' in html
    assert b'aria-label="Excluir foto 1"' in html and b'aria-label="Excluir sobra"' in html


def test_sair_funciona_mesmo_para_quem_ficou_sem_nenhuma_funcao(cliente):
    """Não alcançável pela UI (só editando o banco), mas precisa continuar podendo sair: senão a conta trava."""
    conn = db.conectar()
    conn.execute("DELETE FROM usuarios_funcoes WHERE usuario_id = (SELECT id FROM usuarios WHERE login = ?)", (ADMIN_LOGIN,))
    conn.commit()
    assert cliente.post("/sair").status_code == 302


def test_inventario_desktop_admin_local_entra_na_comissao(cliente_local):
    import db, inventario, usuarios
    conn = db.conectar()
    usuarios.criar(conn, "xis", "Xis", "Senha!234", ["inventariante"])
    r = cliente_local.post("/inventario/abrir", data={"nome": "Inv", "integrantes": ["Xis"], "escopo": "todas"}, follow_redirects=True)
    assert r.status_code == 200
    eid = inventario.evento_aberto(conn)["id"]
    assert inventario.evento(conn, eid)["integrantes"] == ["Administrador local", "Xis"]
    j = cliente_local.post(f"/inventario/{eid}/sala/01 - SALA CCI/ler", json={"numero": "1001"}).get_json()
    assert j["integrante"] == "Administrador local"


def test_inventario_fora_da_comissao_nao_edita_nem_apaga(cliente, monkeypatch):
    """Admin fora da comissão: toda escrita é negada (403) antes da view, sem tocar no banco nem no bucket."""
    import io
    from PIL import Image
    import db, fotos, inventario
    eid = _abrir(cliente, comissao=["Beltrana"])                      # Fulano (logado) não está na comissão
    cliente.post("/sair"); logar(cliente, "beltrana", SENHA_PADRAO)
    cliente.post(f"/inventario/{eid}/sala/01 - SALA CCI/ler", json={"numero": "1001"})
    cliente.post(f"/inventario/{eid}/sala/01 - SALA CCI/sobra", data={"descricao": "VENTILADOR", "observacao": "sem plaqueta"})
    sobra_id = inventario.bens_da_sala(db.conectar(), eid, "01 - SALA CCI")["sobras"][0]["id"]
    cliente.post("/sair"); logar(cliente, ADMIN_LOGIN, ADMIN_SENHA)
    r = cliente.post(f"/inventario/{eid}/leitura/1001", json={"conservacao": "Bom"})
    assert r.status_code == 403 and r.get_json()["erro"] == NEGADO
    for v in fotos.VARIAVEIS:
        monkeypatch.setenv(v, "x")
    monkeypatch.setenv("R2_PUBLIC_URL", "https://f.exemplo.org")
    tocou = []                                     # negado não fala com o bucket
    monkeypatch.setattr(fotos, "enviar", lambda chave, dados: tocou.append(("enviar", chave)))
    monkeypatch.setattr(fotos, "apagar", lambda url: tocou.append(("apagar", url)))
    buf = io.BytesIO(); Image.new("RGB", (30, 20), (1, 2, 3)).save(buf, "PNG"); imagem = buf.getvalue()
    r = cliente.post(f"/inventario/{eid}/leitura/1001/foto", data={"foto": (io.BytesIO(imagem), "a.png")}, content_type="multipart/form-data")
    assert r.status_code == 403 and r.get_json()["erro"] == NEGADO
    r = cliente.post(f"/inventario/{eid}/leitura/1001/foto/1/excluir", follow_redirects=True)
    assert r.status_code == 403 and NEGADO.encode() in r.data
    r = cliente.post(f"/inventario/{eid}/sobra/{sobra_id}/excluir", follow_redirects=True)
    assert r.status_code == 403 and NEGADO.encode() in r.data
    assert any(s["id"] == sobra_id for s in inventario.bens_da_sala(db.conectar(), eid, "01 - SALA CCI")["sobras"])   # sobra continua
    n_antes = db.conectar().execute("SELECT count(*) FROM inventario_leituras WHERE evento_id = ? AND numero = 1001", (eid,)).fetchone()[0]
    r = cliente.post(f"/inventario/{eid}/sala/01 - SALA CCI/lote", data={"acao": "desmarcar", "numeros": ["1001"]}, follow_redirects=True)
    assert r.status_code == 403 and NEGADO.encode() in r.data
    n_depois = db.conectar().execute("SELECT count(*) FROM inventario_leituras WHERE evento_id = ? AND numero = 1001", (eid,)).fetchone()[0]
    assert n_depois == n_antes                                    # leitura de Beltrana continua
    r = cliente.post(f"/inventario/{eid}/sala/01 - SALA CCI/lote", data={"acao": "marcar", "numeros": ["1002"]}, follow_redirects=True)
    assert r.status_code == 403 and NEGADO.encode() in r.data
    assert tocou == []
    cliente.post(f"/inventario/{eid}/comissao", data={"usuarios": _ids("Fulano", "Beltrana")})
    r = cliente.post(f"/inventario/{eid}/leitura/1001", json={"conservacao": "Bom"})
    assert r.status_code == 200


def test_renomear_usuario_mantem_na_comissao_do_evento_aberto(cliente):
    import db, usuarios
    eid = _abrir(cliente, comissao=["Fulano", "Beltrana"])
    admin_id = usuarios.por_login(db.conectar(), ADMIN_LOGIN)["id"]
    r = cliente.post(f"/usuarios/{admin_id}/editar", data={"nome": "Fulano Silva", "funcoes": ["admin"], "ativo": "1"}, follow_redirects=True)
    assert r.status_code == 200
    j = cliente.post(f"/inventario/{eid}/sala/01 - SALA CCI/ler", json={"numero": "1001"}).get_json()
    assert j["situacao"] == "localizado" and j["integrante"] == "Fulano Silva"


def test_inventario_comissao_so_em_evento_aberto(cliente):
    eid = _abrir(cliente)
    cliente.post(f"/inventario/{eid}/encerrar", data={"confirmar": "1"})
    r = cliente.get(f"/inventario/{eid}/comissao", follow_redirects=True)
    assert b"encerrado" in r.data.lower()


def test_inventario_excluir_evento(cliente, dados, monkeypatch, usuarios_exemplo):
    import fotos, inventario
    eid = _abrir(cliente)
    cliente.post(f"/inventario/{eid}/sala/01 - SALA CCI/ler", json={"numero": "1001"})
    inventario.adicionar_foto(dados, eid, 1001, lambda c: "https://x/a.webp")
    r = cliente.get(f"/inventario/{eid}")
    assert b"Excluir evento" in r.data
    r = cliente.get(f"/inventario/{eid}/excluir")
    assert r.status_code == 200 and b"1 leitura" in r.data and b"1 foto" in r.data and b'name="nome"' in r.data
    monkeypatch.setattr(fotos, "apagar", lambda url: (_ for _ in ()).throw(RuntimeError("bucket fora")))
    r = cliente.post(f"/inventario/{eid}/excluir", data={"nome": "Inv"}, follow_redirects=True)
    assert "Não foi possível apagar as fotos no bucket; o evento foi mantido".encode() in r.data
    assert inventario.evento(dados, eid) is not None
    apagadas = []
    monkeypatch.setattr(fotos, "apagar", apagadas.append)
    r = cliente.post(f"/inventario/{eid}/excluir", data={"nome": "Errado"}, follow_redirects=True)
    assert "não confere".encode() in r.data and inventario.evento(dados, eid) is not None
    r = cliente.post(f"/inventario/{eid}/excluir", data={"nome": "Inv"}, follow_redirects=True)
    assert "Evento Inv excluído".encode() in r.data and apagadas == ["https://x/a.webp"]
    assert inventario.evento(dados, eid) is None and cliente.get(f"/inventario/{eid}").status_code == 404
    eid = _abrir(cliente)
    cliente.post("/sair"); logar(cliente, *usuarios_exemplo["operador"])
    assert cliente.get(f"/inventario/{eid}/excluir").status_code == 403
    assert cliente.post(f"/inventario/{eid}/excluir", data={"nome": "Inv"}).status_code == 403
    assert b"Excluir evento" not in cliente.get(f"/inventario/{eid}").data


def _xlsx_base():
    wb = Workbook()
    ws = wb.active
    ws.append(["Número Bem", "Situação", "Descrição", "Complemento", "Classificação Contábil",
               "Localização", "Data Entrada", "Valor Compra", "Valor Atual"])
    ws.append([1002, "ATIVO", "NOTEBOOK", "DELL", "EQUIP", "01 - SALA CCI", "06/12/2012", 3000, 1500])
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf


def test_inicio_sem_importacao_fica_em_alerta_e_sem_linha_do_robo(cliente):
    r = cliente.get("/")
    assert b"dsgov-kpi-alerta" in r.data and b"Nenhuma" in r.data
    assert "robô".encode() not in r.data
    assert "base desatualizada".encode() in r.data


def test_inicio_mostra_base_desatualizada_quando_alerta_nao_e_erro(cliente, dados):
    db.importar_bens(dados, _xlsx_base(), nome_arquivo="SPW automático")
    dados.execute("UPDATE importacoes SET importado_em = '2026-01-01 04:00:00'")
    dados.commit()
    r = cliente.get("/")
    assert "base desatualizada".encode() in r.data and b"dsgov-kpi-alerta" in r.data


def test_inicio_mostra_robo_ok_sem_alerta(cliente, dados):
    db.importar_bens(dados, _xlsx_base(), nome_arquivo="SPW automático")
    db.registrar_execucao_robo(dados, "2026-09-18 04:00:00", "sem_mudanca", hash="h")
    r = cliente.get("/")
    assert "SPW ok · 18/09".encode() in r.data
    assert b"dsgov-kpi-alerta" not in r.data


def test_inicio_mostra_robo_falhou_em_alerta(cliente, dados):
    db.importar_bens(dados, _xlsx_base(), nome_arquivo="SPW automático")
    db.registrar_execucao_robo(dados, "2026-09-18 04:00:00", "erro", mensagem="Timeout 60000ms exceeded")
    r = cliente.get("/")
    assert "SPW falhou 18/09: Timeout 60000ms exceeded".encode() in r.data
    assert b"dsgov-kpi-alerta" in r.data and b"dsgov-robo-erro" in r.data


def test_inicio_esconde_robo_de_quem_nao_ve_importacao(cliente, dados, usuarios_exemplo):
    db.registrar_execucao_robo(dados, "2026-09-18 04:00:00", "erro", mensagem="senha")
    cliente.post("/sair")
    assert logar(cliente, *usuarios_exemplo["consulta"]).status_code == 302
    assert "robô".encode() not in cliente.get("/").data


def test_upload_lista_execucoes_do_robo(cliente, dados):
    assert "Atualizações com o SPW".encode() not in cliente.get("/upload").data
    resumo = db.importar_bens(dados, _xlsx_base(), nome_arquivo="SPW automático")
    db.registrar_execucao_robo(dados, "2026-09-18 04:00:00", "importado", hash="h", importacao_id=resumo["importacao_id"], mensagem="1 bens")
    db.registrar_execucao_robo(dados, "2026-09-19 04:00:00", "sem_mudanca", hash="h", mensagem="nada mudou")
    db.registrar_execucao_robo(dados, "2026-09-20 04:00:00", "erro", mensagem="SPW fora do ar")
    r = cliente.get("/upload")
    html = r.data.decode()
    assert "Atualizações com o SPW" in html
    assert html.index("SPW fora do ar") < html.index("nada mudou") < html.index("1 bens")   # mais recente primeiro
    assert f'href="/importacoes/{resumo["importacao_id"]}"' in html.split("Atualizações com o SPW")[-1]
    assert "sem mudança" in html and "bg-danger" in html


def test_atualizar_com_spw_sem_acesso_avisa(cliente, dados, chave):
    r = cliente.post("/atualizar-base/spw", follow_redirects=True)
    assert "Cadastre seu acesso ao SPW em Meus acessos para atualizar.".encode() in r.data
    assert db.pedido_spw_ativo(dados) is None
    _acesso_spw_admin(dados, chave)
    assert cliente.post("/atualizar-base/spw").status_code == 302
    assert db.pedido_spw_ativo(dados) is not None


def test_atualizar_com_spw_enfileira_e_mostra_andamento(cliente, dados, chave):
    _acesso_spw_admin(dados, chave)
    r = cliente.get("/upload")
    assert b"Atualizar com SPW" in r.data and b'action="/atualizar-base/spw"' in r.data and "robô".encode() not in r.data
    r = cliente.post("/atualizar-base/spw")
    assert r.status_code == 302
    p = db.pedido_spw_ativo(dados)
    assert p and p["criado_por"] == "admin"
    r = cliente.get("/upload")
    assert b"Atualizando com o SPW" in r.data and b'http-equiv="refresh" content="5"' in r.data and b"disabled" in r.data
    r = cliente.post("/atualizar-base/spw", follow_redirects=True)
    assert "A atualização com o SPW já está em andamento.".encode() in r.data
    db.marcar_passo(dados, p["id"], "rodando")
    r = cliente.get("/upload")
    assert "(atualizando com o SPW)".encode() not in r.data
    db.marcar_passo(dados, p["id"], "erro", "SPW fora do ar")
    r = cliente.get("/upload")
    assert b'http-equiv="refresh"' not in r.data and b"SPW fora do ar" in r.data


def test_desktop_nao_mostra_atualizar_com_spw(cliente_local):
    assert b"Atualizar com SPW" not in cliente_local.get("/upload").data


def test_atualizar_com_spw_so_admin(cliente, usuarios_exemplo):
    cliente.post("/sair")
    assert logar(cliente, *usuarios_exemplo["operador"]).status_code == 302
    assert cliente.post("/atualizar-base/spw").status_code == 403
    assert b"Atualizar com SPW" not in cliente.get("/upload").data


def test_formularios_de_cadastro_tem_unidade_sei(cliente):
    assert b'name="unidade_sei"' in cliente.get("/cadastros/responsaveis/CCI/editar").data
    assert b'name="unidade_sei"' in cliente.get("/cadastros/pessoas/ANA%20SILVA/editar").data
    cliente.post("/cadastros/pessoas/incluir", data={"nome": "BEA", "unidade_sei": "GECONT"})
    assert db.pessoa(db.conectar(), "BEA")["unidade_sei"] == "GECONT"
    cliente.post("/cadastros/responsaveis/CCI/editar", data={"ccustos": "CCI", "responsavel": "J", "unidade_sei": "GAB"})
    assert db.responsavel(db.conectar(), "CCI")["unidade_sei"] == "GAB"


def _processo_ccusto(cliente):
    cliente.post("/cadastros/processos/incluir", data={"tipo": "ccusto", "descricao": "T", "numero_sei": "2222", "vigente": "1"})


def test_botao_emitir_no_sei_na_pagina_do_termo(cliente):
    assert b"Emitir Termo no SEI" not in cliente.get("/termo/ccusto/CCI").data      # sem processo vigente
    _processo_ccusto(cliente)
    r = cliente.get("/termo/ccusto/CCI")
    assert b"Emitir Termo no SEI" in r.data and b'action="/termo/ccusto/CCI/enviar-sei"' in r.data
    assert "robô".encode() not in r.data


def test_emitir_registra_numera_e_enfileira(cliente, dados, chave):
    _acesso_admin(dados, chave)
    _processo_ccusto(cliente)
    r = cliente.post("/termo/ccusto/CCI/enviar-sei")
    assert r.status_code == 302 and r.headers["Location"].endswith("/termos-emitidos/1")
    conn = db.conectar()
    t = db.termo_emitido(conn, 1)
    assert t["numero_termo"].endswith(f"/{db._agora()[:4]}") and t["unidade_sei"] == "CCI"
    p = db.pedido_do_termo(conn, 1)
    assert p["tipo"] == "sei" and p["passo"] == "aguardando" and p["criado_por"] == "admin"
    assert "Termo de Responsabilidade - CCI" in p["html"] and "text-align:justify" in p["html"]
    r = cliente.get("/termos-emitidos/1")
    assert b"Emitindo no SEI" in r.data and b"aguardando a vez" in r.data and b'http-equiv="refresh" content="5"' in r.data
    assert b"Emitir Termo no SEI" not in r.data and b'name="documento_sei"' not in r.data
    r = cliente.post("/termo/ccusto/CCI/enviar-sei", follow_redirects=True)           # segundo clique
    assert "Já há uma emissão deste termo em andamento.".encode() in r.data
    assert b"enviando" in cliente.get("/termos-emitidos").data


def test_emitir_sem_processo_ou_sem_unidade_avisa(cliente):
    r = cliente.post("/termo/ccusto/CCI/enviar-sei", follow_redirects=True)
    assert b"Cadastre um processo SEI vigente" in r.data
    cliente.post("/cadastros/processos/incluir", data={"tipo": "individual", "descricao": "T", "numero_sei": "1111", "vigente": "1"})
    r = cliente.post("/termo/individual/ANA%20SILVA/enviar-sei", follow_redirects=True)
    assert "Cadastre a unidade SEI de ANA SILVA".encode() in r.data
    assert db.pedido_do_termo(db.conectar(), 1) is None
    assert db.termos_emitidos(db.conectar()) == []   # sem unidade: nenhum registro fantasma, sem documento


def test_emitir_sem_acesso_ao_sei_nao_registra(cliente, dados, chave):
    _processo_ccusto(cliente)
    r = cliente.post("/termo/ccusto/CCI/enviar-sei", follow_redirects=True)
    assert "Cadastre seu acesso ao SEI em Meus acessos para emitir.".encode() in r.data
    assert db.termos_emitidos(dados) == [] and db.pedido_do_termo(dados, 1) is None
    _acesso_admin(dados, chave)
    assert cliente.post("/termo/ccusto/CCI/enviar-sei").status_code == 302
    assert db.pedido_do_termo(dados, 1)["criado_por"] == ADMIN_LOGIN


def test_estados_da_pagina_do_termo_emitido(cliente, dados, chave):
    _acesso_admin(dados, chave)
    _processo_ccusto(cliente)
    cliente.post("/termo/ccusto/CCI/enviar-sei")
    p = db.pedido_do_termo(dados, 1)
    db.marcar_passo(dados, p["id"], "documento")
    assert b"criando o documento" in cliente.get("/termos-emitidos/1").data
    # parado: aguardando há mais de 2 min
    db.marcar_passo(dados, p["id"], "aguardando")
    dados.execute("UPDATE robo_pedidos SET criado_em = '2020-01-01 00:00:00' WHERE id = ?", (p["id"],)); dados.commit()
    assert "A emissão ainda não começou; avise o administrador.".encode() in cliente.get("/termos-emitidos/1").data
    # erro sem documento → botão de novo
    db.marcar_passo(dados, p["id"], "erro", "O SEI recusou usuário ou senha.")
    r = cliente.get("/termos-emitidos/1")
    assert b"O SEI recusou" in r.data and b"Emitir Termo no SEI" in r.data and b'action="/termos-emitidos/1/enviar-sei"' in r.data
    assert b'http-equiv="refresh"' not in r.data
    # erro com documento → Incluir no bloco
    db.salvar_documento_sei(dados, 1, "1557099", "")
    db.marcar_passo(dados, p["id"], "erro", "Bloco 'Termos CCI' não existe no SEI; crie o bloco e clique em Incluir no bloco.")
    r = cliente.get("/termos-emitidos/1")
    assert b"Incluir no bloco" in r.data and b"1557099" in r.data and b"Emitir Termo no SEI" not in r.data
    r = cliente.post("/termos-emitidos/1/enviar-sei")
    assert r.status_code == 302 and db.pedido_do_termo(dados, 1)["passo"] == "aguardando"
    # concluído → e-mail
    db.salvar_documento_sei(dados, 1, "1557099", "69766")
    db.marcar_passo(dados, db.pedido_do_termo(dados, 1)["id"], "concluido", "documento 1557099 no bloco 69766")
    r = cliente.get("/termos-emitidos/1")
    assert b"Emitido no SEI em" in r.data and b"documento 1557099, bloco 69766" in r.data
    assert b"Enviar email" in r.data and b"mailto:" in r.data and b"Emitir Termo no SEI" not in r.data


def test_erro_preenchido_a_mao_volta_ao_estado_manual(cliente, dados, chave):
    _acesso_admin(dados, chave)
    _processo_ccusto(cliente)
    cliente.post("/termo/ccusto/CCI/enviar-sei")
    p = db.pedido_do_termo(dados, 1)
    db.marcar_passo(dados, p["id"], "erro", "O SEI recusou usuário ou senha.")
    cliente.post("/termos-emitidos/1/documento", data={"documento_sei": "1557099", "bloco_sei": "69766"})
    r = cliente.get("/termos-emitidos/1")
    assert b"O SEI recusou" not in r.data and "Não foi possível emitir no SEI.".encode() not in r.data
    assert b"Incluir no bloco" not in r.data
    assert b"Enviar email" in r.data and b'class="br-button primary mr-3"' in r.data


def test_erro_bloco_permite_digitar_o_bloco_a_mao(cliente, dados, chave):
    _acesso_admin(dados, chave)
    _processo_ccusto(cliente)
    cliente.post("/termo/ccusto/CCI/enviar-sei")
    p = db.pedido_do_termo(dados, 1)
    db.salvar_documento_sei(dados, 1, "1557099", "")
    db.marcar_passo(dados, p["id"], "erro", "Bloco 'Termos CCI' não existe no SEI; crie o bloco e clique em Incluir no bloco.")
    r = cliente.get("/termos-emitidos/1")
    assert b'name="bloco_sei"' in r.data and b"Incluir no bloco" in r.data
    assert b'name="documento_sei" type="text" value="1557099" readonly' in r.data
    r = cliente.post("/termos-emitidos/1/documento", data={"documento_sei": "1557099", "bloco_sei": "69766"}, follow_redirects=True)
    assert b"Enviar email" in r.data and b"Incluir no bloco" not in r.data
    assert "Não foi possível emitir no SEI.".encode() not in r.data


def test_pagina_sem_pedido_mantem_campos_manuais_e_numero_editavel(cliente, dados, chave):
    _acesso_admin(dados, chave)
    _processo_ccusto(cliente)
    j = cliente.post("/termo/ccusto/CCI/registrar").get_json()
    r = cliente.get(f"/termos-emitidos/{j['id']}")
    assert b'name="documento_sei"' in r.data and b'name="numero_termo"' in r.data and b"Emitir Termo no SEI" in r.data
    cliente.post(f"/termos-emitidos/{j['id']}/documento", data={"documento_sei": "", "bloco_sei": "", "numero_termo": "07/2026"})
    assert db.termo_emitido(dados, j["id"])["numero_termo"] == "07/2026"
    r = cliente.post(f"/termos-emitidos/{j['id']}/documento", data={"documento_sei": "555", "bloco_sei": "8", "numero_termo": "x"}, follow_redirects=True)
    assert b"NN/AAAA" in r.data
    t = db.termo_emitido(dados, j["id"])
    assert t["documento_sei"] == "555" and t["bloco_sei"] == "8" and t["numero_termo"] == "07/2026"   # número inválido não perde documento/bloco já digitados
    cliente.post(f"/termos-emitidos/{j['id']}/documento", data={"documento_sei": "123", "bloco_sei": "9", "numero_termo": "07/2026"})
    r = cliente.get(f"/termos-emitidos/{j['id']}")
    assert b"Emitir Termo no SEI" not in r.data and b"Enviar email" in r.data          # preenchido à mão: como hoje


def test_desktop_nao_mostra_emitir_no_sei(cliente_local):
    cliente_local.post("/cadastros/processos/incluir", data={"tipo": "ccusto", "descricao": "T", "numero_sei": "2222", "vigente": "1"})
    assert b"Emitir Termo no SEI" not in cliente_local.get("/termo/ccusto/CCI").data


def test_desktop_recusa_post_que_dependem_da_fila(cliente_local):
    """Sem TERMOS_LOGIN não há trabalhador atendendo robo_pedidos: estas três rotas ficariam com pedidos parados."""
    cliente_local.post("/cadastros/processos/incluir", data={"tipo": "ccusto", "descricao": "T", "numero_sei": "2222", "vigente": "1"})
    assert cliente_local.post("/termo/ccusto/CCI/enviar-sei").status_code == 403
    assert db.pedido_do_termo(db.conectar(), 1) is None
    j = cliente_local.post("/termo/ccusto/CCI/registrar").get_json()
    assert cliente_local.post(f"/termos-emitidos/{j['id']}/enviar-sei").status_code == 403
    assert db.pedido_do_termo(db.conectar(), j["id"]) is None
    assert cliente_local.post("/atualizar-base/spw").status_code == 403
    assert db.pedido_spw_ativo(db.conectar()) is None
