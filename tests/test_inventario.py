"""Módulo de inventário: eventos, salas, leituras, sobras, relatório, xlsx."""
import pytest

import db
import inventario
from tests.conftest import semear


def semear_inventario(conn):
    """semear() + 2ª sala com 2 bens e 1 sala vazia de escopo; devolve o id do evento aberto."""
    semear(conn)
    conn.execute("INSERT INTO bens VALUES (2001,'ATIVO','MONITOR','LG','EQUIPAMENTOS','02 - SALA B','01/01/2020',900,800)")
    conn.execute("INSERT INTO bens VALUES (2002,'ATIVO','IMPRESSORA','HP','EQUIPAMENTOS','02 - SALA B','01/01/2020',1200,1000)")
    conn.commit()
    return inventario.abrir_evento(conn, "Inventário 2026", "Portaria 1/2026", ["Fulano", "Beltrana"])


def foto_falsa(conn, eid, numero, url=None):
    """adicionar_foto com envio falso; devolve a lista de fotos do bem no evento."""
    return inventario.adicionar_foto(conn, eid, numero, lambda chave: url or f"https://x/{chave}")


def test_localizacoes_ativas(dados):
    semear(dados)
    assert db.localizacoes_ativas(dados) == ["01 - SALA CCI", "99 - SEM MAPA"]   # 1003 é BAIXADO, não muda nada


def test_abrir_evento_todas_as_salas_e_integrantes(dados):
    eid = semear_inventario(dados)
    e = inventario.evento(dados, eid)
    assert e["nome"] == "Inventário 2026" and e["encerrado_em"] is None and e["integrantes"] == ["Beltrana", "Fulano"]
    assert [s["localizacao"] for s in inventario.salas(dados, eid)] == ["01 - SALA CCI", "02 - SALA B", "99 - SEM MAPA"]
    assert inventario.evento_aberto(dados)["id"] == eid
    assert e["resumo"] == {"salas": 3, "salas_iniciadas": 0, "bens": 5, "lidos": 0, "divergentes": 0,
                           "pendentes": 5, "sobras": 0, "pct_bens": 0.0}


def test_abrir_evento_amostragem_e_validacoes(dados):
    semear(dados)
    with pytest.raises(db.ErroDeNegocio):
        inventario.abrir_evento(dados, "", "", ["A"])
    with pytest.raises(db.ErroDeNegocio):
        inventario.abrir_evento(dados, "X", "", [" ", ""])
    with pytest.raises(db.ErroDeNegocio):
        inventario.abrir_evento(dados, "X", "", ["A"], salas=["NÃO EXISTE"])
    eid = inventario.abrir_evento(dados, "Amostra", None, ["A", "A ", "b"], salas=["99 - SEM MAPA"])
    assert [s["localizacao"] for s in inventario.salas(dados, eid)] == ["99 - SEM MAPA"]
    assert inventario.evento(dados, eid)["integrantes"] == ["A", "b"]
    inventario.encerrar_evento(dados, eid)
    assert inventario.evento_aberto(dados) is None
    e2 = inventario.abrir_evento(dados, "Outro", "", ["A"])
    assert [x["id"] for x in inventario.eventos(dados)] == [e2, eid]   # aberto primeiro


def test_abrir_evento_com_elegiveis_e_editar_comissao(dados):
    semear(dados)
    with pytest.raises(db.ErroDeNegocio, match="não pode compor"):
        inventario.abrir_evento(dados, "X", "", ["Fulano", "Zé"], elegiveis=["Fulano", "Beltrana"])
    eid = inventario.abrir_evento(dados, "X", "", ["Fulano"], elegiveis=["Fulano", "Beltrana"])
    assert inventario.evento(dados, eid)["integrantes"] == ["Fulano"]
    with pytest.raises(db.ErroDeNegocio, match="não pode compor"):
        inventario.editar_comissao(dados, eid, ["Zé"], elegiveis=["Fulano", "Beltrana"])
    with pytest.raises(db.ErroDeNegocio, match="ao menos um"):
        inventario.editar_comissao(dados, eid, [])
    inventario.ler(dados, eid, "01 - SALA CCI", 1001, "Fulano")
    inventario.editar_comissao(dados, eid, ["Beltrana", " Fulano "], elegiveis=["Fulano", "Beltrana"])
    assert inventario.evento(dados, eid)["integrantes"] == ["Beltrana", "Fulano"]
    inventario.editar_comissao(dados, eid, ["Beltrana"])
    assert inventario.evento(dados, eid)["integrantes"] == ["Beltrana"]
    assert dados.execute("SELECT integrante FROM inventario_leituras WHERE evento_id = ?", (eid,)).fetchone()[0] == "Fulano"  # leitura fica
    with pytest.raises(db.ErroDeNegocio, match="não faz parte da comissão"):
        inventario.ler(dados, eid, "01 - SALA CCI", 1002, "Fulano")
    inventario.encerrar_evento(dados, eid)
    with pytest.raises(db.ErroDeNegocio, match="encerrado"):
        inventario.editar_comissao(dados, eid, ["Fulano"])


def test_renomear_integrante_mantem_comissao_do_evento_aberto(dados):
    eid = semear_inventario(dados)
    inventario.ler(dados, eid, "01 - SALA CCI", 1001, "Fulano")
    n = inventario.renomear_integrante(dados, "Fulano", "Fulano Silva")
    assert n == 1
    assert inventario.evento(dados, eid)["integrantes"] == ["Beltrana", "Fulano Silva"]
    assert dados.execute("SELECT integrante FROM inventario_leituras WHERE evento_id = ? AND numero = 1001",
                         (eid,)).fetchone()[0] == "Fulano"                          # leitura mantém o nome antigo
    inventario.encerrar_evento(dados, eid)
    n2 = inventario.renomear_integrante(dados, "Beltrana", "Beltrana Silva")
    assert n2 == 0                                                                  # evento encerrado: não muda
    assert inventario.evento(dados, eid)["integrantes"] == ["Beltrana", "Fulano Silva"]


def test_salas_contadores_e_resumo(dados):
    eid = semear_inventario(dados)
    s = {x["localizacao"]: x for x in inventario.salas(dados, eid)}
    assert s["01 - SALA CCI"]["total"] == 2 and s["01 - SALA CCI"]["ccustos"] == "CCI" and s["02 - SALA B"]["ccustos"] is None
    assert (s["01 - SALA CCI"]["localizados"], s["01 - SALA CCI"]["pendentes"], s["01 - SALA CCI"]["divergentes"]) == (0, 2, 0)
    assert "concluida_em" not in s["01 - SALA CCI"]
    dados.execute("INSERT INTO inventario_leituras (evento_id, numero, localizacao, lido_em, integrante) VALUES (?,?,?,?,?)",
                  (eid, 2001, "02 - SALA B", "2026-09-15 10:00:00", "Fulano"))
    dados.commit()
    r = inventario.resumo(dados, eid)
    assert (r["salas_iniciadas"], r["lidos"], r["pendentes"], r["pct_bens"]) == (1, 1, 4, 20.0)
    inventario.encerrar_evento(dados, eid)
    assert inventario.evento(dados, eid)["encerrado_em"] is not None
    inventario.encerrar_evento(dados, eid)                                        # idempotente
    with pytest.raises(db.ErroDeNegocio):
        inventario.encerrar_evento(dados, 999)


def test_ler_localizado_divergente_reler_e_erros(dados):
    eid = semear_inventario(dados)
    r = inventario.ler(dados, eid, "01 - SALA CCI", 1001, "Fulano")
    assert r["situacao"] == "localizado" and r["reler"] is False and r["ativo"] and r["bem"]["descricao"] == "CADEIRA"
    r = inventario.ler(dados, eid, "01 - SALA CCI", 2001, "Fulano")             # MONITOR é da SALA B
    assert r["situacao"] == "divergente" and r["cadastrado_em"] == "02 - SALA B"
    r = inventario.ler(dados, eid, "02 - SALA B", 2001, "Beltrana")             # reler: atualiza a mesma linha
    assert r["situacao"] == "localizado" and r["reler"] and r["leitura_anterior"]["localizacao"] == "01 - SALA CCI"
    assert dados.execute("SELECT count(*) FROM inventario_leituras WHERE evento_id = ?", (eid,)).fetchone()[0] == 2
    r = inventario.ler(dados, eid, "01 - SALA CCI", 1003, "Fulano")             # BAIXADO: registra, avisa
    assert r["ativo"] is False and r["situacao"] == "localizado"
    with pytest.raises(inventario.BemNaoEncontrado) as e:
        inventario.ler(dados, eid, "01 - SALA CCI", 99999, "Fulano")
    assert e.value.numero == 99999
    with pytest.raises(db.ErroDeNegocio):
        inventario.ler(dados, eid, "SALA QUE NÃO EXISTE", 1001, "Fulano")
    with pytest.raises(db.ErroDeNegocio):
        inventario.ler(dados, eid, "01 - SALA CCI", 1001, "Ninguém")
    s = {x["localizacao"]: x for x in inventario.salas(dados, eid)}
    assert (s["01 - SALA CCI"]["localizados"], s["01 - SALA CCI"]["pendentes"]) == (1, 1)
    assert (s["02 - SALA B"]["localizados"], s["02 - SALA B"]["divergentes"]) == (1, 0)
    assert s["01 - SALA CCI"]["divergentes"] == 0

    inventario.ler(dados, eid, "02 - SALA B", 1003, "Fulano")                   # BAIXADO relido em outra sala
    s = {x["localizacao"]: x for x in inventario.salas(dados, eid)}
    assert s["02 - SALA B"]["divergentes"] == 0                                # BAIXADO não conta nem aqui

    inventario.ler(dados, eid, "02 - SALA B", 1001, "Fulano")                  # controle: ATIVO de outra sala conta
    s = {x["localizacao"]: x for x in inventario.salas(dados, eid)}
    assert s["02 - SALA B"]["divergentes"] == 1
    assert (s["01 - SALA CCI"]["localizados"], s["01 - SALA CCI"]["pendentes"]) == (0, 2)   # 1001 saiu da CCI
    assert s["02 - SALA B"]["localizados"] == 1                                             # 2001 continua localizado aqui

    inventario.encerrar_evento(dados, eid)
    with pytest.raises(db.ErroDeNegocio):
        inventario.ler(dados, eid, "01 - SALA CCI", 1002, "Fulano")


def test_bens_da_sala_e_atualizar_leitura(dados):
    eid = semear_inventario(dados)
    inventario.ler(dados, eid, "01 - SALA CCI", 1001, "Fulano")
    inventario.ler(dados, eid, "01 - SALA CCI", 2002, "Fulano")                 # trazido da SALA B
    inventario.ler(dados, eid, "02 - SALA B", 1002, "Fulano")                   # bem da CCI lido na SALA B
    inventario.ler(dados, eid, "01 - SALA CCI", 1003, "Fulano")                 # BAIXADO da própria sala
    d = inventario.bens_da_sala(dados, eid, "01 - SALA CCI")
    por = {b["numero"]: b for b in d["bens"]}
    assert set(por) == {1001, 1002}                                              # só ativos da sala
    assert por[1001]["situacao_inv"] == "localizado" and por[1002]["situacao_inv"] == "divergente" and por[1002]["lido_em_sala"] == "02 - SALA B"
    assert [t["numero"] for t in d["trazidos"]] == [1003, 2002] and d["sobras"] == []
    inventario.atualizar_leitura(dados, eid, 1001, conservacao="Ruim", quem_usa="Ciclana", observacao="pé quebrado")
    b = {x["numero"]: x for x in inventario.bens_da_sala(dados, eid, "01 - SALA CCI")["bens"]}[1001]
    assert (b["conservacao"], b["quem_usa"], b["observacao"]) == ("Ruim", "Ciclana", "pé quebrado")
    inventario.atualizar_leitura(dados, eid, 1001, conservacao="")                # limpa
    assert {x["numero"]: x for x in inventario.bens_da_sala(dados, eid, "01 - SALA CCI")["bens"]}[1001]["conservacao"] is None
    with pytest.raises(db.ErroDeNegocio):
        inventario.atualizar_leitura(dados, eid, 1001, conservacao="Ótimo")
    with pytest.raises(db.ErroDeNegocio):
        inventario.atualizar_leitura(dados, eid, 2001, quem_usa="x")             # sem leitura


def test_sobras(dados):
    eid = semear_inventario(dados)
    with pytest.raises(db.ErroDeNegocio):
        inventario.registrar_sobra(dados, eid, "01 - SALA CCI", "", "", "achado", "http://x/1.webp", "Fulano")
    with pytest.raises(db.ErroDeNegocio):
        inventario.registrar_sobra(dados, eid, "01 - SALA CCI", "VENTILADOR", "", "", "http://x/1.webp", "Fulano")
    with pytest.raises(db.ErroDeNegocio):
        inventario.registrar_sobra(dados, eid, "01 - SALA CCI", "VENTILADOR", "", "achado", "", "Fulano")
    with pytest.raises(db.ErroDeNegocio):
        inventario.registrar_sobra(dados, eid, "01 - SALA CCI", "VENTILADOR", "", "achado", "", "Ninguém", exigir_foto=False)
    sid = inventario.registrar_sobra(dados, eid, "01 - SALA CCI", "VENTILADOR", "", "achado", "", "Fulano", exigir_foto=False)
    inventario.definir_foto_sobra(dados, sid, "http://x/1.webp")
    s = inventario.bens_da_sala(dados, eid, "01 - SALA CCI")["sobras"]
    assert len(s) == 1 and s[0]["foto_url"] == "http://x/1.webp" and inventario.resumo(dados, eid)["sobras"] == 1
    with pytest.raises(db.ErroDeNegocio):
        inventario.excluir_sobra(dados, eid, 999)
    assert inventario.excluir_sobra(dados, eid, sid)["descricao"] == "VENTILADOR"
    assert inventario.resumo(dados, eid)["sobras"] == 0
    sid2 = inventario.registrar_sobra(dados, eid, "01 - SALA CCI", "TABLET", "", "achado", "", "Fulano", exigir_foto=False)
    inventario.encerrar_evento(dados, eid)
    with pytest.raises(db.ErroDeNegocio):
        inventario.definir_foto_sobra(dados, sid2, "http://x/2.webp")


def test_relatorio_e_xlsx(dados, tmp_path):
    eid = semear_inventario(dados)
    inventario.ler(dados, eid, "01 - SALA CCI", 1001, "Fulano")
    inventario.ler(dados, eid, "01 - SALA CCI", 2001, "Fulano")                     # divergente (é da SALA B)
    inventario.atualizar_leitura(dados, eid, 1001, conservacao="Bom", quem_usa="Ciclana")
    inventario.registrar_sobra(dados, eid, "02 - SALA B", "VENTILADOR", "ARNO", "sem plaqueta", "http://x/s.webp", "Beltrana")
    inventario.ler(dados, eid, "01 - SALA CCI", 1003, "Fulano")                     # BAIXADO lido na própria sala
    r = inventario.relatorio(dados, eid)
    b1003 = next(x for x in r if x["numero"] == 1003)
    assert b1003["situacao_bem"] == "BAIXADO" and b1003["situacao_inv"] == "localizado"
    assert [x["numero"] for x in inventario.relatorio(dados, eid, localizacao="01 - SALA CCI")] == [1001, 1002, 1003, 2001]
    with pytest.raises(db.ErroDeNegocio):
        inventario.relatorio(dados, 999)
    with pytest.raises(db.ErroDeNegocio):
        inventario.exportar_xlsx(dados, 999, tmp_path / "x.xlsx")
    assert [(x["numero"], x["situacao_inv"]) for x in r] == [(1001, "localizado"), (1002, "pendente"), (1003, "localizado"), (2001, "divergente"), (2002, "pendente"), (1004, "pendente")]
    assert r[3]["local_sistema"] == "02 - SALA B" and r[3]["local_inventario"] == "01 - SALA CCI"
    assert [x["numero"] for x in inventario.relatorio(dados, eid, situacao="pendente")] == [1002, 2002, 1004]
    from openpyxl import load_workbook
    wb = load_workbook(inventario.exportar_xlsx(dados, eid, tmp_path / "inv.xlsx"))
    assert wb.sheetnames == ["Bens", "Sobras"]
    linhas = list(wb["Bens"].iter_rows(values_only=True))
    assert linhas[0][0] == "Inventário 2026" and linhas[4] == tuple(inventario.COLUNAS_XLSX)
    assert linhas[5][:3] == (1001, "CADEIRA", "GIRATÓRIA") and linhas[5][6] == "Localizado" and linhas[5][7] == "Bom"
    sobras = list(wb["Sobras"].iter_rows(values_only=True))
    assert sobras[1][0] == "Sala" and sobras[2][:2] == ("02 - SALA B", "VENTILADOR") and sobras[2][6] == "http://x/s.webp"
    so_cci = load_workbook(inventario.exportar_xlsx(dados, eid, tmp_path / "cci.xlsx", localizacao="01 - SALA CCI"))["Bens"]
    assert so_cci.max_row == 5 + 4


def test_xlsx_cabecalho_filtros_e_fotos(dados, tmp_path):
    from openpyxl import load_workbook
    eid = semear_inventario(dados)
    inventario.ler(dados, eid, "01 - SALA CCI", 1001, "Fulano")
    foto_falsa(dados, eid, 1001, "https://x/1001.webp")
    inventario.registrar_sobra(dados, eid, "01 - SALA CCI", "VENTILADOR", None, "sem plaqueta", "https://x/s.webp", "Fulano")
    wb = load_workbook(inventario.exportar_xlsx(dados, eid, tmp_path / "a.xlsx", situacao="localizado", fotos=True))
    ws = wb["Bens"]
    linhas = list(ws.iter_rows(values_only=True))
    assert linhas[0][0] == "Inventário 2026" and linhas[1][0].startswith("Gerado em ")
    assert linhas[2][0] == "Todas as salas · Situação Localizado" and linhas[3][0] == "Total de bens: 1"
    assert linhas[4] == tuple(inventario.COLUNAS_XLSX) and linhas[5][0] == 1001 and len(linhas) == 6
    assert linhas[5][12] == '=_xlfn.IMAGE("https://x/1001.webp")' and ws.row_dimensions[6].height == 60
    assert wb["Sobras"].cell(row=3, column=7).value == '=_xlfn.IMAGE("https://x/s.webp")'
    ws = load_workbook(inventario.exportar_xlsx(dados, eid, tmp_path / "b.xlsx", fotos=True))["Bens"]
    assert ws.cell(row=6, column=13).value.startswith("=_xlfn") and ws.cell(row=7, column=13).value == "-"   # 1002 sem foto
    ws = load_workbook(inventario.exportar_xlsx(dados, eid, tmp_path / "c.xlsx"))["Bens"]
    assert ws.cell(row=6, column=13).value == "https://x/1001.webp" and ws.row_dimensions[6].height is None


def test_celula_foto_escapa_aspas_na_url():
    from openpyxl import Workbook
    ws = Workbook().active
    inventario._celula_foto(ws, 1, 1, 'https://x/foto "1".webp', True)
    assert ws.cell(row=1, column=1).value == '=_xlfn.IMAGE("https://x/foto ""1"".webp")'


def test_planilha_de_cadastros_exporta_e_importa_abas_de_inventario(dados, tmp_path):
    eid = semear_inventario(dados)
    inventario.ler(dados, eid, "01 - SALA CCI", 1001, "Fulano")
    inventario.registrar_sobra(dados, eid, "02 - SALA B", "VENTILADOR", None, "achado", "http://x/s.webp", "Fulano")
    from openpyxl import load_workbook
    caminho = db.exportar_cadastros(dados, tmp_path / "c.xlsx")
    wb = load_workbook(caminho)
    assert wb.sheetnames == ["responsaveis", "localizacoes", "pessoas", "atribuicoes", "inv_eventos", "inv_integrantes", "inv_salas", "inv_leituras", "inv_sobras", "inv_bens_encerrados", "inv_fotos"]
    assert list(wb["inv_leituras"].iter_rows(values_only=True))[1][:3] == (eid, 1001, "01 - SALA CCI")
    # editar: encerra o evento, acrescenta uma leitura migrada de outro sistema, e reimporta
    ws = wb["inv_eventos"]
    ws.cell(row=2, column=5, value="2026-01-31")                                       # encerrado_em só com data
    wb["inv_leituras"].append([eid, 2001, "02 - SALA B", "2026-01-20 10:00:00", "Antigo", "Regular", "", "migrado"])
    wb.save(tmp_path / "c2.xlsx")
    with open(tmp_path / "c2.xlsx", "rb") as f:
        r = db.importar_cadastros(dados, f)
    assert r["inv_leituras"] == 2 and r["inv_eventos"] == 1
    assert inventario.evento(dados, eid)["encerrado_em"] == "2026-01-31 00:00:00"
    assert {x["numero"]: x["situacao_inv"] for x in inventario.relatorio(dados, eid)}[2001] == "localizado"
    assert inventario.resumo(dados, eid)["sobras"] == 1


def test_planilha_sem_abas_de_inventario_nao_toca_nas_tabelas(dados, tmp_path):
    eid = semear_inventario(dados)
    inventario.ler(dados, eid, "01 - SALA CCI", 1001, "Fulano")
    from openpyxl import load_workbook
    wb = load_workbook(db.exportar_cadastros(dados, tmp_path / "c.xlsx"))
    for aba in list(inventario.ABAS):
        wb.remove(wb[aba])
    wb.save(tmp_path / "so4.xlsx")
    with open(tmp_path / "so4.xlsx", "rb") as f:
        r = db.importar_cadastros(dados, f)
    assert "inv_leituras" not in r and inventario.resumo(dados, eid)["lidos"] == 1


def test_planilha_com_abas_de_inventario_incompletas_e_recusada(dados, tmp_path):
    eid = semear_inventario(dados)
    inventario.ler(dados, eid, "01 - SALA CCI", 1001, "Fulano")
    from openpyxl import load_workbook
    wb = load_workbook(db.exportar_cadastros(dados, tmp_path / "c.xlsx"))
    for aba in ("inv_integrantes", "inv_salas", "inv_leituras", "inv_sobras"):   # sobra só inv_eventos
        wb.remove(wb[aba])
    wb.save(tmp_path / "parcial.xlsx")
    with open(tmp_path / "parcial.xlsx", "rb") as f, pytest.raises(db.ImportacaoInvalida) as e:
        db.importar_cadastros(dados, f)
    assert "incompletas" in str(e.value) and "inv_leituras" in str(e.value)
    assert inventario.resumo(dados, eid)["lidos"] == 1 and inventario.evento(dados, eid) is not None


def test_planilha_de_inventario_validacoes(dados, tmp_path):
    eid = semear_inventario(dados)
    from openpyxl import load_workbook
    wb = load_workbook(db.exportar_cadastros(dados, tmp_path / "c.xlsx"))
    wb["inv_leituras"].append([eid, 99999, "01 - SALA CCI", "2026-01-20 10:00:00", "Fulano", "", "", ""])   # bem inexistente
    wb["inv_leituras"].append([eid, 1001, "01 - SALA CCI", "x", "Fulano", "Ótimo", "", ""])                  # data e conservação
    wb["inv_salas"].append([77, "01 - SALA CCI"])                                                                # evento inexistente
    wb["inv_eventos"].append([2, "Outro aberto", None, "2026-02-01 00:00:00", None])                             # 2 abertos
    wb.save(tmp_path / "ruim.xlsx")
    with open(tmp_path / "ruim.xlsx", "rb") as f, pytest.raises(db.ImportacaoInvalida) as e:
        db.importar_cadastros(dados, f)
    msg = str(e.value)
    assert "99999" in msg and "conservação" in msg and "77" in msg and "aberto" in msg and "data" in msg
    assert inventario.resumo(dados, eid)["salas"] == 3                                                            # nada mudou


def test_titulo_do_xlsx_do_inventario_e_texto_literal(dados, tmp_path):
    from openpyxl import load_workbook
    eid = semear_inventario(dados)
    dados.execute("UPDATE inventario_eventos SET nome = '=1+1' WHERE id = ?", (eid,))
    dados.commit()
    wb = load_workbook(inventario.exportar_xlsx(dados, eid, tmp_path / "inv.xlsx"))
    assert wb["Bens"]["A1"].data_type == "s" and wb["Bens"]["A1"].value.startswith("=1+1")
    assert wb["Sobras"]["A1"].data_type == "s"


def test_encerrar_grava_snapshot_e_congela_o_evento(dados):
    eid = semear_inventario(dados)
    inventario.ler(dados, eid, "01 - SALA CCI", 1001, "Fulano")
    inventario.ler(dados, eid, "01 - SALA CCI", 1003, "Fulano")          # BAIXADO lido: entra no snapshot por ter leitura
    inventario.encerrar_evento(dados, eid)
    snap = {r[0] for r in dados.execute("SELECT numero FROM inventario_bens_encerrados WHERE evento_id = ?", (eid,))}
    assert snap == {1001, 1002, 1003, 1004, 2001, 2002}
    inventario.encerrar_evento(dados, eid)                                 # idempotente: não regrava
    assert dados.execute("SELECT count(*) FROM inventario_bens_encerrados").fetchone()[0] == 6
    # o export do SPW do ano seguinte muda `bens`; o evento encerrado não muda
    dados.execute("UPDATE bens SET localizacao = '02 - SALA B', descricao = 'CADEIRA NOVA' WHERE numero = 1001")
    dados.execute("DELETE FROM bens WHERE numero = 2002")
    dados.execute("INSERT INTO bens VALUES (3001,'ATIVO','TV','LG','EQUIPAMENTOS','01 - SALA CCI','01/01/2027',1,1)")
    dados.commit()
    assert [(s["localizacao"], s["total"], s["localizados"]) for s in inventario.salas(dados, eid)] == \
        [("01 - SALA CCI", 2, 1), ("02 - SALA B", 2, 0), ("99 - SEM MAPA", 1, 0)]
    r = {x["numero"]: x for x in inventario.relatorio(dados, eid)}
    assert set(r) == {1001, 1002, 1003, 1004, 2001, 2002}
    assert r[1001]["descricao"] == "CADEIRA" and r[1001]["situacao_inv"] == "localizado" and r[1003]["situacao_bem"] == "BAIXADO"
    assert inventario.resumo(dados, eid)["bens"] == 5
    assert [b["numero"] for b in inventario.bens_da_sala(dados, eid, "01 - SALA CCI")["bens"]] == [1001, 1002]
    # reabrir (planilha: limpar encerrado_em e reimportar sem a aba inv_bens_encerrados), bens mudam de novo,
    # encerrar outra vez: o snapshot velho não pode sobreviver (INSERT OR IGNORE ignoraria as linhas repetidas)
    dados.execute("UPDATE inventario_eventos SET encerrado_em = NULL WHERE id = ?", (eid,))
    dados.execute("UPDATE bens SET descricao = 'CADEIRA REFORMADA' WHERE numero = 1001")
    dados.execute("DELETE FROM bens WHERE numero = 1004")          # 1002 fica: tem atribuição com FK deferida
    dados.commit()
    inventario.encerrar_evento(dados, eid)
    snap2 = {r["numero"]: r["descricao"] for r in dados.execute(
        "SELECT numero, descricao FROM inventario_bens_encerrados WHERE evento_id = ?", (eid,))}
    assert snap2[1001] == "CADEIRA REFORMADA"
    assert 1004 not in snap2


def test_evento_encerrado_sem_snapshot_le_bens(dados):
    eid = semear_inventario(dados)
    dados.execute("UPDATE inventario_eventos SET encerrado_em = '2026-01-01 00:00:00' WHERE id = ?", (eid,))
    dados.commit()
    assert inventario.resumo(dados, eid)["bens"] == 5 and inventario._fonte_bens(dados, eid) == "bens"


def test_aba_inv_bens_encerrados_exporta_importa_e_valida(dados, tmp_path):
    from openpyxl import load_workbook
    eid = semear_inventario(dados)
    inventario.ler(dados, eid, "01 - SALA CCI", 1001, "Fulano")
    inventario.encerrar_evento(dados, eid)
    wb = load_workbook(db.exportar_cadastros(dados, tmp_path / "c.xlsx"))
    assert wb.sheetnames[-2:] == ["inv_bens_encerrados", "inv_fotos"]
    linhas = list(wb["inv_bens_encerrados"].iter_rows(values_only=True))
    assert linhas[0] == tuple(inventario.ABAS["inv_bens_encerrados"])
    assert linhas[1] == (eid, 1001, "ATIVO", "CADEIRA", "GIRATÓRIA", "MÓVEIS", "01 - SALA CCI") and len(linhas) == 6
    dados.execute("DELETE FROM bens WHERE numero = 1004")          # o SPW mudou; o snapshot importado preserva 1004
    dados.commit()
    with open(tmp_path / "c.xlsx", "rb") as f:
        r = db.importar_cadastros(dados, f)
    assert r["inv_bens_encerrados"] == 5
    assert {x["numero"] for x in inventario.relatorio(dados, eid)} == {1001, 1002, 1004, 2001, 2002}
    # aba ausente (planilha exportada por versão anterior, com 5 abas inv_*): snapshot não é tocado
    wb.remove(wb["inv_bens_encerrados"])
    wb.save(tmp_path / "cinco.xlsx")
    with open(tmp_path / "cinco.xlsx", "rb") as f:
        r = db.importar_cadastros(dados, f)
    assert "inv_bens_encerrados" not in r and dados.execute("SELECT count(*) FROM inventario_bens_encerrados").fetchone()[0] == 5
    # validações: evento aberto, número inválido, repetido
    wb = load_workbook(tmp_path / "c.xlsx")
    wb["inv_eventos"].append([2, "Outro aberto", None, "2026-02-01 00:00:00", None])   # evento 2 existe mas não está encerrado
    wb["inv_bens_encerrados"].append([2, 1001, "", "", "", "", ""])
    wb["inv_bens_encerrados"].append([eid, "abc", "", "", "", "", ""])
    wb["inv_bens_encerrados"].append([eid, 1001, "", "", "", "", ""])                  # repete a linha 2 (eid, 1001)
    wb.save(tmp_path / "ruim.xlsx")
    with open(tmp_path / "ruim.xlsx", "rb") as f, pytest.raises(db.ImportacaoInvalida) as ex:
        db.importar_cadastros(dados, f)
    msg = str(ex.value)
    assert "não está encerrado" in msg and "número inválido" in msg and "repetido" in msg


def test_andar():
    assert inventario.andar("07 - COAD - SALA DE REUNIÃO") == "07"
    assert inventario.andar("02 - CCOM") == "02"
    assert inventario.andar("TERMOS INDIVIDUAIS") == inventario.ANDAR_SEM
    assert inventario.andar("") == inventario.ANDAR_SEM
    assert inventario.andar(" - X") == inventario.ANDAR_SEM


def test_ler_lote_e_desfazer_leituras(dados):
    eid = semear_inventario(dados)
    r = inventario.ler_lote(dados, eid, "01 - SALA CCI", [1001, 2001, 99999], "Fulano")
    assert r == {"lidos": 2, "nao_encontrados": [99999]}
    assert {b["numero"]: b["situacao_inv"] for b in inventario.bens_da_sala(dados, eid, "01 - SALA CCI")["bens"]} == {1001: "localizado", 1002: "pendente"}
    assert inventario.resumo(dados, eid)["divergentes"] == 1
    with pytest.raises(db.ErroDeNegocio):
        inventario.ler_lote(dados, eid, "01 - SALA CCI", [1002], "Ninguém")
    foto_falsa(dados, eid, 1001, "http://x/1001.webp")
    assert inventario.desfazer_leituras(dados, eid, [1001, 2001, 1004]) == (["http://x/1001.webp"], 2)   # 1004 sem leitura: ignorado
    assert inventario.resumo(dados, eid)["lidos"] == 0 and inventario.resumo(dados, eid)["divergentes"] == 0
    assert inventario.desfazer_leituras(dados, eid, []) == ([], 0)
    assert inventario.desfazer_leituras(dados, eid, [1004]) == ([], 0)   # nenhuma leitura apagada
    inventario.encerrar_evento(dados, eid)
    with pytest.raises(db.ErroDeNegocio):
        inventario.desfazer_leituras(dados, eid, [1002])
    with pytest.raises(db.ErroDeNegocio):
        inventario.ler_lote(dados, eid, "01 - SALA CCI", [1002], "Fulano")


def test_relatorio_filtros_busca_e_ordem(dados):
    eid = semear_inventario(dados)
    inventario.ler(dados, eid, "01 - SALA CCI", 1001, "Fulano")
    inventario.ler(dados, eid, "02 - SALA B", 2001, "Beltrana")
    inventario.atualizar_leitura(dados, eid, 1001, conservacao="Ruim", quem_usa="José")
    foto_falsa(dados, eid, 1001, "https://x/1001.webp")
    inventario.atualizar_leitura(dados, eid, 2001, observacao="tela quebrada")
    num = lambda **f: [x["numero"] for x in inventario.relatorio(dados, eid, **f)]
    assert num() == [1001, 1002, 2001, 2002, 1004]
    assert num(integrante="Fulano") == [1001] and num(integrante="Beltrana") == [2001]
    assert num(conservacao="Ruim") == [1001] and num(conservacao="-") == [1002, 2001, 2002, 1004]
    assert num(foto="com") == [1001] and num(foto="sem") == [1002, 2001, 2002, 1004]
    assert num(busca="jose") == [1001] and num(busca="QUEBRADA tela") == [2001]
    assert num(busca="cadeira giratoria") == [1001] and num(busca="20") == [2001, 2002] and num(busca="   ") == num()
    assert num(ordem="numero", dir="desc") == [2002, 2001, 1004, 1002, 1001]
    assert num(ordem="descricao") == [1004, 1001, 2002, 2001, 1002]            # ARMÁRIO, CADEIRA, IMPRESSORA, MONITOR, NOTEBOOK
    assert num(ordem="integrante") == [2001, 1001, 1002, 2002, 1004]           # vazios sempre por último
    assert num(ordem="integrante", dir="desc") == [1001, 2001, 1002, 2002, 1004]
    assert num(ordem="inexistente") == num()
    assert num(localizacao="02 - SALA B", situacao="pendente") == [2002]
    assert inventario.contar_fotos(inventario.relatorio(dados, eid)) == 1
    assert inventario.descrever_filtros({}) == "Todas as salas"
    assert inventario.descrever_filtros({"localizacao": "01 - SALA CCI", "situacao": "divergente", "integrante": "Fulano",
                                         "conservacao": "-", "foto": "com", "busca": " cadeira ", "ordem": "numero"}) == \
        'Sala 01 - SALA CCI · Situação Divergente · Integrante Fulano · Conservação Não informada · Com foto · Busca "cadeira"'


def test_painel_dados(dados):
    eid = semear_inventario(dados)
    inventario.ler(dados, eid, "01 - SALA CCI", 1001, "Fulano")
    inventario.ler(dados, eid, "01 - SALA CCI", 2001, "Beltrana")      # divergente (é da SALA B)
    inventario.ler(dados, eid, "01 - SALA CCI", 1002, "Fulano")
    inventario.atualizar_leitura(dados, eid, 1001, conservacao="Ruim")
    p = inventario.painel(dados, eid)
    assert p["resumo"]["lidos"] == 2
    assert [(x["chave"], x["rotulo"], x["quantidade"]) for x in p["situacao"]] == \
        [("localizado", "Localizado", 2), ("divergente", "Divergente", 1), ("pendente", "Não localizado", 3)]
    assert [(x["chave"], x["quantidade"]) for x in p["integrantes"]] == [("Fulano", 2), ("Beltrana", 1)]
    assert [(x["chave"], x["rotulo"], x["quantidade"]) for x in p["conservacao"]] == [("Ruim", "Ruim", 1), ("-", "Não informada", 2)]
    assert p["andares"] == [{"andar": "01", "total": 2, "localizados": 2, "pendentes": 0, "divergentes": 1, "salas": 1},
                            {"andar": "02", "total": 2, "localizados": 0, "pendentes": 2, "divergentes": 0, "salas": 1},
                            {"andar": "99", "total": 1, "localizados": 0, "pendentes": 1, "divergentes": 0, "salas": 1}]
    assert p["salas_do_andar"] == []
    assert inventario.painel(dados, eid, "02")["salas_do_andar"] == \
        [{"localizacao": "02 - SALA B", "total": 2, "localizados": 0, "pendentes": 2, "divergentes": 0}]
    with pytest.raises(db.ErroDeNegocio):
        inventario.painel(dados, 999)


def test_painel_sala_sem_andar_vai_por_ultimo(dados):
    semear(dados)
    dados.execute("INSERT INTO bens VALUES (5001,'ATIVO','QUADRO','','MÓVEIS','TERMOS INDIVIDUAIS','01/01/2020',1,1)")
    dados.commit()
    eid = inventario.abrir_evento(dados, "Inv", None, ["Fulano"])
    assert [a["andar"] for a in inventario.painel(dados, eid)["andares"]] == ["01", "99", inventario.ANDAR_SEM]
    assert [x["chave"] for x in inventario.painel(dados, eid)["integrantes"]] == [] and inventario.painel(dados, eid)["conservacao"] == []


def test_migracao_foto_url_para_inventario_fotos(dados):
    """Banco anterior à Fase 3: inventario_leituras tinha foto_url. criar_esquema move para inventario_fotos
    (nfoto 1), apaga a coluna e é idempotente."""
    eid = semear_inventario(dados)
    inventario.ler(dados, eid, "01 - SALA CCI", 1001, "Fulano")
    inventario.ler(dados, eid, "01 - SALA CCI", 1002, "Fulano")
    dados.execute("ALTER TABLE inventario_leituras ADD COLUMN foto_url TEXT")
    dados.execute("UPDATE inventario_leituras SET foto_url = 'https://x/inventario/INV1_BEM_1001_1.webp' WHERE numero = 1001")
    dados.commit()
    db.criar_esquema(dados)
    assert "foto_url" not in db._colunas(dados, "inventario_leituras")
    assert [tuple(r) for r in dados.execute("SELECT evento_id, numero, nfoto, url FROM inventario_fotos")] == [(eid, 1001, 1, "https://x/inventario/INV1_BEM_1001_1.webp")]
    db.criar_esquema(dados)                                                          # de novo: nada muda
    assert dados.execute("SELECT count(*) FROM inventario_fotos").fetchone()[0] == 1
    dados.execute("DELETE FROM inventario_leituras WHERE numero = 1001")            # cascata
    assert dados.execute("SELECT count(*) FROM inventario_fotos").fetchone()[0] == 0


def test_adicionar_e_apagar_fotos_do_bem(dados):
    eid = semear_inventario(dados)
    assert inventario.pasta_do_evento(dados, eid) == "inventario2026"
    with pytest.raises(db.ErroDeNegocio):                                            # sem leitura
        foto_falsa(dados, eid, 1001)
    inventario.ler(dados, eid, "01 - SALA CCI", 1001, "Fulano")
    chaves = []
    f1 = inventario.adicionar_foto(dados, eid, 1001, lambda c: chaves.append(c) or "https://x/" + c)
    f2 = inventario.adicionar_foto(dados, eid, 1001, lambda c: chaves.append(c) or "https://x/" + c)
    assert chaves == ["inventario2026/1-1001.webp", "inventario2026/2-1001.webp"]
    assert [f["nfoto"] for f in f1] == [1] and [(f["nfoto"], f["url"]) for f in f2] == [(1, "https://x/inventario2026/1-1001.webp"), (2, "https://x/inventario2026/2-1001.webp")]
    assert inventario.fotos_do_bem_no_evento(dados, eid, 1001) == f2 and f2[0]["criado_em"]
    # envio falhou: nada gravado
    with pytest.raises(RuntimeError):
        inventario.adicionar_foto(dados, eid, 1001, lambda c: (_ for _ in ()).throw(RuntimeError("bucket")))
    assert len(inventario.fotos_do_bem_no_evento(dados, eid, 1001)) == 2
    # apagar a 2 e tirar outra → 3 (número nunca reaproveitado)
    assert inventario.apagar_foto(dados, eid, 1001, 2) == "https://x/inventario2026/2-1001.webp"
    assert inventario.apagar_foto(dados, eid, 1001, 2) is None
    f3 = foto_falsa(dados, eid, 1001)
    assert [f["nfoto"] for f in f3] == [1, 3] and chaves[-1] == "inventario2026/2-1001.webp"   # chaves só tem os 2 primeiros envios
    # bens_da_sala/relatorio: primeira foto + contagem; lista completa na sala
    b = {x["numero"]: x for x in inventario.bens_da_sala(dados, eid, "01 - SALA CCI")["bens"]}
    assert b[1001]["foto_url"] == "https://x/inventario2026/1-1001.webp" and b[1001]["n_fotos"] == 2 and [f["nfoto"] for f in b[1001]["fotos"]] == [1, 3]
    assert b[1002]["foto_url"] is None and b[1002]["n_fotos"] == 0 and b[1002]["fotos"] == []
    r = {x["numero"]: x for x in inventario.relatorio(dados, eid)}
    assert r[1001]["n_fotos"] == 2 and r[1001]["foto_url"].endswith("/1-1001.webp") and inventario.contar_fotos(r.values()) == 1
    assert [x["numero"] for x in inventario.relatorio(dados, eid, foto="com")] == [1001]
    # trazido de outra sala também traz a lista
    inventario.ler(dados, eid, "01 - SALA CCI", 2001, "Fulano")
    foto_falsa(dados, eid, 2001)
    t = inventario.bens_da_sala(dados, eid, "01 - SALA CCI")["trazidos"]
    assert "fotos" not in t[0]
    # desfazer leva todas as urls; cascata limpa a tabela
    urls, n = inventario.desfazer_leituras(dados, eid, [1001])
    assert sorted(urls) == ["https://x/inventario2026/1-1001.webp", "https://x/inventario2026/3-1001.webp"] and n == 1
    assert inventario.fotos_do_bem_no_evento(dados, eid, 1001) == []
    with pytest.raises(db.ErroDeNegocio):
        inventario.atualizar_leitura(dados, eid, 2001, foto_url="x")                 # campo saiu
    inventario.encerrar_evento(dados, eid)
    with pytest.raises(db.ErroDeNegocio):
        foto_falsa(dados, eid, 2001)
    with pytest.raises(db.ErroDeNegocio):
        inventario.apagar_foto(dados, eid, 2001, 1)


def test_fotos_do_bem_agrupadas_por_evento(dados):
    eid1 = semear_inventario(dados)
    inventario.ler(dados, eid1, "01 - SALA CCI", 1001, "Fulano")
    foto_falsa(dados, eid1, 1001); foto_falsa(dados, eid1, 1001)
    inventario.encerrar_evento(dados, eid1)
    eid2 = inventario.abrir_evento(dados, "Inventário 2027", None, ["Fulano"])
    inventario.ler(dados, eid2, "01 - SALA CCI", 1001, "Fulano")
    foto_falsa(dados, eid2, 1001)
    inventario.ler(dados, eid2, "01 - SALA CCI", 1002, "Fulano")                       # lido sem foto: não aparece
    g = inventario.fotos_do_bem(dados, 1001)
    assert [x["evento"] for x in g] == ["Inventário 2027", "Inventário 2026"]           # mais recente primeiro
    assert [[f["nfoto"] for f in x["fotos"]] for x in g] == [[1], [1, 2]]
    assert g[1]["encerrado_em"] and g[0]["encerrado_em"] is None and g[0]["lido_em"] and g[0]["evento_id"] == eid2
    assert g[0]["fotos"][0]["url"] == "https://x/inventario2027/1-1001.webp"
    assert inventario.fotos_do_bem(dados, 1002) == [] and inventario.fotos_do_bem(dados, 99999) == []


def test_abrir_evento_recusa_pasta_de_fotos_repetida(dados):
    eid = semear_inventario(dados)
    inventario.encerrar_evento(dados, eid)
    with pytest.raises(db.ErroDeNegocio) as e:
        inventario.abrir_evento(dados, "INVENTÁRIO 2026", None, ["Fulano"])            # mesma pasta: inventario2026
    assert "inventario2026" in str(e.value)
    assert inventario.abrir_evento(dados, "Inventário 2026 B", None, ["Fulano"])


def test_aba_inv_fotos_exporta_importa_e_valida(dados, tmp_path):
    from openpyxl import load_workbook
    eid = semear_inventario(dados)
    inventario.ler(dados, eid, "01 - SALA CCI", 1001, "Fulano")
    foto_falsa(dados, eid, 1001); foto_falsa(dados, eid, 1001)
    inventario.apagar_foto(dados, eid, 1001, 1)                                        # fica só a 2
    wb = load_workbook(db.exportar_cadastros(dados, tmp_path / "c.xlsx"))
    assert wb.sheetnames[-1] == "inv_fotos"
    assert list(wb["inv_leituras"].iter_rows(values_only=True))[0] == tuple(inventario.ABAS["inv_leituras"]) and "foto_url" not in inventario.ABAS["inv_leituras"]
    linhas = list(wb["inv_fotos"].iter_rows(values_only=True))
    assert linhas[0] == ("evento_id", "numero", "nfoto", "url", "criado_em") and linhas[1][:4] == (eid, 1001, 2, "https://x/inventario2026/2-1001.webp") and len(linhas) == 2
    with open(tmp_path / "c.xlsx", "rb") as f:
        r = db.importar_cadastros(dados, f)
    assert r["inv_fotos"] == 1 and [x["nfoto"] for x in inventario.fotos_do_bem_no_evento(dados, eid, 1001)] == [2]
    assert [x["nfoto"] for x in foto_falsa(dados, eid, 1001)] == [2, 3]                # contador continua do maior
    # aba ausente + inv_leituras com foto_url (planilha anterior à Fase 3): vira foto 1
    wb.remove(wb["inv_fotos"])
    ws = wb["inv_leituras"]
    ws.cell(row=1, column=9, value="foto_url")
    ws.cell(row=2, column=9, value="https://x/inventario/INV1_BEM_1001_1.webp")
    wb.save(tmp_path / "antiga.xlsx")
    with open(tmp_path / "antiga.xlsx", "rb") as f:
        r = db.importar_cadastros(dados, f)
    assert r["inv_fotos"] == 1 and inventario.fotos_do_bem_no_evento(dados, eid, 1001) [0]["url"] == "https://x/inventario/INV1_BEM_1001_1.webp"
    # aba ausente e sem foto_url: tabela fica vazia (a planilha é a fonte de verdade)
    ws.cell(row=2, column=9).value = None
    wb.save(tmp_path / "vazia.xlsx")
    with open(tmp_path / "vazia.xlsx", "rb") as f:
        r = db.importar_cadastros(dados, f)
    assert r["inv_fotos"] == 0 and inventario.fotos_do_bem_no_evento(dados, eid, 1001) == []
    # validações: leitura inexistente, nfoto inválido, repetido, url vazia, pasta repetida
    wb = load_workbook(tmp_path / "c.xlsx")
    wb["inv_fotos"].append([eid, 1002, 1, "https://x/a.webp", "2026-01-01 00:00:00"])      # 1002 não foi lido
    wb["inv_fotos"].append([eid, 1001, "x", "https://x/b.webp", "2026-01-01 00:00:00"])
    wb["inv_fotos"].append([eid, 1001, 2, "https://x/c.webp", "2026-01-01 00:00:00"])        # repete (eid, 1001, 2)
    wb["inv_fotos"].append([eid, 1001, 5, "", "2026-01-01 00:00:00"])
    wb["inv_eventos"].append([9, "INVENTARIO 2026", None, "2026-02-01 00:00:00", "2026-02-02 00:00:00"])
    wb.save(tmp_path / "ruim.xlsx")
    with open(tmp_path / "ruim.xlsx", "rb") as f, pytest.raises(db.ImportacaoInvalida) as ex:
        db.importar_cadastros(dados, f)
    msg = str(ex.value)
    assert "não tem leitura" in msg and "nfoto inválido" in msg and "repetida" in msg and "url vazia" in msg and "pasta de fotos repetida" in msg


def test_importar_cadastros_recusa_inv_sobras_sem_foto_url(dados, tmp_path):
    """`opcionais` da leitura da aba é só para inv_leituras (foto_url/fotos_seq); em inv_sobras foto_url
    continua obrigatória no cabeçalho — planilha sem ela não pode ser aceita com foto vazia em silêncio."""
    from openpyxl import load_workbook
    eid = semear_inventario(dados)
    inventario.registrar_sobra(dados, eid, "01 - SALA CCI", "VENTILADOR", None, "achado", "http://x/s.webp", "Fulano")
    wb = load_workbook(db.exportar_cadastros(dados, tmp_path / "c.xlsx"))
    ws = wb["inv_sobras"]
    coluna = inventario.ABAS["inv_sobras"].index("foto_url") + 1
    ws.delete_cols(coluna)
    wb.save(tmp_path / "sem_foto_url.xlsx")
    with open(tmp_path / "sem_foto_url.xlsx", "rb") as f, pytest.raises(db.ImportacaoInvalida) as ex:
        db.importar_cadastros(dados, f)
    msg = str(ex.value)
    assert "inv_sobras" in msg and "foto_url" in msg


def test_fotos_seq_sobrevive_a_exportar_e_importar(dados, tmp_path):
    """nfoto nunca reaproveitado mesmo depois de um ciclo de exportar/importar a planilha de cadastros:
    fotos_seq viaja na última coluna de inv_leituras."""
    from openpyxl import load_workbook
    eid = semear_inventario(dados)
    inventario.ler(dados, eid, "01 - SALA CCI", 1001, "Fulano")
    foto_falsa(dados, eid, 1001); foto_falsa(dados, eid, 1001)
    inventario.apagar_foto(dados, eid, 1001, 2)                                        # fica só a 1
    wb = load_workbook(db.exportar_cadastros(dados, tmp_path / "c.xlsx"))
    cabecalho = list(wb["inv_leituras"].iter_rows(values_only=True))[0]
    assert cabecalho[-1] == "fotos_seq"
    linha_1001 = next(l for l in wb["inv_leituras"].iter_rows(min_row=2, values_only=True) if l[1] == 1001)
    assert linha_1001[-1] == 2                                                         # maior nfoto já usado
    with open(tmp_path / "c.xlsx", "rb") as f:
        db.importar_cadastros(dados, f)
    f3 = foto_falsa(dados, eid, 1001)
    assert [x["nfoto"] for x in f3] == [1, 3]                                          # não voltou a ser 2


def test_apagar_no_bucket_antes_do_banco(dados):
    """apagar_foto / desfazer_leituras / excluir_sobra chamam `apagar(url)` antes do DELETE: se falhar, nada sai do banco."""
    eid = semear_inventario(dados)
    inventario.ler(dados, eid, "01 - SALA CCI", 1001, "Fulano")
    foto_falsa(dados, eid, 1001, "https://x/1.webp"); foto_falsa(dados, eid, 1001, "https://x/2.webp")

    def falha(url):
        raise RuntimeError("bucket fora")
    with pytest.raises(RuntimeError):
        inventario.apagar_foto(dados, eid, 1001, 1, apagar=falha)
    assert [f["nfoto"] for f in inventario.fotos_do_bem_no_evento(dados, eid, 1001)] == [1, 2]
    with pytest.raises(RuntimeError):
        inventario.desfazer_leituras(dados, eid, [1001], apagar=falha)
    assert inventario.resumo(dados, eid)["lidos"] == 1
    apagadas = []
    assert inventario.apagar_foto(dados, eid, 1001, 1, apagar=apagadas.append) == "https://x/1.webp" and apagadas == ["https://x/1.webp"]
    assert inventario.desfazer_leituras(dados, eid, [1001], apagar=apagadas.append) == (["https://x/2.webp"], 1) and apagadas[-1] == "https://x/2.webp"
    sid = inventario.registrar_sobra(dados, eid, "01 - SALA CCI", "VENTILADOR", "", "achado", "https://x/s.webp", "Fulano")
    with pytest.raises(RuntimeError):
        inventario.excluir_sobra(dados, eid, sid, apagar=falha)
    assert inventario.resumo(dados, eid)["sobras"] == 1
    inventario.excluir_sobra(dados, eid, sid, apagar=apagadas.append)
    assert apagadas[-1] == "https://x/s.webp" and inventario.resumo(dados, eid)["sobras"] == 0


def test_excluir_evento_apaga_tudo_com_fotos_primeiro(dados):
    eid = semear_inventario(dados)
    inventario.ler(dados, eid, "01 - SALA CCI", 1001, "Fulano")
    foto_falsa(dados, eid, 1001, "https://x/1.webp")
    foto_falsa(dados, eid, 1001, "https://x/2.webp")
    sid = inventario.registrar_sobra(dados, eid, "01 - SALA CCI", "VENT", "", "obs", "https://x/s.webp", "Fulano")
    inventario.registrar_sobra(dados, eid, "01 - SALA CCI", "SEM FOTO", "", "obs", "", "Fulano", exigir_foto=False)
    c = inventario.contagem_para_exclusao(dados, eid)
    assert c == {"leituras": 1, "fotos": 2, "sobras": 2, "sobras_com_foto": 1, "integrantes": 2, "bens_encerrados": 0, "salas": 3}
    assert inventario.urls_das_fotos(dados, eid) == ["https://x/1.webp", "https://x/2.webp", "https://x/s.webp"]
    with pytest.raises(db.ErroDeNegocio, match="não confere"):
        inventario.excluir_evento(dados, eid, "Inventario 2026")
    apagadas = []
    def apagar_falha(url):
        apagadas.append(url)
        if url.endswith("2.webp"):
            raise db.ErroDeNegocio("bucket fora")
    with pytest.raises(db.ErroDeNegocio, match="bucket fora"):
        inventario.excluir_evento(dados, eid, "Inventário 2026", apagar=apagar_falha)
    assert inventario.evento(dados, eid) and inventario.contagem_para_exclusao(dados, eid)["fotos"] == 2   # nada mudou
    apagadas.clear()
    inventario.encerrar_evento(dados, eid)                                       # encerrado também pode ser excluído
    assert inventario.contagem_para_exclusao(dados, eid)["bens_encerrados"] == 5
    e = inventario.excluir_evento(dados, eid, "  Inventário  2026 ", apagar=apagadas.append)
    assert e["id"] == eid and apagadas == ["https://x/1.webp", "https://x/2.webp", "https://x/s.webp"]
    assert inventario.evento(dados, eid) is None and inventario.eventos(dados) == []
    for t in ("inventario_leituras", "inventario_fotos", "inventario_sobras", "inventario_integrantes", "inventario_bens_encerrados", "inventario_salas"):
        assert dados.execute(f"SELECT count(*) FROM {t} WHERE evento_id = ?", (eid,)).fetchone()[0] == 0, t
    with pytest.raises(db.ErroDeNegocio, match="não encontrado"):
        inventario.excluir_evento(dados, eid, "x")


def test_esquema_normaliza_encerrado_em_vazio(dados):
    """Evento migrado com encerrado_em = '' (texto vazio) ficava invisível: nem aberto (IS NULL) nem encerrado (falsy)."""
    semear(dados)
    eid = inventario.abrir_evento(dados, "Migrado", "", ["Fulano"])
    dados.execute("UPDATE inventario_eventos SET encerrado_em = '' WHERE id = ?", (eid,))
    dados.commit()
    assert inventario.evento_aberto(dados) is None                     # o limbo
    db.criar_esquema(dados)                                             # roda a cada abertura do programa
    assert inventario.evento_aberto(dados)["id"] == eid


def test_criar_evento_nasce_fechado_e_a_chave_fecha_o_outro(dados):
    eid = semear_inventario(dados)                                             # abrir_evento: nasce aberto
    assert inventario.evento(dados, eid)["estado"] == "aberto"
    e2 = inventario.criar_evento(dados, "Inventário 2027", "", ["Fulano"])
    assert inventario.evento(dados, e2)["estado"] == "fechado" and inventario.evento_aberto(dados)["id"] == eid
    fechado = inventario.ligar_chave(dados, e2)
    assert fechado["id"] == eid
    assert inventario.evento_aberto(dados)["id"] == e2 and inventario.evento(dados, eid)["estado"] == "fechado"
    assert inventario.ligar_chave(dados, e2) is None                          # já estava aberto: nada muda
    inventario.desligar_chave(dados, e2)
    assert inventario.evento_aberto(dados) is None and inventario.evento_corrente(dados)["id"] == e2   # fechado mais recente
    inventario.desligar_chave(dados, e2)                                       # idempotente
    e3 = inventario.abrir_evento(dados, "Inventário 2028", "", ["Fulano"])    # sem erro mesmo com outros fechados
    assert inventario.evento_aberto(dados)["id"] == e3 and inventario.eventos(dados)[0]["id"] == e3
    assert [e["estado"] for e in map(lambda e: inventario.evento(dados, e["id"]), inventario.eventos(dados))] == ["aberto", "fechado", "fechado"]


def test_evento_fechado_bloqueia_leitura_e_reabrir_mantem(dados):
    eid = semear_inventario(dados)
    inventario.ler(dados, eid, "01 - SALA CCI", 1001, "Fulano")
    inventario.desligar_chave(dados, eid)
    with pytest.raises(db.ErroDeNegocio, match="fechado"):
        inventario.ler(dados, eid, "01 - SALA CCI", 1002, "Fulano")
    with pytest.raises(db.ErroDeNegocio, match="fechado"):
        inventario.registrar_sobra(dados, eid, "01 - SALA CCI", "X", "", "", "", "Fulano", exigir_foto=False)
    with pytest.raises(db.ErroDeNegocio, match="fechado"):
        inventario.atualizar_leitura(dados, eid, 1001, conservacao="Bom")
    inventario.editar_comissao(dados, eid, ["Fulano", "Beltrana"])           # comissão muda com o evento fechado
    inventario.ligar_chave(dados, eid)
    assert inventario.resumo(dados, eid)["lidos"] == 1
    inventario.ler(dados, eid, "01 - SALA CCI", 1002, "Fulano")
    assert inventario.resumo(dados, eid)["lidos"] == 2


def test_finalizar_fechado_congela_e_nao_reabre(dados):
    eid = semear_inventario(dados)
    inventario.desligar_chave(dados, eid)
    inventario.encerrar_evento(dados, eid)
    e = inventario.evento(dados, eid)
    assert e["estado"] == "finalizado" and e["suspenso_em"] is None and e["encerrado_em"]
    assert dados.execute("SELECT count(*) FROM inventario_bens_encerrados WHERE evento_id=?", (eid,)).fetchone()[0] == 5
    with pytest.raises(db.ErroDeNegocio, match="finalizado"):
        inventario.ligar_chave(dados, eid)
    with pytest.raises(db.ErroDeNegocio, match="finalizado"):
        inventario.desligar_chave(dados, eid)
    with pytest.raises(db.ErroDeNegocio, match="encerrado"):
        inventario.editar_comissao(dados, eid, ["Fulano"])
    assert inventario.evento_corrente(dados) is None


def test_esquema_acrescenta_suspenso_em_em_banco_antigo(dados):
    semear(dados)
    eid = inventario.abrir_evento(dados, "Antigo", "", ["Fulano"])
    dados.execute("ALTER TABLE inventario_eventos DROP COLUMN suspenso_em")
    dados.commit()
    db.criar_esquema(dados)
    assert "suspenso_em" in db._colunas(dados, "inventario_eventos")
    assert inventario.evento_aberto(dados)["id"] == eid                        # o aberto de antes continua aberto


def test_exportar_e_importar_cadastros_levam_suspenso_em(dados, tmp_path):
    from openpyxl import load_workbook
    eid = semear_inventario(dados)
    e2 = inventario.criar_evento(dados, "Preparado", "", ["Fulano"])          # fechado
    destino = tmp_path / "cadastros.xlsx"
    db.exportar_cadastros(dados, destino)
    ws = load_workbook(destino)["inv_eventos"]
    cab = [c.value for c in ws[1]]
    assert cab == ["id", "nome", "descricao", "aberto_em", "encerrado_em", "suspenso_em"]
    linhas = {r[0]: r for r in ws.iter_rows(min_row=2, values_only=True)}
    assert linhas[eid][5] is None and linhas[e2][5]                            # aberto sem suspenso_em; fechado com
    db.importar_cadastros(dados, destino)                                       # round-trip mantém os estados
    assert inventario.evento(dados, eid)["estado"] == "aberto" and inventario.evento(dados, e2)["estado"] == "fechado"


def test_importar_cadastros_rejeita_dois_abertos_e_aceita_planilha_antiga(dados, tmp_path):
    from openpyxl import load_workbook
    eid = semear_inventario(dados)
    e2 = inventario.criar_evento(dados, "Preparado", "", ["Fulano"])
    destino = tmp_path / "cadastros.xlsx"
    db.exportar_cadastros(dados, destino)
    wb = load_workbook(destino)
    ws = wb["inv_eventos"]
    for r in ws.iter_rows(min_row=2):
        r[5].value = None                                                       # os dois sem suspenso_em: dois abertos
    wb.save(destino)
    with pytest.raises(db.ImportacaoInvalida, match="mais de um evento aberto"):
        db.importar_cadastros(dados, destino)
    ws.delete_cols(6)                                                           # planilha antiga: sem a coluna
    for r in ws.iter_rows(min_row=2):
        if r[0].value == e2:
            r[4].value = "2026-01-01 10:00:00"                                  # e2 finalizado para sobrar um aberto
    wb.save(destino)
    db.importar_cadastros(dados, destino)
    assert inventario.evento(dados, eid)["estado"] == "aberto" and inventario.evento(dados, e2)["estado"] == "finalizado"
