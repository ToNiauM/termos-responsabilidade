"""TERMOS INDIVIDUAIS (localização do SPW dos bens com termo individual): nunca tem centro de custo,
não é pendência de mapeamento, continua como sala virtual do inventário; bem ativo nela sem pessoa é pendência própria."""
import io

import pytest
from openpyxl import Workbook

import db
from tests.conftest import semear


def _semear_individuais(conn):
    semear(conn)
    conn.execute("INSERT INTO bens VALUES (2003,'ATIVO','TABLET','','EQUIPAMENTOS','TERMOS INDIVIDUAIS','01/02/2023',900,800)")
    conn.execute("INSERT INTO bens VALUES (2004,'ATIVO','NOTEBOOK','','EQUIPAMENTOS','TERMOS INDIVIDUAIS','01/02/2023',900,800)")
    conn.execute("INSERT INTO bens VALUES (2005,'BAIXADO','MONITOR','','EQUIPAMENTOS','TERMOS INDIVIDUAIS','01/02/2023',900,800)")
    conn.execute("INSERT INTO atribuicoes VALUES ('ANA SILVA', 2003)")
    conn.commit()


@pytest.mark.parametrize("nome", ["TERMOS INDIVIDUAIS", "termos individuais", "  Termos   Indivíduais "])
def test_reconhece_sem_diferenciar_maiusculas_acentos_e_espacos(nome):
    assert db.localizacao_individual(nome)


@pytest.mark.parametrize("nome", ["", None, "01 - SALA CCI", "TERMOS", "TERMOS INDIVIDUAIS 2"])
def test_outras_localizacoes_nao_sao_individuais(nome):
    assert not db.localizacao_individual(nome)


def test_nao_vincula_a_centro(dados):
    _semear_individuais(dados)
    with pytest.raises(db.ErroDeNegocio, match="termo individual"):
        db.incluir_localizacao(dados, "TERMOS INDIVIDUAIS", "CCI")
    assert not any(db.localizacao_individual(l["localizacao"]) for l in db.localizacoes_mapeadas(dados))


def test_de_para_recusa(dados):
    _semear_individuais(dados)
    with pytest.raises(db.ErroDeNegocio, match="termo individual"):
        db.mover_localizacoes(dados, ["01 - SALA CCI", "TERMOS INDIVIDUAIS"], "CCI")


def test_vinculo_antigo_desfeito_ao_subir(dados):
    _semear_individuais(dados)
    dados.execute("INSERT INTO localizacoes VALUES ('TERMOS INDIVIDUAIS', 'CCI')")
    dados.commit()
    db.criar_esquema(dados)
    assert [l["localizacao"] for l in db.localizacoes_mapeadas(dados)] == ["01 - SALA CCI"]
    # o termo do centro deixa de receber o bem sem pessoa dessa localização
    assert 2004 not in [b["numero"] for b in db.bens_do_centro(dados, "CCI")]


def test_nao_e_pendencia_de_mapeamento(dados):
    _semear_individuais(dados)
    assert db.localizacoes_sem_centro(dados) == ["99 - SEM MAPA"]


def test_bens_sem_pessoa_viram_pendencia_propria(dados):
    _semear_individuais(dados)
    assert [b["numero"] for b in db.bens_individuais_sem_pessoa(dados)] == [2004]   # 2003 tem pessoa; 2005 baixado


def test_continua_sala_virtual_do_inventario(dados):
    """A comissão confere os bens de termo individual numa sala virtual com esse nome (decisão de 2026-09-23)."""
    _semear_individuais(dados)
    assert "TERMOS INDIVIDUAIS" in db.localizacoes_ativas(dados)


def test_aparece_na_lista_de_localizacoes_marcada_e_sem_centro(dados):
    _semear_individuais(dados)
    itens = db.listar_cadastros(dados, "localizacoes", {})["itens"]
    ti = [i for i in itens if i["localizacao"] == "TERMOS INDIVIDUAIS"]
    assert len(ti) == 1 and ti[0]["ccustos"] == "" and ti[0]["individual"] == 1
    assert all(i["individual"] == 0 for i in itens if i["localizacao"] != "TERMOS INDIVIDUAIS")


def test_aparece_na_lista_mesmo_com_todos_os_bens_atribuidos(dados):
    _semear_individuais(dados)
    dados.execute("INSERT INTO atribuicoes VALUES ('ANA SILVA', 2004)")
    itens = db.listar_cadastros(dados, "localizacoes", {})["itens"]
    assert "TERMOS INDIVIDUAIS" in [i["localizacao"] for i in itens]


def _planilha_cadastros(localizacoes):
    wb = Workbook()
    wb.remove(wb.active)
    ws = wb.create_sheet("responsaveis"); ws.append(["ccustos", "responsavel", "email", "matricula", "funcao"])
    ws.append(["CCI", "JAQUELINE PORTELA", "", "", ""])
    ws = wb.create_sheet("localizacoes"); ws.append(["localizacao", "ccustos"])
    for linha in localizacoes:
        ws.append(linha)
    ws = wb.create_sheet("pessoas"); ws.append(["nome", "email", "matricula"])
    ws = wb.create_sheet("atribuicoes"); ws.append(["nome", "numero"])
    arq = io.BytesIO(); wb.save(arq); arq.seek(0)
    return arq


def test_importacao_de_cadastros_ignora_e_avisa(dados):
    _semear_individuais(dados)
    r = db.importar_cadastros(dados, _planilha_cadastros([["01 - SALA CCI", "CCI"], ["TERMOS INDIVIDUAIS", "CCI"]]))
    assert r["localizacoes"] == 1 and r["localizacoes_ignoradas"] == ["TERMOS INDIVIDUAIS"]
    assert [l["localizacao"] for l in db.localizacoes_mapeadas(dados)] == ["01 - SALA CCI"]


# ---------------------------------------------------------------- telas (DSGov e Tabler)
@pytest.fixture(params=["dsgov", "tabler"])
def tela(request, cliente, dados):
    """Cliente logado com o cenário de termo individual, nos dois designs."""
    import app as modulo
    dados.execute("INSERT INTO bens VALUES (2004,'ATIVO','NOTEBOOK','','EQUIPAMENTOS','TERMOS INDIVIDUAIS','01/02/2023',900,800)")
    dados.commit()
    original = modulo.app.jinja_loader
    modulo.app.jinja_loader = modulo.carregador_templates(tabler=request.param == "tabler")
    modulo.app.jinja_env.cache.clear()
    yield cliente
    modulo.app.jinja_loader = original
    modulo.app.jinja_env.cache.clear()


def test_tela_localizacoes_marca_e_nao_oferece_vinculo(tela):
    html = tela.get("/cadastros/localizacoes").get_data(as_text=True)
    assert "Termo individual · sem centro de custo" in html
    assert "localizacao=TERMOS+INDIVIDUAIS" not in html          # nenhum "Vincular centro" para ela
    assert "bem(ns) de termo individual sem pessoa atribuída" in html and "2004" in html


def test_tela_importacao_avisa_bens_sem_pessoa(tela):
    html = tela.get("/upload").get_data(as_text=True)
    assert "bem(ns) de termo individual sem pessoa atribuída" in html and "NOTEBOOK" in html
