from openpyxl import Workbook

import db
import importar_planilhas as ip
from tests.test_db import CABECALHO


def planilhas(tmp_path):
    acervo = Workbook()
    ws = acervo.active
    ws.title = "acervo"
    ws.append(["numero", "situacao"])  # aba ignorada
    r = acervo.create_sheet("responsavel")
    r.append(["ccustos", "tratamento", "responsavel", "email", "matricula", "funcao"])
    r.append(["CCI", "Prezada", "JAQUELINE", "j@cfc", 46, "coordenadora"])
    c = acervo.create_sheet("ccustos")
    c.append(["localizacao", "ccustos"])
    c.append(["01 - SALA CCI", "CCI"])
    c.append(["02 - GAB", "TERMOS INDIVIDUAIS"])
    c.append(["03 - DEPOSITO", "SEPAT"])  # sem responsável cadastrado
    pa = tmp_path / "acervo.xlsx"
    acervo.save(pa)

    geral = Workbook()
    d = geral.active
    d.title = "dados"
    d.append(["Nome", "Patrimônio", "Situação", "Descrição", "Complemento", "Valor Atual"])
    d.append(["ANA SILVA", 1002, None, None, None, None])
    d.append(["CARLOS ", 1001, None, None, None, None])
    b = geral.create_sheet("base")
    b.append(CABECALHO)
    b.append([1001, "ATIVO", "CADEIRA", "", "MÓVEIS", "01 - SALA CCI", "31/12/1996", 75.94, 64.54])
    b.append([1002, "ATIVO", "NOTEBOOK", "DELL", "EQUIP", "02 - GAB", "06/12/2012", 3000, 1500])
    n = geral.create_sheet("nomes")
    n.append(["Nº", "responsavel"])
    n.append([1, "ANA SILVA"])
    n.append([2, "DANIEL"])
    pg = tmp_path / "geral.xlsx"
    geral.save(pg)
    return pa, pg


def test_migrar_popula_cinco_tabelas(dados, tmp_path):
    pa, pg = planilhas(tmp_path)
    resumo = ip.migrar(dados, pa, pg)
    assert resumo["bens"] == 2 and resumo["responsaveis"] == 2 and resumo["localizacoes"] == 2
    assert db.pessoas(dados) == ["ANA SILVA", "CARLOS", "DANIEL"]
    assert db.pessoa_do_bem(dados, 1001) == "CARLOS"
    assert db.responsavel(dados, "SEPAT")["responsavel"] == "(preencher)"
    assert [l["localizacao"] for l in db.localizacoes_mapeadas(dados)] == ["01 - SALA CCI", "03 - DEPOSITO"]
    assert "SEPAT" in resumo["avisos"][0]
