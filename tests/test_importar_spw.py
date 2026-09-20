"""Partes do robô do SPW que não precisam de rede: env, hash, conversão e orquestração."""
import sqlite3
from datetime import datetime

import pytest
from openpyxl import load_workbook

import db
import importar_spw as robo
from tests.conftest import semear

CABECALHO = ["Número Bem", "Situação", "Descrição", "Complemento", "Classificação Contábil",
             "Localização", "Data Entrada", None, "Valor Compra", "Valor Atual"]   # coluna vazia como no SPW


def linhas(*bens):
    """bens: (numero, situacao, descricao, localizacao). Demais colunas fixas."""
    return [CABECALHO] + [[n, s, d, "", "MÓVEIS", loc, "06/12/2012", None, 100.0, 50.0] for n, s, d, loc in bens]


BASE = [(1001, "ATIVO", "CADEIRA", "01 - SALA CCI"), (1002, "ATIVO", "NOTEBOOK", "01 - SALA CCI"),
        (1003, "BAIXADO", "MESA", "01 - SALA CCI"), (1004, "ATIVO", "ARMÁRIO", "99 - SEM MAPA")]


def test_ler_env_exige_arquivo_e_chaves(tmp_path):
    try:
        robo.ler_env(tmp_path / "nao-existe.env")
        assert False, "should have raised"
    except robo.RoboErro as e:
        msg = str(e)
        assert "spw.env" in msg  # label says spw.env
        assert "nao-existe.env" in msg  # but real path appears somewhere
    arq = tmp_path / "spw.env"
    arq.write_text("SPW_USUARIO=u\nSPW_SENHA=s=com=igual\n# comentário\nSPW_LOGIN_URL=http://l\n")
    with pytest.raises(robo.RoboErro, match="SPW_CONSULTA_URL"):
        robo.ler_env(arq)
    arq.write_text(arq.read_text() + "SPW_CONSULTA_URL=http://c\n")
    env = robo.ler_env(arq)
    assert env["SPW_SENHA"] == "s=com=igual" and env["SPW_CONSULTA_URL"] == "http://c"


def test_hash_linhas_estavel_e_sensivel():
    a = robo.hash_linhas(linhas(*BASE))
    assert a == robo.hash_linhas(linhas(*BASE)) and len(a) == 64
    assert a != robo.hash_linhas(linhas(*BASE[:3]))
    mudado = linhas(*BASE); mudado[1][2] = "CADEIRA NOVA"
    assert a != robo.hash_linhas(mudado)
    # espaços duplicados e datetime x texto normalizam do mesmo jeito que db._texto
    com_espacos = linhas(*BASE); com_espacos[1][2] = "CADEIRA "
    assert a == robo.hash_linhas(com_espacos)
    com_data = linhas(*BASE); com_data[1][6] = datetime(2012, 12, 6)
    assert a == robo.hash_linhas(com_data)


def test_linhas_para_xlsx_e_aceito_por_importar_bens(dados):
    semear(dados)
    buf = robo.linhas_para_xlsx(linhas(*BASE, (2001, "ATIVO", "MONITOR", "01 - SALA CCI")))
    ws = load_workbook(buf).active
    assert [c.value for c in ws[1]][:2] == ["Número Bem", "Situação"]
    buf.seek(0)
    resumo = db.importar_bens(dados, buf, nome_arquivo="teste")
    assert resumo["total"] == 5 and resumo["novos"] == 1


def test_executar_importa_depois_ve_sem_mudanca(dados):
    semear(dados)
    relogio = iter(["2026-09-18 04:00:00", "2026-09-19 04:00:00", "2026-09-20 04:00:00"])
    dados_spw = linhas(*BASE, (2001, "ATIVO", "MONITOR", "01 - SALA CCI"))
    r = robo.executar(dados, baixar=lambda: dados_spw, agora=lambda: next(relogio))
    assert r["resultado"] == "importado" and r["importacao_id"] and "1 novo" in r["mensagem"]
    imp = db.importacoes(dados)
    assert len(imp) == 1 and imp[0]["arquivo"] == "SPW automático" and imp[0]["novos"] == 1

    r = robo.executar(dados, baixar=lambda: dados_spw, agora=lambda: next(relogio))
    assert r["resultado"] == "sem_mudanca" and r["importacao_id"] is None
    assert len(db.importacoes(dados)) == 1                       # não repetiu
    e = db.execucoes_robo(dados)
    assert [x["resultado"] for x in e] == ["sem_mudanca", "importado"] and e[0]["hash"] == e[1]["hash"]
    assert e[1]["importacao_id"] == imp[0]["id"] and e[1]["iniciado_em"] == "2026-09-18 04:00:00"

    dados_spw[1][1] = "BAIXADO"
    r = robo.executar(dados, baixar=lambda: dados_spw, agora=lambda: next(relogio))
    assert r["resultado"] == "importado" and len(db.importacoes(dados)) == 2


def test_executar_repassa_env_a_baixar(dados):
    """Fila do site: env é a credencial de quem pediu; baixar(env) recebe exatamente esse dict."""
    semear(dados)
    recebido = []
    env = {"SPW_USUARIO": "maria.spw", "SPW_SENHA": "S3nha"}
    r = robo.executar(dados, baixar=lambda e: (recebido.append(e), linhas(*BASE))[1], env=env)
    assert recebido == [env] and r["resultado"] == "importado"


def test_executar_sem_env_chama_baixar_sem_argumento(dados):
    """cron / scripts/atualizar_base.sh: sem env, baixar() é chamado como hoje, sem argumento nenhum."""
    semear(dados)
    r = robo.executar(dados, baixar=lambda: linhas(*BASE))
    assert r["resultado"] == "importado"


def test_executar_erro_de_negocio_nao_altera_bens(dados):
    semear(dados)                                                # 1002 atribuído a ANA SILVA
    # troca 1002 por um bem novo: 4 linhas (>= 3,6 = 4*0,9) não aciona o piso de C1, mas ainda
    # deixa 1002 órfão da atribuição, então o erro de negócio dispara normalmente.
    sem_1002 = [b for b in BASE if b[0] != 1002] + [(2001, "ATIVO", "MONITOR", "01 - SALA CCI")]
    r = robo.executar(dados, baixar=lambda: linhas(*sem_1002))
    assert r["resultado"] == "erro" and "1002" in r["mensagem"] and "atribuídos" in r["mensagem"]
    assert dados.execute("SELECT count(*) FROM bens").fetchone()[0] == 4
    assert db.importacoes(dados) == [] and db.execucoes_robo(dados)[0]["resultado"] == "erro"
    assert db.ultimo_hash_robo(dados) is None


def test_executar_export_curto_nao_apaga_bens(dados):
    semear(dados)                                                 # 4 bens na base
    r = robo.executar(dados, baixar=lambda: linhas())             # só cabeçalho
    assert r["resultado"] == "erro" and r["mensagem"].startswith("export do SPW veio curto")
    assert dados.execute("SELECT count(*) FROM bens").fetchone()[0] == 4
    assert db.importacoes(dados) == []


def test_executar_export_curto_com_poucos_bens_tambem_e_erro(dados):
    semear(dados)                                                 # 4 bens na base
    # cabeçalho + 3 dos 4 bens semeados (some 1001): 3 < 4*0,9
    r = robo.executar(dados, baixar=lambda: linhas(*[b for b in BASE if b[0] != 1001]))
    assert r["resultado"] == "erro" and r["mensagem"].startswith("export do SPW veio curto")
    assert dados.execute("SELECT count(*) FROM bens").fetchone()[0] == 4
    assert db.importacoes(dados) == []


def test_executar_com_base_vazia_ignora_o_piso(dados):
    r = robo.executar(dados, baixar=lambda: linhas(*BASE))        # bens vazia: sem semear()
    assert r["resultado"] == "importado"
    assert dados.execute("SELECT count(*) FROM bens").fetchone()[0] == 4


def test_executar_cabecalho_mudou(dados):
    semear(dados)
    sem_valor = [l[:-1] for l in linhas(*BASE)]                  # some "Valor Atual"
    r = robo.executar(dados, baixar=lambda: sem_valor)
    assert r["resultado"] == "erro" and r["mensagem"] == "cabeçalho do SPW mudou: faltam Valor Atual"
    assert dados.execute("SELECT count(*) FROM bens").fetchone()[0] == 4


def test_executar_excecao_do_download_vira_erro(dados):
    def falha():
        raise TimeoutError("Timeout 60000ms exceeded")
    r = robo.executar(dados, baixar=falha)
    assert r["resultado"] == "erro" and r["mensagem"] == "Timeout 60000ms exceeded"
    assert db.execucoes_robo(dados)[0]["mensagem"] == "Timeout 60000ms exceeded"


def test_executar_mensagem_de_erro_e_limitada_a_500(dados):
    def falha():
        raise RuntimeError("x" * 900)
    assert len(robo.executar(dados, baixar=falha)["mensagem"]) == 500


def test_executar_export_vazio_e_erro(dados):
    r = robo.executar(dados, baixar=lambda: [])
    assert r["resultado"] == "erro" and r["mensagem"].startswith("cabeçalho do SPW mudou")


def test_executar_banco_travado_ao_registrar_nao_propaga(dados, monkeypatch):
    def travado(*a, **k):
        raise sqlite3.OperationalError("database is locked")
    monkeypatch.setattr(db, "registrar_execucao_robo", travado)
    assert robo.executar(dados, baixar=lambda: [])["resultado"] == "erro"
