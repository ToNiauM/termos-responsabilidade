"""Migração do inventário do sistema antigo (planilha INVENTARIO_SEPAT, aba BASE) para este sistema. Roda UMA vez.

Não escreve no banco. Lê a planilha antiga e a planilha de cadastros baixada em Atualizar base → Cadastros,
e gera uma cópia dela com as abas inv_* preenchidas. Você confere no Excel e carrega em Atualizar base →
Cadastros → Carregar, que valida tudo (tudo ou nada) antes de gravar.

    .venv/bin/python migrar_inventario.py "INVENTARIO_SEPAT .xlsx" cadastros.xlsx [saida.xlsx] [--pular-inexistentes]

Antes: carregue em Atualizar base → Bens um export ATUAL do SPW. Leituras de bens que não existem na base
são listadas e interrompem a geração (ou saem do arquivo, com --pular-inexistentes).

Regras: uma linha da BASE com local_verificado vira uma leitura; linha sem número vira sobra; um único
evento (aberto, id 1) com aberto_em = primeira leitura; salas = localizações ativas hoje + salas lidas.
"""
import sys
from collections import Counter
from datetime import datetime
from pathlib import Path

from openpyxl import load_workbook

import db
import inventario

EVENTO_NOME = "Inventário 2026"
EVENTO_DESCRICAO = "Migrado do sistema de inventário anterior (planilha INVENTARIO_SEPAT)."
SEM_FOTO, SEM_OBS = "(migrado sem foto)", "(migrado sem observação)"


def _data(valor) -> str | None:
    """'dd/mm/aaaa HH:MM:SS' (ou datetime do Excel) → 'AAAA-MM-DD HH:MM:SS'."""
    if isinstance(valor, datetime):
        return valor.strftime("%Y-%m-%d %H:%M:%S")
    t = db._texto(valor)
    for fmt in ("%d/%m/%Y %H:%M:%S", "%d/%m/%Y %H:%M", "%d/%m/%Y", "%Y-%m-%d %H:%M:%S", "%Y-%m-%d"):
        try:
            return datetime.strptime(t, fmt).strftime("%Y-%m-%d %H:%M:%S")
        except ValueError:
            pass
    return None


def ler_base(caminho: Path) -> list[dict]:
    wb = load_workbook(caminho, read_only=True, data_only=True)
    if "BASE" not in wb.sheetnames:
        raise SystemExit(f"{caminho}: aba BASE não encontrada (abas: {', '.join(wb.sheetnames)})")
    it = wb["BASE"].iter_rows(values_only=True)
    cab = [db._texto(c) for c in next(it)]
    faltando = [c for c in ("numero", "local_verificado", "estado_conservacao", "usuario", "data_hora") if c not in cab]
    if faltando:
        raise SystemExit(f"aba BASE sem coluna(s): {', '.join(faltando)}")
    linhas = [{"_linha": n, **dict(zip(cab, r))} for n, r in enumerate(it, start=2) if any(v not in (None, "") for v in r)]
    wb.close()
    return linhas


def transformar(conn, linhas: list[dict], pular_inexistentes: bool = False) -> tuple[dict, dict]:
    """Devolve ({aba: [linhas]}, resumo). Só linhas com local_verificado (lidas) entram."""
    t = db._texto
    leituras, sobras, problemas, inexistentes, datas = [], [], [], [], []
    vistos: dict[int, int] = {}
    fotos: dict[int, str | None] = {}   # numero -> imagem da leitura que "ganhou" (mais recente), ligada 1:1 à leitura
    quem_leu: set[str] = set()
    for r in linhas:
        loc = t(r.get("local_verificado"))
        if not loc:
            continue
        quando = _data(r.get("data_hora"))
        integrante = " ".join(t(r.get("usuario")).split())
        rot = f"BASE linha {r['_linha']}"
        if not quando:
            problemas.append(f"{rot}: data_hora inválida ({t(r.get('data_hora')) or 'vazia'})"); continue
        if not integrante:
            problemas.append(f"{rot}: usuário vazio"); continue
        datas.append(quando)
        quem_leu.add(integrante)
        num = db._numero(r.get("numero"))
        if num is None:                                                        # sobra (bem sem cadastro)
            desc = t(r.get("descricao"))
            if not desc:
                problemas.append(f"{rot}: sobra sem descrição"); continue
            sobras.append([1, loc, desc, t(r.get("complemento")) or None, t(r.get("observacao")) or SEM_OBS,
                           t(r.get("imagem")) or SEM_FOTO, integrante, quando])
            continue
        if num != int(num):
            problemas.append(f"{rot}: número inválido ({num})"); continue
        num = int(num)
        if not db.buscar_bem(conn, num):
            inexistentes.append(num); continue
        cons = t(r.get("estado_conservacao")) or None
        if cons and cons not in inventario.CONSERVACAO:
            problemas.append(f"{rot}: conservação inválida ({cons})"); cons = None
        linha = [1, num, loc, quando, integrante, cons, t(r.get("usuario_bem")) or None, t(r.get("observacao")) or None]
        imagem = t(r.get("imagem")) or None
        if num in vistos:                                                      # fica a leitura mais recente
            if quando > leituras[vistos[num]][3]:
                leituras[vistos[num]] = linha
                fotos[num] = imagem
            continue
        vistos[num] = len(leituras)
        leituras.append(linha)
        fotos[num] = imagem
    if inexistentes and not pular_inexistentes:
        raise SystemExit(f"{len(inexistentes)} leitura(s) de bens que não existem na base atual "
                         f"(ex.: {', '.join(map(str, sorted(inexistentes)[:10]))}).\n"
                         "Carregue um export atual do SPW em Atualizar base → Bens e rode de novo, "
                         "ou use --pular-inexistentes para deixá-las de fora.")
    salas = sorted(set(db.localizacoes_ativas(conn)) | {l[2] for l in leituras} | {s[1] for s in sobras})
    integrantes = sorted(quem_leu)
    fotos_leituras = [[1, l[1], 1, fotos[l[1]], l[3]] for l in leituras if fotos.get(l[1])]
    abas = {
        "inv_eventos": [[1, EVENTO_NOME, EVENTO_DESCRICAO, min(datas) if datas else db._agora(), None]],
        "inv_integrantes": [[1, n] for n in integrantes],
        "inv_salas": [[1, s] for s in salas],
        "inv_leituras": leituras,
        "inv_sobras": sobras,
        "inv_bens_encerrados": [],   # evento migrado nasce aberto; snapshot só existe depois de encerrado
        "inv_fotos": fotos_leituras,
    }
    resumo = {"leituras": len(leituras), "sobras": len(sobras), "salas": len(salas), "integrantes": integrantes,
              "inexistentes": sorted(inexistentes), "problemas": problemas,
              "divergentes": sum(1 for l in leituras if db.buscar_bem(conn, l[1])["localizacao"] != l[2]),
              "por_integrante": Counter(l[4] for l in leituras)}
    return abas, resumo


def gravar(cadastros: Path, saida: Path, abas: dict) -> Path:
    """Copia a planilha de cadastros e acrescenta (ou substitui) as abas inv_*."""
    wb = load_workbook(cadastros)
    for aba in ("responsaveis", "localizacoes", "pessoas", "atribuicoes"):
        if aba not in wb.sheetnames:
            raise SystemExit(f"{cadastros}: aba {aba} não encontrada; baixe a planilha em Atualizar base → Cadastros.")
    for aba, colunas in inventario.ABAS.items():
        if aba in wb.sheetnames:
            wb.remove(wb[aba])
        ws = wb.create_sheet(aba)
        ws.append(colunas)
        for linha in abas[aba]:
            ws.append(linha)
    wb.save(saida)
    return saida


def main(argv: list[str]) -> None:
    pular = "--pular-inexistentes" in argv
    args = [a for a in argv if not a.startswith("--")]
    if len(args) < 2:
        raise SystemExit(__doc__)
    antiga, cadastros = Path(args[0]), Path(args[1])
    saida = Path(args[2]) if len(args) > 2 else cadastros.with_name("cadastros_com_inventario.xlsx")
    conn = db.conectar()
    abas, r = transformar(conn, ler_base(antiga), pular)
    gravar(cadastros, saida, abas)
    print(f"Gerado: {saida}")
    print(f"  evento: {EVENTO_NOME} (aberto), aberto_em {abas['inv_eventos'][0][3]}")
    print(f"  leituras: {r['leituras']} ({r['divergentes']} divergentes) | sobras: {r['sobras']} | salas: {r['salas']}")
    for nome, n in sorted(r["por_integrante"].items()):
        print(f"    {nome}: {n}")
    if r["inexistentes"]:
        print(f"  DEIXADAS DE FORA: {len(r['inexistentes'])} leitura(s) de bens fora da base: "
              f"{', '.join(map(str, r['inexistentes'][:20]))}{' ...' if len(r['inexistentes']) > 20 else ''}")
    for p in r["problemas"]:
        print(f"  aviso: {p}")
    print("Próximo passo: confira o arquivo e carregue em Atualizar base → Cadastros → Carregar.")


if __name__ == "__main__":
    main(sys.argv[1:])
