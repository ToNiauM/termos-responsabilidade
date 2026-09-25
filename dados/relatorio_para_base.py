"""Converte o relatório exportado pelo sistema de inventário (relatorio_*.xlsx) para o formato da
planilha INVENTARIO_SEPAT (aba BASE), só para testar o fluxo de migração antes de ter a planilha do Drive."""
import sys
from openpyxl import Workbook, load_workbook

MAPA = {"Patrimonio": "numero", "Situacao": "situacao", "Descricao": "descricao", "Complemento": "complemento",
        "Classificacao Contabil": "classificacao_contabil", "Local Sistema": "localizacao_sistema",
        "Data Entrada": "data_entrada", "Valor Compra": "valor_compra", "Valor Atual": "valor_atual",
        "Local Inventariado": "local_verificado", "Conservacao": "estado_conservacao", "Integrante": "usuario",
        "Data/Hora Inventario": "data_hora", "Usuario do Bem": "usuario_bem"}
COLS = ["numero", "situacao", "descricao", "complemento", "classificacao_contabil", "localizacao_sistema", "data_entrada",
        "valor_compra", "valor_atual", "local_verificado", "estado_conservacao", "usuario", "data_hora", "usuario_bem",
        "imagem", "observacao", "contagem"]

src, dst = sys.argv[1], sys.argv[2]
ws = load_workbook(src, read_only=True, data_only=True).worksheets[0]
rows = ws.iter_rows(values_only=True)
for r in rows:
    if r and r[0] == "Foto":
        cab = list(r); break
idx = {MAPA[c]: i for i, c in enumerate(cab) if c in MAPA}
wb = Workbook(); out = wb.active; out.title = "BASE"; out.append(COLS)
n = 0
for r in rows:
    if not any(v not in (None, "") for v in r):
        continue
    out.append([r[idx[c]] if c in idx else None for c in COLS]); n += 1
wb.save(dst)
print(f"{dst}: {n} linhas na aba BASE")
